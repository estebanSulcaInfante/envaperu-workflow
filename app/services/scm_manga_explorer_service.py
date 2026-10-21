"""Read-only Fabricación manga explorer.

This projection deliberately keeps the physical manga as the row identity.  The
origin OT supplies the operational date (continuity is explained by the detail
view), while the effective work/assignment projection supplies the current OF,
OT and work context.
"""

from __future__ import annotations

import base64
import hashlib
import json
import re
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO
from types import SimpleNamespace

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.styles import Alignment
from sqlalchemy import String, and_, cast, exists, func, or_, select
from sqlalchemy.orm import aliased, selectinload

from app.models.registro import RegistroDiarioProduccion
from app.models.maquina import Maquina, TipoMaquina
from app.models.scm_ot import (
    ScmAnulacionPesajeManga,
    ScmAsignacionPersonalTrabajoOt,
    ScmCorreccionAsignacionManga,
    ScmCorreccionPesajeManga,
    ScmLoteArticulo,
    ScmManga,
    ScmPesajeManga,
    ScmReaperturaManga,
    ScmTramoMangaTrabajo,
    ScmTrabajoOt,
    ScmTrabajoColor,
)
from app.models.scm_production_orders import (
    ScmCorridaFabricacion,
    ScmOrdenFabricacion,
    ScmOrdenOperacion,
    ScmOrdenOperacionSalida,
    ScmPlanProduccion,
)
from app.models.molde import Molde
from app.models.receta_color import RecetaColorLinea, RecetaColorMaestra
from app.models.scm_catalogos import ScmMaterial
from app.services.scm_manga_assignment_projection import (
    effective_assignment,
    effective_assignment_for_segment,
    effective_work,
)
from app.services.scm_manga_detail_service import _number
from app.services.scm_production_observability_service import _iso
from app.services.scm_service_support import ScmServiceError, load_actor


DEFAULT_LIMIT = 25
MAX_LIMIT = 100
COLOR_HEX_PATTERN = re.compile(r"#[0-9a-fA-F]{6}")
OPEN_MANGA_STATES = {
    "PLANIFICADA", "PREETIQUETADA", "EN_ARMADO", "CONTINUIDAD_PENDIENTE", "EN_LLENADO",
}
QUANTITY_UN_SOURCES = {
    "RESPONSABLE_ARMADO",
    "CORRECCION_AUTORIZADA",
    "CIERRE_PARCIAL_SUPERVISADO",
}


def _parse_date(value, field):
    if value in (None, ""):
        return None
    try:
        return date.fromisoformat(str(value).strip())
    except (TypeError, ValueError) as error:
        raise ScmServiceError(
            "INVALID_MANGA_EXPLORER_DATE",
            f"El filtro {field} debe usar YYYY-MM-DD.",
            status_code=400,
        ) from error


def _filters(raw=None):
    raw = dict(raw or {})
    start = _parse_date(raw.get("fecha_desde") or raw.get("desde"), "fecha_desde")
    end = _parse_date(raw.get("fecha_hasta") or raw.get("hasta"), "fecha_hasta")
    if start and end and start > end:
        raise ScmServiceError(
            "INVALID_MANGA_EXPLORER_DATE_RANGE",
            "fecha_desde no puede ser posterior a fecha_hasta.",
            status_code=400,
        )
    try:
        limit = int(raw.get("limit", DEFAULT_LIMIT))
    except (TypeError, ValueError) as error:
        raise ScmServiceError(
            "INVALID_MANGA_EXPLORER_LIMIT",
            "limit debe ser un entero entre 1 y 100.",
            status_code=400,
        ) from error
    if not 1 <= limit <= MAX_LIMIT:
        raise ScmServiceError(
            "INVALID_MANGA_EXPLORER_LIMIT",
            "limit debe estar entre 1 y 100.",
            status_code=400,
        )
    tipo = str(raw.get("tipo_ot") or raw.get("tipo") or "FABRICACION").strip().upper()
    fecha_tipo = str(
        raw.get("fecha_tipo")
        or raw.get("criterio_fecha")
        or raw.get("modo_fecha")
        or "OT"
    ).strip().upper()
    if fecha_tipo in {"FINAL", "PESAJE", "PESAJE_FINAL", "FECHA_PESAJE_FINAL"}:
        fecha_tipo = "PESAJE_FINAL"
    elif fecha_tipo in {"OT", "FECHA_OT", "OPERATIVA"}:
        fecha_tipo = "OT"
    else:
        raise ScmServiceError(
            "INVALID_MANGA_EXPLORER_DATE_TYPE",
            "fecha_tipo debe ser OT o PESAJE_FINAL.",
            status_code=400,
        )
    return {
        "fecha_desde": start,
        "fecha_hasta": end,
        "tipo_ot": tipo,
        "fecha_tipo": fecha_tipo,
        "estado_manga": str(raw.get("estado_manga") or "").strip().upper() or None,
        "q": str(raw.get("q") or raw.get("busqueda") or "").strip().casefold() or None,
        "cursor": str(raw.get("cursor") or "").strip() or None,
        "limit": limit,
        "sort": str(raw.get("sort") or "FECHA_DESC").strip().upper(),
    }


def _scope(filters):
    return hashlib.sha256(json.dumps({
        key: value.isoformat() if isinstance(value, date) else value
        for key, value in filters.items() if key not in {"cursor", "limit"}
    }, sort_keys=True).encode()).hexdigest()[:16]


def _encode_cursor(*, operational_date, manga_id, as_of, scope):
    payload = {
        "d": operational_date.isoformat(),
        "id": int(manga_id),
        "as_of": _iso(as_of),
        "scope": scope,
    }
    return base64.urlsafe_b64encode(json.dumps(payload, separators=(",", ":")).encode()).decode().rstrip("=")


def _decode_cursor(value, *, scope):
    try:
        padded = value + "=" * (-len(value) % 4)
        payload = json.loads(base64.urlsafe_b64decode(padded).decode())
        if payload.get("scope") != scope:
            raise ValueError("scope")
        return date.fromisoformat(payload["d"]), int(payload["id"]), datetime.fromisoformat(payload["as_of"])
    except (TypeError, ValueError, KeyError, json.JSONDecodeError) as error:
        raise ScmServiceError(
            "INVALID_MANGA_EXPLORER_CURSOR",
            "El cursor de mangas no es válido para estos filtros.",
            status_code=400,
        ) from error


def _final_cursor_has_mutations(session, as_of):
    """Reject final-date cursors after append-only weighing events.

    The report is deliberately conservative: any later final/correction,
    reopening, or annulment in the fabrication scope invalidates the cursor.
    This avoids silently continuing with duplicate or missing rows when the
    effective final projection changed between pages.
    """
    return bool(session.scalar(select(or_(
        exists(select(ScmManga.id).where(ScmManga.anulada_at > as_of)),
        exists(select(ScmPesajeManga.id).where(ScmPesajeManga.created_at > as_of)),
        exists(select(ScmCorreccionPesajeManga.id).where(
            ScmCorreccionPesajeManga.estado == "APLICADA",
            ScmCorreccionPesajeManga.resolved_at.is_not(None),
            ScmCorreccionPesajeManga.resolved_at > as_of,
        )),
        exists(select(ScmReaperturaManga.id).where(ScmReaperturaManga.reabierta_at > as_of)),
        exists(select(ScmAnulacionPesajeManga.id).where(ScmAnulacionPesajeManga.anulada_at > as_of)),
    ))))


def _canonical_color_hex(value):
    if not isinstance(value, str):
        return None
    value = value.strip()
    return value.upper() if COLOR_HEX_PATTERN.fullmatch(value) else None


def _worker_payload(assignment):
    worker = getattr(assignment, "trabajador", None) if assignment else None
    if worker is None:
        return None
    return {
        "id": worker.id,
        "codigo": worker.codigo,
        "nombre": worker.nombre_corto or worker.nombre_completo,
    }


def _worker_projection(manga):
    """Return productive manga assignments once, including hand-offs.

    The weighing actor is intentionally not consulted: this report describes
    who was assigned to each physical work segment.  A correction without a
    segment anchor replaces the whole current projection; an anchored
    correction only replaces that segment.  Names are allowed to be absent so
    the UI/export can state the omission explicitly.
    """
    correction = getattr(manga, "correccion_asignacion", None)
    segments = list(getattr(manga, "tramos_trabajo", ()) or ())
    assignments = []
    incomplete = False
    correction_is_global = correction is not None and getattr(correction, "tramo_objetivo_id", None) is None
    latest_segment = max(segments, key=lambda item: item.secuencia) if segments else None
    global_correction_superseded = bool(
        correction_is_global
        and latest_segment is not None
        and latest_segment.created_at is not None
        and correction.created_at is not None
        and _comparable_datetime(latest_segment.created_at) > _comparable_datetime(correction.created_at)
    )
    if correction_is_global and not global_correction_superseded:
        assignment = effective_assignment(manga)
        if assignment is not None:
            assignments.append(assignment)
        else:
            incomplete = True
    else:
        for segment in segments:
            assignment = effective_assignment_for_segment(manga, segment)
            if assignment is not None:
                assignments.append(assignment)
            else:
                incomplete = True
        if not assignments and not segments:
            assignment = effective_assignment(manga)
            if assignment is not None:
                assignments.append(assignment)
            else:
                incomplete = True
    workers = []
    seen = set()
    for assignment in assignments:
        worker = getattr(assignment, "trabajador", None)
        worker_id = getattr(worker, "id", None)
        if worker_id in seen:
            continue
        if worker_id is not None:
            seen.add(worker_id)
        payload = _worker_payload(assignment)
        if payload is None:
            incomplete = True
        else:
            workers.append(payload)
    return {
        "items": workers,
        "cobertura": "INCOMPLETA" if incomplete and workers else ("SIN_ASIGNACION" if incomplete else "COMPLETA"),
    }


def _comparable_datetime(value):
    """Normalize SQLite-naive and PostgreSQL-aware UTC timestamps."""
    if value.tzinfo is not None and value.utcoffset() is not None:
        return value.astimezone(timezone.utc).replace(tzinfo=None)
    return value


def _worker_payloads(manga):
    """Backward-compatible list-only view of the productive assignment projection."""
    return _worker_projection(manga)["items"]


def _final_date_expression():
    """Correlated effective final date, portable across SQLite and PostgreSQL."""
    pesaje = aliased(ScmPesajeManga)
    correction = aliased(ScmCorreccionPesajeManga)
    correction_date = (
        select(correction.result_projection_json["fecha_local_pesaje"].as_string())
        .where(
            correction.pesaje_id == pesaje.id,
            correction.estado == "APLICADA",
        )
        .order_by(correction.id.desc())
        .limit(1)
        .scalar_subquery()
    )
    return (
        # Compare ISO dates as text so JSON projection and DATE columns have
        # the same type in both SQLite and PostgreSQL.
        select(func.coalesce(correction_date, cast(pesaje.fecha_local_pesaje, String)))
        .where(
            pesaje.manga_id == ScmManga.id,
            pesaje.estado == "VIGENTE",
        )
        .order_by(pesaje.id.desc())
        .limit(1)
        .correlate(ScmManga)
        .scalar_subquery()
    )


def _resource(ot):
    machine = getattr(ot, "maquina", None)
    if machine is None:
        return None
    return {
        "tipo": "MAQUINA",
        "id": machine.id,
        "codigo": machine.codigo,
        "nombre": machine.nombre,
    }


def _decimal(value):
    """Convert ORM numerics without losing a zero or raising on legacy NULL."""
    if value is None:
        return None
    try:
        return Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError):
        return None


def _material_recipe_context(recipe, *, recipe_id, recipe_hash, conflict=False):
    if conflict:
        return {
            "id": recipe_id,
            "hash": recipe_hash,
            "revision": None,
            "materials": [],
            "status": "CONFLICTO_REFERENCIA_RECETA",
            "fraccion_virgen_receta": None,
            "fraccion_virgen_estado": "CONFLICTO_REFERENCIA_RECETA",
        }
    if recipe is None:
        return {
            "id": recipe_id,
            "hash": recipe_hash,
            "revision": None,
            "materials": [],
            "status": "RECETA_AUSENTE",
            "fraccion_virgen_receta": None,
            "fraccion_virgen_estado": "RECETA_AUSENTE",
        }
    lines = [line for line in (getattr(recipe, "lineas", ()) or ())
             if getattr(line, "tipo_componente", None) == "MATERIA_PRIMA"]
    lines.sort(key=lambda line: (getattr(line, "orden", 0), getattr(line, "id", 0)))
    materials = []
    for line in lines:
        material = getattr(line, "material", None)
        materials.append({
            "tipo": getattr(getattr(material, "materia_prima", None), "tipo", None),
            "codigo": getattr(material, "codigo", None),
            "nombre": getattr(material, "nombre", None),
            "cantidad": _decimal(getattr(line, "cantidad", None)),
            "unidad": getattr(line, "unidad", None),
            "base_kg": _decimal(getattr(line, "base_kg", None)),
        })
    missing_type = any(item.get("tipo") is None for item in materials)
    complete_fraction = bool(materials) and all(
        item.get("tipo") in {"VIRGEN", "SEGUNDA"}
        and item.get("cantidad") is not None
        and item.get("cantidad") >= 0
        and item.get("unidad") == "FRACCION"
        for item in materials
    )
    total_fraction = sum(
        (item["cantidad"] for item in materials if item.get("cantidad") is not None),
        Decimal("0"),
    )
    virgin_fraction = sum(
        (item["cantidad"] for item in materials
         if item.get("tipo") == "VIRGEN" and item.get("cantidad") is not None),
        Decimal("0"),
    )
    fraction = (
        virgin_fraction
        if complete_fraction and abs(total_fraction - Decimal("1")) <= Decimal("0.000001")
        else None
    )
    return {
        "id": recipe.id,
        "hash": recipe_hash,
        "revision": recipe.revision,
        "materials": materials,
        "status": (
            "SIN_MATERIA_PRIMA" if not materials
            else ("TIPO_MATERIA_PRIMA_AUSENTE" if missing_type else "OK")
        ),
        "fraccion_virgen_receta": fraction,
        "fraccion_virgen_estado": (
            "OK" if fraction is not None else (
                "RECETA_SIN_PROPORCIONES_COMPLETAS" if materials
                else "SIN_MATERIA_PRIMA"
            )
        ),
    }


def _production_context(
    work, article, *, recipes=None, molds=None, fabrications=None, outputs_by_operation=None,
    machine_types=None, machines=None,
):
    """Project immutable work references used by both Excel views.

    This helper consumes the effective work already selected by the manga
    projection.  Recipe lines and mold masters are supplied by bulk caches so
    export does not turn into one query per manga.
    """
    color = getattr(work, "trabajo_color", None)
    corrida = getattr(color, "corrida", None)
    snapshot_recipe_id = getattr(color, "receta_revision_id_snapshot", None)
    corrida_recipe_id = getattr(corrida, "receta_revision_id", None)
    recipe_id = snapshot_recipe_id or corrida_recipe_id
    conflict = bool(snapshot_recipe_id and corrida_recipe_id and snapshot_recipe_id != corrida_recipe_id)
    recipe_hash = getattr(color, "receta_hash_snapshot", None) or getattr(corrida, "receta_hash", None)
    recipe = (recipes or {}).get(recipe_id)
    recipe_context = _material_recipe_context(
        recipe, recipe_id=recipe_id, recipe_hash=recipe_hash, conflict=conflict
    )

    operation = getattr(work, "orden_operacion", None)
    fabrication = (
        (fabrications or {}).get(getattr(operation, "id", None))
        if operation else None
    )
    if fabrication is None and operation is not None and fabrications is None:
        fabrication = getattr(operation, "fabricacion", None)
    mold_code = getattr(color, "molde_codigo_snapshot", None) or getattr(fabrication, "molde_id", None)
    mold = (molds or {}).get(mold_code)
    mold_name = getattr(mold, "nombre", None)
    mold_status = "OK" if mold_code and mold_name else ("MOLDE_AUSENTE" if not mold_code else "NOMBRE_MOLDE_NO_RESOLUBLE")
    mold_name_source = "MAESTRO_ACTUAL_SIN_SNAPSHOT" if mold_name else "AUSENTE"

    article_id = getattr(article, "id", None)
    operation_outputs = (
        (outputs_by_operation or {}).get(getattr(operation, "id", None), [])
        if outputs_by_operation is not None and operation is not None
        else (getattr(operation, "salidas", ()) or [])
    )
    outputs = [output for output in operation_outputs
               if article_id is not None and output.articulo_scm_id == article_id]
    corrida_id = getattr(corrida, "id", None)
    if corrida_id is not None:
        outputs = [output for output in outputs if output.corrida_fabricacion_id == corrida_id]
    output = outputs[0] if len(outputs) == 1 else None
    output_status = "OK" if output is not None else (
        "SALIDA_AMBIGUA" if len(outputs) > 1 else "SALIDA_AUSENTE"
    )
    quantity = _decimal(getattr(output, "cantidad_por_ciclo_snapshot", None)) if output else None
    unit_weight = _decimal(getattr(output, "peso_unitario_snapshot_g", None)) if output else None
    return {
        "tipo_materia_prima": ", ".join(dict.fromkeys(
            item["tipo"] for item in recipe_context["materials"] if item.get("tipo")
        )) or None,
        "materias_primas": recipe_context["materials"],
        "tipo_maquina": (machine_types or {}).get(getattr(getattr(work, "orden_trabajo", None), "id", None)),
        "maquina": (machines or {}).get(getattr(getattr(work, "orden_trabajo", None), "id", None)),
        "tipo_maquina_nombre": (
            (machines or {}).get(getattr(getattr(work, "orden_trabajo", None), "id", None)) or {}
        ).get("tipo_nombre"),
        "tipo_maquina_proceso": (
            (machines or {}).get(getattr(getattr(work, "orden_trabajo", None), "id", None)) or {}
        ).get("proceso"),
        "receta": recipe_context,
        "fraccion_virgen_receta": recipe_context["fraccion_virgen_receta"],
        "fraccion_virgen_estado": recipe_context["fraccion_virgen_estado"],
        "molde": {"codigo": mold_code, "nombre": mold_name, "status": mold_status, "nombre_fuente": mold_name_source},
        "salida": {
            "id": str(output.id) if output else None,
            "status": output_status,
            "cantidad_por_ciclo_snapshot": quantity,
            "peso_unitario_snapshot_g": unit_weight,
            "articulo_id": article_id,
        },
    }


def _bulk_production_context(session, rows):
    works = []
    seen = set()
    operation_ids = set()
    work_ot_ids = set()
    for manga, origin_ot in rows:
        work_ot_ids.add(origin_ot.id)
        work = effective_work(manga)
        if work is not None:
            work_ot_ids.add(getattr(work, "orden_trabajo_id", None))
            operation_ids.add(getattr(work, "orden_operacion_id", None))
    operation_ids.discard(None)
    work_ot_ids.discard(None)
    for manga, _origin_ot in rows:
        work = effective_work(manga)
        if work is not None and work.id not in seen:
            seen.add(work.id)
            works.append(work)
    fabrications = {}
    if operation_ids:
        fabrications = {
            row.orden_operacion_id: SimpleNamespace(molde_id=row.molde_id)
            for row in session.execute(
                select(
                    ScmOrdenFabricacion.orden_operacion_id,
                    ScmOrdenFabricacion.molde_id,
                ).where(ScmOrdenFabricacion.orden_operacion_id.in_(operation_ids))
            ).all()
        }
    recipe_ids = set()
    mold_codes = set()
    for work in works:
        color = getattr(work, "trabajo_color", None)
        corrida = getattr(color, "corrida", None)
        snapshot_id = getattr(color, "receta_revision_id_snapshot", None)
        corrida_id = getattr(corrida, "receta_revision_id", None)
        if snapshot_id:
            recipe_ids.add(snapshot_id)
        elif corrida_id:
            recipe_ids.add(corrida_id)
        code = getattr(color, "molde_codigo_snapshot", None)
        fabrication = fabrications.get(getattr(work, "orden_operacion_id", None))
        code = code or getattr(fabrication, "molde_id", None)
        if code:
            mold_codes.add(code)
    recipes = {}
    if recipe_ids:
        recipes = {
            recipe.id: recipe for recipe in session.scalars(
                select(RecetaColorMaestra)
                .where(RecetaColorMaestra.id.in_(recipe_ids))
                .options(
                    selectinload(RecetaColorMaestra.lineas)
                    .selectinload(RecetaColorLinea.material)
                    .selectinload(ScmMaterial.materia_prima)
                )
            ).all()
        }
    molds = {}
    if mold_codes:
        molds = {
            mold.codigo: mold for mold in session.scalars(
                select(Molde).where(Molde.codigo.in_(mold_codes))
            ).all()
        }
    outputs_by_operation = {}
    if operation_ids:
        for row in session.execute(
            select(
                ScmOrdenOperacionSalida.id,
                ScmOrdenOperacionSalida.orden_operacion_id,
                ScmOrdenOperacionSalida.corrida_fabricacion_id,
                ScmOrdenOperacionSalida.articulo_scm_id,
                ScmOrdenOperacionSalida.cantidad_por_ciclo_snapshot,
                ScmOrdenOperacionSalida.peso_unitario_snapshot_g,
            ).where(ScmOrdenOperacionSalida.orden_operacion_id.in_(operation_ids))
        ).all():
            output = SimpleNamespace(**row._mapping)
            outputs_by_operation.setdefault(output.orden_operacion_id, []).append(output)
    machine_types = {}
    machines = {}
    if work_ot_ids:
        for row in session.execute(
            select(
                RegistroDiarioProduccion.id.label("ot_id"),
                Maquina.id.label("machine_id"),
                Maquina.codigo.label("machine_code"),
                Maquina.nombre.label("machine_name"),
                TipoMaquina.nombre.label("machine_type_name"),
                TipoMaquina.proceso.label("machine_process"),
            )
            .join(Maquina, Maquina.id == RegistroDiarioProduccion.maquina_id)
            .join(TipoMaquina, TipoMaquina.id == Maquina.tipo_maquina_id)
            .where(RegistroDiarioProduccion.id.in_(work_ot_ids))
        ).all():
            machines[row.ot_id] = {
                "tipo": "MAQUINA", "id": row.machine_id,
                "codigo": row.machine_code, "nombre": row.machine_name,
                "tipo_nombre": row.machine_type_name,
                "proceso": row.machine_process,
            }
            machine_types[row.ot_id] = row.machine_process
    context = {}
    for manga, _origin_ot in rows:
        work = effective_work(manga)
        if work is not None:
            article = getattr(getattr(manga, "lote_articulo", None), "articulo", None)
            context[manga.id] = _production_context(
                work, article, recipes=recipes, molds=molds,
                fabrications=fabrications, outputs_by_operation=outputs_by_operation,
                machine_types=machine_types, machines=machines,
            )
    return context


def _upstream(work):
    operation = getattr(work, "orden_operacion", None) if work else None
    if operation is None:
        return {"op": None, "orden": None, "ordenes": [], "plan": None, "warnings": [
            {"codigo": "UPSTREAM_ORDEN_AUSENTE", "mensaje": "La manga no tiene OF resoluble."}
        ]}
    plan = getattr(operation, "plan_produccion", None)
    op = getattr(plan, "orden_produccion", None) if plan else None
    order = {
        "id": str(operation.id),
        "codigo": operation.codigo,
        "tipo": operation.tipo,
    }
    return {
        "op": {"id": str(op.id), "codigo": op.codigo} if op else None,
        "orden": order,
        "ordenes": [order],
        "plan": {"revision": plan.revision, "estado": plan.estado} if plan else None,
        "warnings": [] if plan is not None else [{
            "codigo": "UPSTREAM_PLAN_AUSENTE",
            "mensaje": "La OF no tiene un plan resoluble.",
        }],
    }


def _row(manga, origin_ot, *, vigente, visible_weights, production_context=None):
    work = effective_work(manga)
    effective_ot = getattr(work, "orden_trabajo", None) or origin_ot
    trabajo_color = getattr(work, "trabajo_color", None)
    color = getattr(trabajo_color, "color_nombre_snapshot", None)
    if color is None:
        color = manga.color_snapshot
    color_produccion = getattr(getattr(trabajo_color, "corrida", None), "color_produccion", None)
    color_hex = getattr(color_produccion, "hex_referencia", None)
    color_hex = _canonical_color_hex(color_hex)
    mold_code = getattr(trabajo_color, "molde_codigo_snapshot", None)
    article = getattr(getattr(manga, "lote_articulo", None), "articulo", None)
    production = production_context if production_context is not None else _production_context(work, article)
    machine_type = production.get("tipo_maquina") if production_context else (
        getattr(getattr(getattr(effective_ot, "maquina", None), "tipo_maquina", None), "nombre", None)
        if production_context is None else None
    )
    labels = getattr(manga, "etiquetas", ()) or ()
    label = (
        {
            "id": str(labels[-1].public_id),
            "tipo": labels[-1].tipo,
            "version": labels[-1].version,
            "estado": labels[-1].estado,
            "plantilla_version": labels[-1].plantilla_version,
            "generated_at": _iso(labels[-1].generated_at),
            "printed_at": _iso(labels[-1].printed_at),
        }
        if labels and visible_weights else None
    )
    assignment = effective_assignment(manga)
    worker_projection = _worker_projection(manga)
    workers = worker_projection["items"]
    return {
        "manga": {
            "public_id": str(manga.public_id),
            "codigo": manga.codigo,
            "tipo": manga.tipo,
            "articulo": {
                "codigo": manga.articulo_codigo_snapshot,
                "nombre": manga.articulo_nombre_snapshot,
                "sku_pieza_color": manga.pieza_color_sku_snapshot,
                "unidad_inventario": getattr(article, "unidad_inventario", None),
            },
            "color": color,
            "color_hex": color_hex,
            "contenedor": {
                "codigo": manga.tipo_contenedor_codigo_snapshot,
                "nombre": manga.tipo_contenedor_nombre_snapshot,
            },
            "created_at": _iso(manga.created_at),
            "estado_operativo": manga.estado,
                "estado_logistico": "RECIBIDA_ALMACEN" if manga.estado == "RECIBIDA" else "NO_INGRESADA",
            "cantidad_objetivo_un": float(manga.cantidad_asignada_un) if manga.cantidad_asignada_un is not None else None,
            "cantidad_confirmada_un": float(manga.cantidad_confirmada_un) if manga.cantidad_confirmada_un is not None else None,
            "responsable": _worker_payload(assignment),
            "maquinistas": workers,
            "maquinistas_cobertura": worker_projection["cobertura"],
            "etiqueta": label,
            "pesaje": vigente,
            "almacen": None,
        },
        "ot": {
            "public_id": str(effective_ot.public_id),
            "codigo": effective_ot.codigo_ot,
            "tipo": "FABRICACION",
            "estado_documental": effective_ot.estado,
            "estado_operativo": getattr(work, "estado", None) or effective_ot.estado,
            # The explorer's date is explicitly the physical/origin OT date.
            "fecha_operativa": origin_ot.fecha.isoformat(),
            "turno": origin_ot.turno,
            "version": effective_ot.version,
        },
        "recurso": production.get("maquina") or _resource(effective_ot),
        "produccion": {
            **production,
            "tipo_maquina": machine_type,
            "tipo_maquina_estado": "OK" if machine_type else "TIPO_MAQUINA_AUSENTE",
        },
        "upstream": _upstream(work),
        "trabajo": {
            "id": str(work.id) if work else None,
            "codigo": getattr(work, "codigo", None),
            "tipo": "COLOR" if work else None,
            "color": color,
            "molde_codigo": mold_code,
            "estado": getattr(work, "estado", None),
            "responsable": _worker_payload(assignment),
        },
        "maquinistas": workers,
        "maquinistas_cobertura": worker_projection["cobertura"],
        "visibilidad": {"pesaje": visible_weights, "inventario": False, "genealogia": False},
        "fecha_operativa_origen": origin_ot.fecha,
        "_manga_id": manga.id,
    }


def _collect(session, *, actor_id, raw_filters=None, paginate=True, include_production=False):
    filters = _filters(raw_filters)
    if filters["sort"] != "FECHA_DESC":
        raise ScmServiceError(
            "INVALID_MANGA_EXPLORER_SORT",
            "sort debe ser FECHA_DESC.",
            status_code=400,
        )
    if filters["tipo_ot"] not in {"FABRICACION", "OF"}:
        return {"items": [], "page": {"next_cursor": None, "limit": filters["limit"], "has_more": False}, "as_of": _iso(datetime.now(timezone.utc))}
    actor = load_actor(session, actor_id, capability="OT_VER")
    visible_weights = actor.tiene_capacidad("MANGA_PESAJE_VER")
    if filters["fecha_tipo"] == "PESAJE_FINAL" and not visible_weights:
        raise ScmServiceError(
            "MANGA_PESAJE_PERMISSION_REQUIRED",
            "Se requiere MANGA_PESAJE_VER para consultar por Fecha del pesaje final.",
            status_code=403,
        )
    date_type = filters["fecha_tipo"] if visible_weights else "OT"
    as_of = datetime.now(timezone.utc)
    scope = _scope(filters)
    cursor = _decode_cursor(filters["cursor"], scope=scope) if filters["cursor"] else None
    if cursor:
        as_of = cursor[2]
        if date_type == "PESAJE_FINAL" and _final_cursor_has_mutations(session, as_of):
            raise ScmServiceError(
                "STALE_MANGA_EXPLORER_CURSOR",
                "Los datos cambiaron; actualiza la consulta antes de continuar.",
                status_code=409,
            )
    stmt = select(ScmManga, RegistroDiarioProduccion).join(
        RegistroDiarioProduccion,
        RegistroDiarioProduccion.id == ScmManga.ot_id,
    ).outerjoin(
        ScmTrabajoOt,
        ScmTrabajoOt.id == ScmManga.trabajo_ot_id,
    ).outerjoin(
        ScmOrdenOperacion,
        ScmOrdenOperacion.id == ScmTrabajoOt.orden_operacion_id,
    ).where(
        RegistroDiarioProduccion.codigo_ot_sintetico.is_(False),
        RegistroDiarioProduccion.tipo_ot == "FABRICACION",
        ScmManga.created_at <= as_of,
    ).options(
        selectinload(ScmManga.etiquetas),
        selectinload(ScmManga.lote_articulo).selectinload(ScmLoteArticulo.articulo),
        selectinload(ScmManga.trabajo)
        .selectinload(ScmTrabajoOt.orden_trabajo),
        selectinload(ScmManga.trabajo)
        .selectinload(ScmTrabajoOt.orden_operacion)
        .selectinload(ScmOrdenOperacion.plan_produccion)
        .selectinload(ScmPlanProduccion.orden_produccion),
        selectinload(ScmManga.trabajo)
        .selectinload(ScmTrabajoOt.trabajo_color)
        .selectinload(ScmTrabajoColor.corrida)
        .selectinload(ScmCorridaFabricacion.color_produccion),
        selectinload(ScmManga.trabajo)
        .selectinload(ScmTrabajoOt.asignaciones_personal)
        .selectinload(ScmAsignacionPersonalTrabajoOt.trabajador),
        selectinload(ScmManga.asignacion_personal_trabajo)
        .selectinload(ScmAsignacionPersonalTrabajoOt.trabajador),
        selectinload(ScmManga.tramos_trabajo)
        .selectinload(ScmTramoMangaTrabajo.trabajo)
        .selectinload(ScmTrabajoOt.orden_trabajo),
        selectinload(ScmManga.tramos_trabajo)
        .selectinload(ScmTramoMangaTrabajo.trabajo)
        .selectinload(ScmTrabajoOt.orden_trabajo),
        selectinload(ScmManga.tramos_trabajo)
        .selectinload(ScmTramoMangaTrabajo.trabajo)
        .selectinload(ScmTrabajoOt.orden_operacion)
        .selectinload(ScmOrdenOperacion.plan_produccion)
        .selectinload(ScmPlanProduccion.orden_produccion),
        selectinload(ScmManga.tramos_trabajo)
        .selectinload(ScmTramoMangaTrabajo.trabajo)
        .selectinload(ScmTrabajoOt.trabajo_color)
        .selectinload(ScmTrabajoColor.corrida)
        .selectinload(ScmCorridaFabricacion.color_produccion),
        selectinload(ScmManga.tramos_trabajo)
        .selectinload(ScmTramoMangaTrabajo.asignacion_personal_trabajo)
        .selectinload(ScmAsignacionPersonalTrabajoOt.trabajador),
        selectinload(ScmManga.correccion_asignacion)
        .selectinload(ScmCorreccionAsignacionManga.destino_trabajo)
        .selectinload(ScmTrabajoOt.orden_trabajo),
        selectinload(ScmManga.correccion_asignacion)
        .selectinload(ScmCorreccionAsignacionManga.destino_trabajo)
        .selectinload(ScmTrabajoOt.orden_trabajo)
        .selectinload(RegistroDiarioProduccion.maquina),
        selectinload(ScmManga.correccion_asignacion)
        .selectinload(ScmCorreccionAsignacionManga.destino_trabajo)
        .selectinload(ScmTrabajoOt.orden_operacion)
        .selectinload(ScmOrdenOperacion.plan_produccion)
        .selectinload(ScmPlanProduccion.orden_produccion),
        selectinload(ScmManga.correccion_asignacion)
        .selectinload(ScmCorreccionAsignacionManga.destino_trabajo)
        .selectinload(ScmTrabajoOt.trabajo_color)
        .selectinload(ScmTrabajoColor.corrida)
        .selectinload(ScmCorridaFabricacion.color_produccion),
        selectinload(ScmManga.correccion_asignacion)
        .selectinload(ScmCorreccionAsignacionManga.destino_asignacion)
        .selectinload(ScmAsignacionPersonalTrabajoOt.trabajador),
        selectinload(ScmManga.correccion_asignacion)
        .selectinload(ScmCorreccionAsignacionManga.tramo_objetivo)
        .selectinload(ScmTramoMangaTrabajo.trabajo)
        .selectinload(ScmTrabajoOt.orden_trabajo),
        selectinload(ScmManga.correccion_asignacion)
        .selectinload(ScmCorreccionAsignacionManga.tramo_objetivo)
        .selectinload(ScmTramoMangaTrabajo.asignacion_personal_trabajo)
        .selectinload(ScmAsignacionPersonalTrabajoOt.trabajador),
    )
    sort_date = RegistroDiarioProduccion.fecha
    if date_type == "PESAJE_FINAL":
        # This correlated expression is part of the SQL scope, before LIMIT,
        # and resolves the canonical final or its applied correction.  It does
        # not inspect controls, sync timestamps, or the weighing actor.
        sort_date = _final_date_expression()
        stmt = stmt.where(sort_date.is_not(None), ScmManga.estado != "ANULADA")
    if filters["fecha_desde"]:
        date_value = filters["fecha_desde"].isoformat() if date_type == "PESAJE_FINAL" else filters["fecha_desde"]
        stmt = stmt.where(sort_date >= date_value)
    if filters["fecha_hasta"]:
        date_value = filters["fecha_hasta"].isoformat() if date_type == "PESAJE_FINAL" else filters["fecha_hasta"]
        stmt = stmt.where(sort_date <= date_value)
    if filters["estado_manga"]:
        stmt = stmt.where(ScmManga.estado == filters["estado_manga"])
    if filters["q"]:
        term = f"%{filters['q']}%"
        def work_match(work_model, operation_model, ot_model, color_model):
            return or_(
                work_model.codigo.ilike(term),
                operation_model.codigo.ilike(term),
                ot_model.codigo_ot.ilike(term),
                ot_model.maquina_codigo_snapshot.ilike(term),
                ot_model.maquina_nombre_snapshot.ilike(term),
                color_model.color_nombre_snapshot.ilike(term),
                color_model.molde_codigo_snapshot.ilike(term),
            )

        direct_work = aliased(ScmTrabajoOt)
        direct_operation = aliased(ScmOrdenOperacion)
        direct_ot = aliased(RegistroDiarioProduccion)
        direct_color = aliased(ScmTrabajoColor)
        direct_segment = aliased(ScmTramoMangaTrabajo)
        direct_correction = aliased(ScmCorreccionAsignacionManga)
        direct_context = exists(
            select(direct_work.id)
            .join(direct_operation, direct_operation.id == direct_work.orden_operacion_id)
            .join(direct_ot, direct_ot.id == direct_work.orden_trabajo_id)
            .outerjoin(direct_color, direct_color.trabajo_ot_id == direct_work.id)
            .where(
                direct_work.id == ScmManga.trabajo_ot_id,
                ~exists(select(direct_segment.id).where(
                    direct_segment.manga_id == ScmManga.id,
                ).correlate(ScmManga)),
                ~exists(select(direct_correction.id).where(
                    direct_correction.manga_id == ScmManga.id,
                ).correlate(ScmManga)),
                work_match(direct_work, direct_operation, direct_ot, direct_color),
            )
        )

        correction = aliased(ScmCorreccionAsignacionManga)
        correction_work = aliased(ScmTrabajoOt)
        correction_operation = aliased(ScmOrdenOperacion)
        correction_ot = aliased(RegistroDiarioProduccion)
        correction_color = aliased(ScmTrabajoColor)
        correction_target = aliased(ScmTramoMangaTrabajo)
        newer_target = aliased(ScmTramoMangaTrabajo)
        correction_active = or_(
            and_(
                correction.tramo_objetivo_id.is_(None),
                ~exists(select(newer_target.id).where(
                    newer_target.manga_id == correction.manga_id,
                    newer_target.created_at > correction.created_at,
                )),
            ),
            and_(
                correction.tramo_objetivo_id.is_not(None),
                ~exists(select(newer_target.id).where(
                    newer_target.manga_id == correction.manga_id,
                    newer_target.secuencia > correction_target.secuencia,
                )),
            ),
        )
        correction_context = exists(
            select(correction.id)
            .join(correction_work, correction_work.id == correction.destino_trabajo_ot_id)
            .join(correction_operation, correction_operation.id == correction_work.orden_operacion_id)
            .join(correction_ot, correction_ot.id == correction_work.orden_trabajo_id)
            .outerjoin(correction_color, correction_color.trabajo_ot_id == correction_work.id)
            .outerjoin(
                correction_target,
                correction_target.id == correction.tramo_objetivo_id,
            )
            .where(
                correction.manga_id == ScmManga.id,
                correction_active,
                work_match(correction_work, correction_operation, correction_ot, correction_color),
            )
        )

        latest_segment = aliased(ScmTramoMangaTrabajo)
        newer_segment = aliased(ScmTramoMangaTrabajo)
        latest_work = aliased(ScmTrabajoOt)
        latest_operation = aliased(ScmOrdenOperacion)
        latest_ot = aliased(RegistroDiarioProduccion)
        latest_color = aliased(ScmTrabajoColor)
        latest_correction = aliased(ScmCorreccionAsignacionManga)
        latest_correction_target = aliased(ScmTramoMangaTrabajo)
        latest_context = exists(
            select(latest_segment.id)
            .join(latest_work, latest_work.id == latest_segment.trabajo_ot_id)
            .join(latest_operation, latest_operation.id == latest_work.orden_operacion_id)
            .join(latest_ot, latest_ot.id == latest_work.orden_trabajo_id)
            .outerjoin(latest_color, latest_color.trabajo_ot_id == latest_work.id)
            .where(
                latest_segment.manga_id == ScmManga.id,
                ~exists(select(newer_segment.id).where(
                    newer_segment.manga_id == latest_segment.manga_id,
                    newer_segment.secuencia > latest_segment.secuencia,
                )),
                # A later real segment supersedes the correction overlay only
                # when its sequence is strictly newer than the anchored
                # segment.  Keep this correlation local to the manga so search
                # follows the same effective projection as the serializer.
                ~exists(
                    select(latest_correction.id)
                    .outerjoin(
                        latest_correction_target,
                        latest_correction_target.id == latest_correction.tramo_objetivo_id,
                    )
                    .where(
                        latest_correction.manga_id == ScmManga.id,
                        or_(
                            and_(
                                latest_correction.tramo_objetivo_id.is_not(None),
                                latest_correction_target.secuencia >= latest_segment.secuencia,
                            ),
                            and_(
                                latest_correction.tramo_objetivo_id.is_(None),
                                latest_correction.created_at >= latest_segment.created_at,
                            ),
                        ),
                    )
                ),
                work_match(latest_work, latest_operation, latest_ot, latest_color),
            )
        )
        origin_context = and_(
            ~exists(select(direct_segment.id).where(
                direct_segment.manga_id == ScmManga.id,
            ).correlate(ScmManga)),
            ~exists(select(direct_correction.id).where(
                direct_correction.manga_id == ScmManga.id,
            ).correlate(ScmManga)),
            or_(
                RegistroDiarioProduccion.codigo_ot.ilike(term),
                RegistroDiarioProduccion.maquina_codigo_snapshot.ilike(term),
                RegistroDiarioProduccion.maquina_nombre_snapshot.ilike(term),
            ),
        )
        stmt = stmt.where(or_(
            cast(ScmManga.public_id, String).ilike(term),
            ScmManga.codigo.ilike(term),
            ScmManga.articulo_codigo_snapshot.ilike(term),
            ScmManga.articulo_nombre_snapshot.ilike(term),
            ScmManga.pieza_color_sku_snapshot.ilike(term),
            ScmManga.color_snapshot.ilike(term),
            origin_context,
            direct_context,
            correction_context,
            latest_context,
        ))
    if cursor:
        cursor_date = cursor[0].isoformat() if date_type == "PESAJE_FINAL" else cursor[0]
        stmt = stmt.where(or_(
            sort_date < cursor_date,
            and_(sort_date == cursor_date, ScmManga.id < cursor[1]),
        ))
    stmt = stmt.order_by(sort_date.desc(), ScmManga.id.desc())
    if paginate:
        stmt = stmt.limit(filters["limit"] + 1)
    rows = session.execute(stmt).all()
    manga_ids = [manga.id for manga, _origin_ot in rows]
    vigente_by_manga = {}
    if visible_weights and manga_ids:
        vigente_by_manga = {
            item.manga_id: item
            for item in session.scalars(
                select(ScmPesajeManga).where(
                    ScmPesajeManga.manga_id.in_(manga_ids),
                    ScmPesajeManga.estado == "VIGENTE",
                )
            ).all()
        }

    correction_by_pesaje = {}
    if vigente_by_manga:
        corrections = session.scalars(
            select(ScmCorreccionPesajeManga)
            .where(
                ScmCorreccionPesajeManga.pesaje_id.in_(
                    [item.id for item in vigente_by_manga.values()]
                ),
                ScmCorreccionPesajeManga.estado == "APLICADA",
            )
            .order_by(ScmCorreccionPesajeManga.id.desc())
        ).all()
        correction_by_pesaje = {
            item.pesaje_id: item
            for item in corrections
            if item.pesaje_id not in correction_by_pesaje
        }

    def serialize_vigente(manga, pesaje):
        if pesaje is None or manga.estado == "ANULADA":
            return None
        correction = correction_by_pesaje.get(pesaje.id)
        if correction:
            projection = dict(correction.result_projection_json or {})
            proposed = correction.proposed_json or {}
            if "cantidad_confirmada" in proposed:
                projection["fuente_cantidad"] = "CORRECCION_AUTORIZADA"
            elif pesaje.fuente_cantidad in QUANTITY_UN_SOURCES:
                # A weight/date-only correction may preserve the immutable
                # weighing source, but it must not invent a new observation.
                projection["fuente_cantidad"] = pesaje.fuente_cantidad
            else:
                projection["fuente_cantidad"] = None
        else:
            projection = {
            "peso_bruto_kg": format(pesaje.peso_bruto_kg, "f"),
            "tara_kg": format(pesaje.tara_kg, "f"),
            "peso_fisico_neto_kg": format(pesaje.peso_fisico_neto_kg, "f"),
            "cantidad_confirmada": format(pesaje.cantidad_confirmada, "f"),
            "fuente_cantidad": pesaje.fuente_cantidad,
            "tara_fuente": pesaje.tara_fuente,
            "kg_produccion_ot": format(pesaje.kg_produccion_ot, "f"),
            "pesada_at": pesaje.pesada_at.isoformat(),
            "fecha_local_pesaje": pesaje.fecha_local_pesaje.isoformat(),
            }
        return {
            "pesaje_id": str(pesaje.public_id),
            "estado": "VIGENTE",
            "peso_bruto_kg": _number(projection.get("peso_bruto_kg")),
            "tara_kg": _number(projection.get("tara_kg")),
            "peso_fisico_neto_kg": _number(projection.get("peso_fisico_neto_kg")),
            "cantidad_confirmada": _number(projection.get("cantidad_confirmada")),
            "fuente_cantidad": projection.get("fuente_cantidad"),
            "tara_fuente": projection.get("tara_fuente"),
            "kg_produccion_ot": _number(projection.get("kg_produccion_ot")),
            "pesada_at": projection.get("pesada_at"),
            "fecha_local_pesaje": projection.get("fecha_local_pesaje"),
        }

    items = []
    production_context_by_manga = _bulk_production_context(session, rows) if include_production else {}
    for manga, origin_ot in rows:
        item = _row(
            manga,
            origin_ot,
            vigente=serialize_vigente(manga, vigente_by_manga.get(manga.id)) if visible_weights else None,
            visible_weights=visible_weights,
            production_context=(production_context_by_manga.get(manga.id) if include_production else {}),
        )
        item["_sort_date"] = (
            date.fromisoformat(item["manga"]["pesaje"]["fecha_local_pesaje"])
            if date_type == "PESAJE_FINAL" and item["manga"].get("pesaje")
            else origin_ot.fecha
        )
        items.append(item)
    # READ COMMITTED can observe a new append-only event after the page query
    # began. Re-check the same cursor watermark after hydration/serialization
    # so we do not return a page or issue a next cursor from a mixed scope.
    if cursor and date_type == "PESAJE_FINAL" and _final_cursor_has_mutations(session, as_of):
        raise ScmServiceError(
            "STALE_MANGA_EXPLORER_CURSOR",
            "Los datos cambiaron; actualiza la consulta antes de continuar.",
            status_code=409,
        )
    has_more = paginate and len(items) > filters["limit"]
    visible = items[: filters["limit"]] if paginate else items
    next_cursor = None
    if has_more and visible:
        last = visible[-1]
        next_cursor = _encode_cursor(
            operational_date=last["_sort_date"],
            manga_id=last["_manga_id"],
            as_of=as_of,
            scope=scope,
        )
    for item in visible:
        item.pop("fecha_operativa_origen", None)
        item.pop("_sort_date", None)
        item.pop("_manga_id", None)
    return {
        "items": visible,
        "page": {"next_cursor": next_cursor, "limit": filters["limit"], "has_more": bool(has_more)},
        "as_of": _iso(as_of),
        "fecha_tipo": filters["fecha_tipo"],
        "fecha_tipo_aplicado": date_type,
    }


def list_manga_explorer(session, *, actor_id, filters=None):
    return _collect(session, actor_id=actor_id, raw_filters=filters, paginate=True)


def _safe_cell(value):
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    return value


def _human_label(value, fallback="Sin estado"):
    if value in (None, ""):
        return fallback
    return str(value).replace("_", " ").capitalize()


def _export_rows(payload):
    detail_rows = []
    groups = {}
    for item in payload["items"]:
        manga = item["manga"]
        ot = item["ot"]
        article = manga["articulo"]
        resource = item.get("recurso") or {}
        weighing = manga.get("pesaje") or {}
        weights_visible = bool(item.get("visibilidad", {}).get("pesaje"))
        final_kg = weighing.get("peso_fisico_neto_kg") if weights_visible and weighing.get("estado") == "VIGENTE" else None
        weighing_state = (
            "RESTRINGIDO" if not weights_visible
            else ("FINAL_VIGENTE" if final_kg is not None else "SIN_PESAJE_VIGENTE")
        )
        worker_names = [
            worker.get("nombre") or worker.get("codigo") or "Sin nombre"
            for worker in (item.get("maquinistas") or [])
            if worker
        ]
        worker_text = ", ".join(dict.fromkeys(worker_names)) or "Sin maquinista asignado"
        final_date = weighing.get("fecha_local_pesaje") if weights_visible and weighing.get("estado") == "VIGENTE" else None
        final_date_value = date.fromisoformat(final_date) if final_date else None
        gross = weighing.get("peso_bruto_kg") if final_kg is not None else None
        tare = weighing.get("tara_kg") if final_kg is not None else None
        quantity_source = weighing.get("fuente_cantidad")
        quantity_unit = article.get("unidad_inventario")
        quantity_is_un = quantity_unit in {"UN", "UND", "UNIDAD"}
        quantity_is_observed = quantity_source in QUANTITY_UN_SOURCES
        confirmed_quantity = (
            weighing.get("cantidad_confirmada")
            if final_kg is not None and quantity_is_un and quantity_is_observed
            else None
        )
        final_date_state = (
            "Fecha final restringida" if not weights_visible
            else ("Fecha final vigente" if final_date else "Sin pesaje final vigente")
        )
        worker_coverage = item.get("maquinistas_cobertura") or ("COMPLETA" if worker_names else "SIN_ASIGNACION")
        production = item.get("produccion") or {}
        recipe = production.get("receta") or {}
        materials = production.get("materias_primas") or []
        mold = production.get("molde") or {}
        output = production.get("salida") or {}
        article_name = article.get("nombre") or "Pieza sin nombre"
        color_name = manga.get("color") or "Color no indicado"
        piece_color = f"{article_name} · {color_name}"
        denominator = (_decimal(output.get("cantidad_por_ciclo_snapshot")) or Decimal("0")) * (
            _decimal(output.get("peso_unitario_snapshot_g")) or Decimal("0")
        )
        if final_kg is None:
            cycles = None
            cycles_state = "SIN_KG_FINAL"
        elif output.get("status") != "OK":
            cycles = None
            cycles_state = output.get("status") or "SALIDA_AUSENTE"
        elif denominator <= 0:
            cycles = None
            cycles_state = "BASE_CICLO_INVALIDA"
        else:
            cycles = float((_decimal(final_kg) or Decimal("0")) * Decimal("1000") / denominator)
            cycles_state = "ESTIMADO"
        material_text = [item.get("nombre") or item.get("codigo") or "Material sin nombre" for item in materials]
        recipe_identity = (
            f"ID={recipe.get('id')};REV={recipe.get('revision')};HASH={recipe.get('hash')}"
            if recipe.get("id") is not None or recipe.get("revision") is not None or recipe.get("hash")
            else recipe.get("status")
        )
        production_status = "; ".join(
            status for status in (recipe.get("status"), mold.get("status"), output.get("status"))
            if status and status != "OK"
        ) or "OK"
        if not weights_visible:
            weighing_observation = "Peso no autorizado"
        elif final_kg is not None:
            weighing_observation = "Pesaje final vigente"
        elif manga.get("estado_operativo") == "ANULADA":
            weighing_observation = "Manga anulada sin pesaje final"
        else:
            weighing_observation = "Pendiente de pesaje final"
        if confirmed_quantity is not None:
            quantity_observation = None
        elif final_kg is not None and (not quantity_is_un or not quantity_is_observed):
            quantity_observation = "Cantidad por bulto no acreditada como UN"
        else:
            quantity_observation = "Cantidad por bulto no registrada"
        row_observations = [observation for observation in (weighing_observation, quantity_observation) if observation]
        fraction = production.get("fraccion_virgen_receta")
        fraction = float(fraction) if fraction is not None else None
        row = {
            "Fecha de OT": date.fromisoformat(ot["fecha_operativa"]),
            "Fecha del pesaje final": final_date_value,
            "Estado fecha final": final_date_state,
            "Turno": ot.get("turno"), "Máquina": resource.get("nombre"), "Código máquina": resource.get("codigo"),
            "N°OF": (item.get("upstream", {}).get("orden") or {}).get("codigo"), "N°OT": ot.get("codigo"),
            "Pieza": article_name, "Código pieza": article.get("codigo"), "Color": color_name,
            "Pieza-Color": piece_color,
            "Maquinistas": worker_text,
            "Cobertura maquinistas": worker_coverage,
            "Manga": manga.get("codigo"), "ID manga": manga.get("public_id"), "Estado": manga.get("estado_operativo"),
            "KG final vigente": final_kg, "Estado pesaje": weighing_state,
            "Tipo máquina": production.get("tipo_maquina"),
            "Nombre tipo máquina": production.get("tipo_maquina_nombre"),
            "Proceso máquina": production.get("tipo_maquina_proceso") or production.get("tipo_maquina"),
            "Tipo materia prima": production.get("tipo_materia_prima"),
            "Receta ID": recipe.get("id"),
            "Receta revisión": recipe.get("revision"),
            "Receta hash": recipe.get("hash"),
            "Identidad receta": recipe_identity,
            "Molde": mold.get("nombre") or mold.get("codigo"),
            "Código molde": mold.get("codigo"),
            "Fuente nombre molde": mold.get("nombre_fuente"),
            "Ciclos aproximados": cycles,
            "Estado ciclos aproximados": cycles_state,
            "Estado datos producción": production_status,
            "Estado receta": recipe.get("status"),
            "Estado molde": mold.get("status"),
            "Estado salida": output.get("status"),
            "Base técnica ciclos": output.get("id"),
            "Peso real (kg)": final_kg,
            "Peso bruto (kg)": gross,
            "Tara (kg)": tare,
            "Cantidad por bulto": confirmed_quantity,
            "Bultos pesados": 1 if final_kg is not None else 0,
            "Fracción virgen (receta)": fraction,
            "Estado fracción virgen": production.get("fraccion_virgen_estado"),
            "Hora del pesaje": weighing.get("pesada_at") if final_kg is not None else None,
            "Fuente cantidad": weighing.get("fuente_cantidad") if final_kg is not None else None,
            "Situación del pesaje": weighing_observation,
            "Observación": "; ".join(row_observations),
            "Estado operativo": _human_label(manga.get("estado_operativo")),
            "_estado_operativo_code": manga.get("estado_operativo"),
        }
        for index, material in enumerate(material_text, start=1):
            row[f"Material {index}"] = material
        detail_rows.append(row)
        key = tuple(row[field] for field in (
            "Fecha de OT", "Turno", "Máquina", "Código máquina", "N°OF", "Pieza",
            "Código pieza", "Color", "Receta ID", "Receta revisión", "Receta hash", "Código molde", "Base técnica ciclos",
        ))
        row["_group_key"] = key
        summary = groups.setdefault(key, {
            field: row[field]
            for field in ("Fecha de OT", "Turno", "Máquina", "Código máquina", "N°OF", "Pieza", "Código pieza", "Color",
                          "Tipo máquina", "Tipo materia prima", "Receta ID", "Receta revisión", "Receta hash", "Identidad receta", "Molde", "Código molde",
                          "Fuente nombre molde", "Estado datos producción", "Base técnica ciclos")
        })
        summary["Pieza-Color"] = row["Pieza-Color"]
        summary["Nombre tipo máquina"] = row["Nombre tipo máquina"]
        summary["Proceso máquina"] = row["Proceso máquina"]
        summary["Fracción virgen (receta)"] = row["Fracción virgen (receta)"]
        for index, material in enumerate(material_text, start=1):
            summary.setdefault(f"Material {index}", material)
        summary.setdefault("Maquinistas", [])
        summary.setdefault("_final_dates", set())
        summary.setdefault("_final_date_states", set())
        summary.setdefault("_worker_coverages", set())
        summary.setdefault("Mangas", 0)
        summary.setdefault("Mangas abiertas", 0)
        summary.setdefault("Mangas con pesaje restringido", 0)
        summary.setdefault("Mangas anuladas", 0)
        summary.setdefault("Bultos pesados", 0)
        summary.setdefault("Pendientes de pesaje", 0)
        summary.setdefault("KG final vigente", 0.0)
        summary.setdefault("_has_final_kg", False)
        summary.setdefault("_gross_values", [])
        summary.setdefault("_tare_values", [])
        summary.setdefault("_quantity_values", [])
        summary.setdefault("_fraction_values", [])
        summary.setdefault("Ciclos aproximados", 0.0)
        summary.setdefault("_has_cycles", False)
        summary.setdefault("_cycle_states", set())
        summary.setdefault("_group_key", key)
        summary["Mangas"] += 1
        if worker_text != "Sin maquinista asignado":
            summary["Maquinistas"] = list(dict.fromkeys(summary["Maquinistas"] + worker_text.split(", ")))
        summary["_worker_coverages"].add(worker_coverage)
        if isinstance(row["Fecha del pesaje final"], date):
            summary["_final_dates"].add(row["Fecha del pesaje final"])
        summary["_final_date_states"].add(row["Estado fecha final"])
        if not weights_visible:
            summary["Mangas con pesaje restringido"] += 1
        elif final_kg is None and manga.get("estado_operativo") in OPEN_MANGA_STATES:
            summary["Mangas abiertas"] += 1
            summary["Pendientes de pesaje"] += 1
        elif final_kg is None and manga.get("estado_operativo") == "ANULADA":
            summary["Mangas anuladas"] += 1
        elif final_kg is None:
            summary["Pendientes de pesaje"] += 1
        elif final_kg is not None:
            summary["KG final vigente"] += float(final_kg)
            summary["_has_final_kg"] = True
            summary["Bultos pesados"] += 1
            if gross is not None:
                summary["_gross_values"].append(float(gross))
            if tare is not None:
                summary["_tare_values"].append(float(tare))
            if confirmed_quantity is not None:
                summary["_quantity_values"].append(float(confirmed_quantity))
        if fraction is not None:
            summary["_fraction_values"].append(fraction)
        summary["_cycle_states"].add(cycles_state)
        if cycles is not None:
            summary["Ciclos aproximados"] += cycles
            summary["_has_cycles"] = True
    summary_rows = list(groups.values())
    for row in summary_rows:
        row["Maquinistas"] = ", ".join(row["Maquinistas"]) or "Sin maquinista asignado"
        final_dates = row.pop("_final_dates", set())
        final_date_states = row.pop("_final_date_states", set())
        if len(final_dates) == 1 and final_date_states <= {"Fecha final vigente"}:
            row["Fecha del pesaje final"] = next(iter(final_dates))
        else:
            row["Fecha del pesaje final"] = None
        if len(final_dates) > 1:
            row["Estado fecha final"] = "Varias fechas finales: " + ", ".join(sorted(item.isoformat() for item in final_dates))
            if final_date_states - {"Fecha final vigente"}:
                row["Estado fecha final"] += "; " + "; ".join(sorted(final_date_states - {"Fecha final vigente"}))
        elif final_date_states:
            row["Estado fecha final"] = "; ".join(sorted(final_date_states))
        else:
            row["Estado fecha final"] = "SIN_PESAJE_FINAL_VIGENTE"
        worker_coverages = row.pop("_worker_coverages", set())
        row["Cobertura maquinistas"] = (
            "INCOMPLETA" if "INCOMPLETA" in worker_coverages
            else ("SIN_ASIGNACION" if worker_coverages and worker_coverages <= {"SIN_ASIGNACION"} else "COMPLETA")
        )
        if row["Mangas con pesaje restringido"]:
            row["Mangas abiertas"] = None
        if not row.pop("_has_final_kg", False):
            row["KG final vigente"] = None
        gross_values = row.pop("_gross_values", [])
        tare_values = row.pop("_tare_values", [])
        quantity_values = row.pop("_quantity_values", [])
        fraction_values = row.pop("_fraction_values", [])
        row["Peso real (kg)"] = row["KG final vigente"]
        row["Peso bruto (kg)"] = sum(gross_values) if len(gross_values) == row["Bultos pesados"] and gross_values else None
        row["Tara (kg)"] = sum(tare_values) if len(tare_values) == row["Bultos pesados"] and tare_values else None
        row["Cantidad por bulto"] = (
            quantity_values[0]
            if quantity_values and len(set(quantity_values)) == 1 and len(quantity_values) == row["Bultos pesados"]
            else None
        )
        row["Fracción virgen (receta)"] = (
            fraction_values[0] if fraction_values and len(set(fraction_values)) == 1 else None
        )
        status_parts = []
        if row["Bultos pesados"]:
            status_parts.append(f"{row['Bultos pesados']} bultos pesados")
        if row["Pendientes de pesaje"]:
            status_parts.append(f"{row['Pendientes de pesaje']} pendientes de pesaje")
        if row["Mangas con pesaje restringido"]:
            status_parts.append(f"{row['Mangas con pesaje restringido']} con peso no autorizado")
        if row["Mangas anuladas"]:
            status_parts.append(f"{row['Mangas anuladas']} anuladas")
        row["Situación del pesaje"] = "; ".join(status_parts) or "Sin pesajes"
        observations = []
        if row["Cantidad por bulto"] is None and row["Bultos pesados"]:
            observations.append("Cantidad por bulto ausente o variable")
        if row["Fracción virgen (receta)"] is None:
            observations.append("Fracción virgen no acreditada")
        row["Observación"] = "; ".join(observations)
        row["Pieza-Color"] = row.get("Pieza-Color") or f"{row.get('Pieza') or 'Pieza sin nombre'} · {row.get('Color') or 'Color no indicado'}"
        cycle_states = row.pop("_cycle_states", set())
        if not row.pop("_has_cycles", False):
            row["Ciclos aproximados"] = None
        row["Estado ciclos aproximados"] = "; ".join(sorted(cycle_states)) or "SIN_KG_FINAL"
    group_refs = {
        key: f"G-{index:04d}"
        for index, key in enumerate(groups, start=1)
    }
    for row in summary_rows:
        row["Referencia grupo"] = group_refs[row["_group_key"]]
        row.pop("_group_key", None)
    for row in detail_rows:
        row["Referencia grupo"] = group_refs[row["_group_key"]]
        row.pop("_group_key", None)
        row["Peso real (kg)"] = row["KG final vigente"]
    return summary_rows, detail_rows


def key_fields(max_materials=2):
    fields = [
        "Tipo máquina", "Máquina", "Fecha de OT", "Maquinistas",
        "Tipo materia prima",
    ]
    fields.extend(f"Material {index}" for index in range(1, max(2, max_materials) + 1))
    fields.extend([
        "Molde", "Pieza-Color", "Ciclos aproximados", "Peso real (kg)", "Turno",
        "Cantidad por bulto", "Bultos pesados", "N°OF", "Fracción virgen (receta)",
        "Fecha del pesaje final", "Estado fecha final", "Peso bruto (kg)", "Tara (kg)", "Situación del pesaje",
        "Mangas", "Pendientes de pesaje", "Mangas con pesaje restringido", "Mangas anuladas",
        "Observación", "Referencia grupo",
    ])
    return fields


def detail_fields(max_materials=2):
    fields = [
        "Tipo máquina", "Máquina", "Fecha de OT", "Maquinistas",
        "Tipo materia prima",
    ]
    fields.extend(f"Material {index}" for index in range(1, max(2, max_materials) + 1))
    fields.extend([
        "Molde", "Pieza-Color", "Ciclos aproximados", "Peso real (kg)", "Turno",
        "Cantidad por bulto", "Bultos pesados", "N°OF", "Fracción virgen (receta)",
        "Fecha del pesaje final", "Estado fecha final", "Peso bruto (kg)", "Tara (kg)", "Situación del pesaje",
        "Manga", "N°OT", "Estado operativo", "Observación", "Referencia grupo",
    ])
    return fields


def audit_fields():
    return [
        "Referencia grupo", "Manga", "ID manga", "N°OF", "N°OT", "Fecha de OT", "Estado operativo",
        "Estado pesaje", "Fecha del pesaje final", "Estado fecha final", "Hora del pesaje", "Peso bruto (kg)", "Tara (kg)",
        "Peso neto (kg)", "Cantidad UN acreditada", "Fuente cantidad", "Tipo máquina", "Proceso máquina",
        "Nombre tipo máquina", "Código máquina", "Tipo materia prima", "Fracción virgen (receta)",
        "Estado fracción virgen", "Receta ID", "Receta revisión", "Receta hash", "Molde", "Código molde",
        "Salida/base ciclos", "Estado receta", "Estado molde", "Estado salida", "Estado datos producción",
        "Cobertura maquinistas", "Observación técnica",
    ]


def _audit_rows(detail_rows):
    """Build one traceability row per physical manga.

    Operational sheets intentionally omit immutable identifiers and technical
    diagnostics.  Keeping them here makes a workbook useful for review
    without making the production-facing table unreadable.
    """
    rows = []
    for source in detail_rows:
        rows.append({
            "Referencia grupo": source.get("Referencia grupo"),
            "Manga": source.get("Manga"),
            "ID manga": source.get("ID manga"),
            "N°OF": source.get("N°OF"),
            "N°OT": source.get("N°OT"),
            "Fecha de OT": source.get("Fecha de OT"),
            "Estado operativo": source.get("_estado_operativo_code") or source.get("Estado operativo"),
            "Estado pesaje": source.get("Estado pesaje"),
            "Fecha del pesaje final": source.get("Fecha del pesaje final"),
            "Estado fecha final": source.get("Estado fecha final"),
            "Hora del pesaje": source.get("Hora del pesaje"),
            "Peso bruto (kg)": source.get("Peso bruto (kg)"),
            "Tara (kg)": source.get("Tara (kg)"),
            "Peso neto (kg)": source.get("KG final vigente"),
            "Cantidad UN acreditada": source.get("Cantidad por bulto"),
            "Fuente cantidad": source.get("Fuente cantidad"),
            "Tipo máquina": source.get("Tipo máquina"),
            "Proceso máquina": source.get("Proceso máquina"),
            "Nombre tipo máquina": source.get("Nombre tipo máquina"),
            "Código máquina": source.get("Código máquina"),
            "Tipo materia prima": source.get("Tipo materia prima"),
            "Fracción virgen (receta)": source.get("Fracción virgen (receta)"),
            "Estado fracción virgen": source.get("Estado fracción virgen"),
            "Receta ID": source.get("Receta ID"),
            "Receta revisión": source.get("Receta revisión"),
            "Receta hash": source.get("Receta hash"),
            "Molde": source.get("Molde"),
            "Código molde": source.get("Código molde"),
            "Salida/base ciclos": source.get("Base técnica ciclos"),
            "Estado receta": source.get("Estado receta"),
            "Estado molde": source.get("Estado molde"),
            "Estado salida": source.get("Estado salida"),
            "Estado datos producción": source.get("Estado datos producción"),
            "Cobertura maquinistas": source.get("Cobertura maquinistas"),
            "Observación técnica": source.get("Observación"),
        })
    return rows


def _style_export_columns(sheet, headers):
    """Apply a small, typed, readable layout to every exported column."""
    widths = {
        "Referencia grupo": 16, "Tipo máquina": 18, "Máquina": 20, "Fecha de OT": 14,
        "Maquinistas": 28, "Tipo materia prima": 18,
        "Receta revisión": 15, "Receta hash": 18, "Identidad receta": 34,
        "Material 1": 22, "Material 2": 22, "Material 3": 22,
        "Molde": 22, "Pieza-Color": 30, "Código molde": 16, "Fuente nombre molde": 28,
        "Ciclos aproximados": 18, "Peso real (kg)": 16, "Cantidad por bulto": 18,
        "Bultos pesados": 16, "Fracción virgen (receta)": 22, "Fecha del pesaje final": 22,
        "Estado fecha final": 34,
        "Turno": 12, "N°OF": 16, "N°OT": 16, "Manga": 18, "ID manga": 38,
        "Peso bruto (kg)": 16, "Tara (kg)": 14, "Peso neto (kg)": 16,
        "Hora del pesaje": 22, "Situación del pesaje": 30,
        "Pendientes de pesaje": 20, "Mangas con pesaje restringido": 28, "Mangas anuladas": 18,
        "Observación": 34, "Observación técnica": 34, "Estado ciclos aproximados": 24,
        "Estado datos producción": 28, "Base técnica ciclos": 24,
        "Salida/base ciclos": 24, "Estado operativo": 18, "Estado pesaje": 20,
        "Estado fracción virgen": 30, "Proceso máquina": 18, "Nombre tipo máquina": 24,
        "Código máquina": 18, "Cantidad UN acreditada": 18, "Fuente cantidad": 22,
        "Estado receta": 28, "Estado molde": 28, "Estado salida": 28,
        "Cobertura maquinistas": 24,
    }
    for index, header in enumerate(headers, start=1):
        width = widths.get(header, 22 if header.startswith("Material ") else 16)
        column = sheet.cell(row=1, column=index).column_letter
        sheet.column_dimensions[column].width = width
        for cell in sheet[column]:
            if cell.row == 1:
                cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                cell.fill = PatternFill(fill_type="solid", fgColor="1F4E78")
            else:
                cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if header == "Ciclos aproximados" and cell.row > 1:
                cell.number_format = "0.0"
            elif header in {"Peso real (kg)", "Peso bruto (kg)", "Tara (kg)", "Peso neto (kg)"} and cell.row > 1:
                cell.number_format = "0.00"
            elif header in {"Fracción virgen (receta)"} and cell.row > 1:
                cell.number_format = "0.0%"
            elif header in {"Fecha de OT", "Fecha del pesaje final"} and cell.row > 1:
                cell.number_format = "dd/mm/yyyy"
            elif header == "Hora del pesaje" and cell.row > 1:
                cell.number_format = "dd/mm/yyyy hh:mm"
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = f"A1:{sheet.cell(row=max(1, sheet.max_row), column=max(1, len(headers))).coordinate}"


def generate_manga_explorer_xlsx(session, *, actor_id, filters=None):
    raw = dict(filters or {})
    mode = str(raw.pop("vista", raw.pop("modo", "AMBOS"))).strip().upper()
    if mode not in {"RESUMIDO", "DETALLADO", "AMBOS"}:
        raise ScmServiceError("INVALID_MANGA_EXPORT_MODE", "vista debe ser RESUMIDO, DETALLADO o AMBOS.", status_code=400)
    # Exports always represent the complete filtered physical set, never the
    # cursor page currently visible in the table.
    raw.pop("cursor", None)
    raw.pop("page", None)
    parsed = _filters(raw)
    date_type = parsed["fecha_tipo"]
    payload = _collect(session, actor_id=actor_id, raw_filters=raw, paginate=False, include_production=True)
    summary_rows, detail_rows = _export_rows(payload)
    audit_rows = _audit_rows(detail_rows)
    max_materials = max(
        (len((item.get("produccion") or {}).get("materias_primas") or []) for item in payload["items"]),
        default=2,
    )
    workbook = Workbook()
    workbook.remove(workbook.active)
    criterion = "Fecha del pesaje final" if date_type == "PESAJE_FINAL" else "Fecha de OT"

    def append_data_sheet(title, headers, rows):
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for row in rows:
            sheet.append([_safe_cell(row.get(header)) for header in headers])
        _style_export_columns(sheet, headers)
        sheet["A1"].comment = Comment(
            f"Criterio de fecha: {criterion}. Cohorte por OT origen; no reparto diario de masa. "
            "Ciclos: base snapshot por salida; no aditivo entre coproductos.",
            "EnvaPeru",
        )
        return sheet

    if mode in {"RESUMIDO", "AMBOS"}:
        append_data_sheet("Resumen", key_fields(max_materials), summary_rows)
    if mode in {"DETALLADO", "AMBOS"}:
        append_data_sheet("Detalle", detail_fields(max_materials), detail_rows)
    append_data_sheet("Auditoría", audit_fields(), audit_rows)

    # Keep the applied scope in a visible, typed worksheet rather than hiding
    # it in a header comment that is easy to miss in an exported file.
    criteria_sheet = workbook.create_sheet("Criterios")
    criteria_rows = [
        ["Criterio de fecha", criterion],
        ["Desde", parsed["fecha_desde"].isoformat() if parsed["fecha_desde"] else "Todas"],
        ["Hasta", parsed["fecha_hasta"].isoformat() if parsed["fecha_hasta"] else "Todas"],
        ["Búsqueda", parsed["q"] or "Todas"],
        ["Estado", _human_label(parsed["estado_manga"], "Todos")],
        ["Campo aplicado", "Fecha del pesaje final (fecha_local_pesaje del final vigente)" if date_type == "PESAJE_FINAL" else "Fecha de OT (origen)"],
        ["Conjunto exportado", "Completo según filtros; el cursor de la tabla no limita el Excel"],
        ["Resumen", "Una fila por cohorte operativa; KG y ciclos no se duplican por maquinista ni coproducto"],
        ["Detalle", "Una fila por manga física; una pesada vigente aporta su KG una sola vez"],
        ["Cantidad por bulto", "Sólo se muestra como Cantidad UN acreditada cuando la fuente observada declara unidades (UN); KG, plan o fuente ambigua quedan vacíos y explicados."],
        ["Ciclos aproximados", "KG final vigente × 1000 / (cantidad por ciclo snapshot × peso unitario snapshot g); no suma coproductos"],
        ["Receta vs consumo", "Materiales y fracción virgen son snapshot de receta; no representan consumo real de materiales."],
        ["Tipo máquina", "Proviene del proceso actual del catálogo TipoMaquina; el nombre/modelo se conserva separado en Auditoría."],
        ["Agrupación de ciclos", "Resumen separa receta, molde y salida/base técnica; los coproductos no se suman entre sí."],
        ["Campos no disponibles", "Se dejan vacíos con observación; no se estiman unidades desde KG ni se inventan horas o mermas."],
        ["Estados de pesaje", "Pesaje final vigente; Pendiente de pesaje final; Peso no autorizado; Manga anulada sin pesaje final."],
        ["Dimensiones", "La cohorte se agrupa por Fecha OT, turno, máquina, N°OF, Pieza-Color, receta, molde y salida/base; Detalle conserva una fila por manga física."],
        ["Ausente vs cero", "Celda vacía significa fuente ausente/no acreditada; cero sólo se conserva cuando la fuente lo registra."],
        ["Columnas omitidas", "Por falta de fuente no se inventan HoraIngreso, HorasTrab, documento, merma ni chancaca; UUID, hash y estados técnicos quedan en Auditoría."],
    ]
    criteria_sheet.append(["Campo", "Descripción"])
    for values in criteria_rows:
        criteria_sheet.append([_safe_cell(value) for value in values])
    _style_export_columns(criteria_sheet, ["Campo", "Descripción"])
    criteria_sheet.column_dimensions["A"].width = 24
    criteria_sheet.column_dimensions["B"].width = 100
    for cell in criteria_sheet["A"][1:]:
        cell.font = Font(name="Arial", size=10, bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in criteria_sheet["B"][1:]:
        cell.font = Font(name="Arial", size=10)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    criteria_sheet.freeze_panes = "A2"
    criteria_sheet.sheet_view.showGridLines = False
    criteria_sheet.auto_filter.ref = f"A1:B{max(1, criteria_sheet.max_row)}"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
