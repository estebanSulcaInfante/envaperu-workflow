from io import BytesIO
from datetime import date, datetime, timezone
from decimal import Decimal
import os
from pathlib import Path
from types import SimpleNamespace
from uuid import UUID, uuid4

from openpyxl import load_workbook
import pytest
from sqlalchemy import event

from app.extensions import db
from app.models.registro import RegistroDiarioProduccion
from app.models.scm_production_orders import ScmOrdenOperacion
from app.models.scm_production_orders import (
    ScmCorridaFabricacion,
    ScmOrdenFabricacion,
    ScmOrdenOperacionSalida,
)
from app.models.scm_ot import (
    ScmAsignacionPersonalTrabajoOt,
    ScmAsignacionPlanMangaOt,
    ScmCorreccionAsignacionManga,
    ScmCorreccionPesajeManga,
    ScmLoteArticulo,
    ScmManga,
    ScmPesajeManga,
    ScmTrabajoColor,
    ScmTrabajoOt,
    ScmTramoMangaTrabajo,
)
from app.models.scm_catalogos import ScmCategoriaRecepcion, ScmMaterial
from app.models.scm_articulos import ScmArticulo
from app.models.materiales import MateriaPrima
from app.models.molde import Molde
from app.models.maquina import Maquina, TipoMaquina
from app.models.receta_color import RecetaColorLinea, RecetaColorMaestra
from app.models.trabajador import Trabajador
from app.services.scm_assignment_correction_service import apply_assignment_correction
from app.services.scm_ot_service import create_fabrication_ot_header
from app.services.scm_manga_explorer_service import (
    _audit_rows,
    _canonical_color_hex,
    _export_rows,
    _material_recipe_context,
    _production_context,
    _safe_cell,
    _worker_projection,
    audit_fields,
    detail_fields,
    key_fields,
)
from tests.scm.test_scm_production_observability import _manga, _seed_observability_graph
from tests.scm.test_scm_kg_production import _auto_final_kg_fixture


def _headers(actor):
    return {"X-Actor-Id": str(actor.id)}


def _add_scaled_manga_graph(seeded, count=120):
    """Build a page/export-sized physical set without changing production data."""
    from app.models.scm_ot import ScmTrabajoOt
    from app.models.trabajador import Trabajador

    worker = Trabajador.query.filter_by(codigo="TRB-01").one()
    works = [
        ScmTrabajoOt.query.filter_by(codigo="TC-OBS-ROJO").one(),
        ScmTrabajoOt.query.filter_by(codigo="TC-OBS-AZUL").one(),
    ]
    secondary_operation = ScmOrdenOperacion(
        codigo="OF-OBS-002",
        tipo="FABRICACION",
        origen_demanda="PLANIFICADA",
        plan_produccion_id=works[0].orden_operacion.plan_produccion_id,
        propuesta_clave="OF-2",
        estado="EN_EJECUCION",
        created_by_id=seeded["full"].id,
    )
    secondary_ot = RegistroDiarioProduccion(
        codigo_ot="OT-OBS-FAB-B",
        codigo_ot_sintetico=False,
        estado="EN_EJECUCION",
        tipo_ot="FABRICACION",
        fecha=seeded["fab"].fecha,
        turno="NOCHE",
        maquina_id=1,
        maquina_codigo_snapshot="MAQ-B",
        maquina_nombre_snapshot="Máquina B",
        responsable_id=worker.id,
        maquinista_previsto_id=worker.id,
        created_by_id=seeded["full"].id,
        created_at=seeded["fab"].created_at,
        updated_at=seeded["fab"].updated_at,
    )
    db.session.add_all([secondary_operation, secondary_ot])
    db.session.flush()
    secondary_work = ScmTrabajoOt(
        orden_trabajo_id=secondary_ot.id,
        codigo="TC-OBS-VERDE-B",
        secuencia=1,
        estado="PLANIFICADO",
        orden_operacion_id=secondary_operation.id,
        cantidad_objetivo_un=250,
        cantidad_confirmada_un=150,
        created_by_id=seeded["full"].id,
    )
    db.session.add(secondary_work)
    db.session.flush()
    db.session.add(ScmTrabajoColor(
        trabajo_ot_id=secondary_work.id,
        corrida_fabricacion_id=uuid4(),
        color_nombre_snapshot="VERDE",
    ))
    db.session.flush()
    works.append(secondary_work)
    ots = [seeded["fab"], seeded["fab"], secondary_ot]
    assignments = [
        ScmAsignacionPersonalTrabajoOt(
            trabajo_ot_id=work.id,
            trabajador_id=worker.id,
            asignada_por_id=seeded["full"].id,
            estado="PREVISTA",
        )
        for work in works
    ]
    db.session.add_all(assignments)
    db.session.flush()
    mangas = []
    for sequence in range(100, 100 + count):
        work = works[sequence % len(works)]
        manga = _manga(
            ot=ots[sequence % len(ots)], work=work, sequence=sequence, state="PREETIQUETADA", worker=worker,
        )
        manga.asignacion_personal_trabajo = assignments[sequence % len(assignments)]
        mangas.append(manga)
    db.session.add_all(mangas)
    db.session.flush()
    for index, manga in enumerate(mangas[::10], start=1):
        assignment = assignments[index % len(assignments)]
        db.session.add(ScmTramoMangaTrabajo(
            manga=manga,
            trabajo=assignment.trabajo,
            asignacion_personal_trabajo=assignment,
            secuencia=1,
            estado="ACTIVO",
            cantidad_inicio_un=0,
            cantidad_atribuida_un=0,
            created_by_id=worker.id,
        ))
    db.session.commit()


def _add_auditable_correction(app, station_code):
    """Create a real correction overlay through the canonical service path."""
    from tests.scm.test_assignment_correction_service import _grant_assignment_correction

    creator, _approver, manga, _station, _label, _weighed = _auto_final_kg_fixture(
        app, station_code=station_code,
    )
    _grant_assignment_correction(creator)
    source = manga.trabajo
    source_color = source.trabajo_color
    second_manga = (
        ScmManga.query.filter_by(trabajo_ot_id=source.id)
        .filter(ScmManga.id != manga.id)
        .order_by(ScmManga.id)
        .first()
    )
    assert second_manga is not None
    for corrected_manga in (manga, second_manga):
        corrected_manga.estado = "PESADA"
        db.session.add(ScmTramoMangaTrabajo(
            manga=corrected_manga,
            trabajo=source,
            asignacion_personal_trabajo=corrected_manga.asignacion_personal_trabajo,
            asignacion_plan_id=corrected_manga.asignacion_id,
            secuencia=1,
            estado="CERRADO",
            cantidad_inicio_un=0,
            cantidad_atribuida_un=0,
            cantidad_fin_un=corrected_manga.cantidad_asignada_un,
            motivo_cierre="PESAJE_FINAL",
            created_by_id=creator.id,
        ))
    db.session.flush()
    def _target_work(target_ot, code):
        target = ScmTrabajoOt(
            orden_trabajo_id=target_ot.id,
            codigo=f"{target_ot.codigo_ot}-{code}",
            tipo="COLOR",
            secuencia=1,
            estado="PAUSADO",
            orden_operacion_id=source.orden_operacion_id,
            cantidad_objetivo_un=0,
            cantidad_confirmada_un=0,
            created_by_id=creator.id,
        )
        target.trabajo_color = ScmTrabajoColor(
            corrida_fabricacion_id=source_color.corrida_fabricacion_id,
            molde_codigo_snapshot=source_color.molde_codigo_snapshot,
            color_id_snapshot=source_color.color_id_snapshot,
            color_nombre_snapshot=source_color.color_nombre_snapshot,
            receta_revision_id_snapshot=source_color.receta_revision_id_snapshot,
            receta_hash_snapshot=source_color.receta_hash_snapshot,
            cavidades_snapshot=source_color.cavidades_snapshot,
            peso_neto_snapshot_g=source_color.peso_neto_snapshot_g,
            peso_colada_snapshot_g=source_color.peso_colada_snapshot_g,
        )
        assignment = ScmAsignacionPersonalTrabajoOt(
            trabajo=target,
            trabajador_id=creator.id,
            estado="PREVISTA",
            asignada_por_id=creator.id,
        )
        db.session.add_all([target, assignment])
        db.session.flush()
        target_plan = ScmAsignacionPlanMangaOt(
            plan_linea_id=manga.plan_linea_id,
            ot_id=target_ot.id,
            trabajo_ot_id=target.id,
            cantidad_asignada_un=0,
            mangas_asignadas=0,
            asignada_por_id=creator.id,
        )
        db.session.add(target_plan)
        db.session.flush()
        return target, assignment, target_plan

    active_header = create_fabrication_ot_header(
        db.session,
        actor_id=creator.id,
        operation_id=uuid4(),
        data={
            "maquina_id": source.orden_trabajo.maquina_id,
            "fecha_operativa": "2026-09-19",
            "turno": "DIA",
            "maquinista_predeterminado_id": creator.id,
        },
    )["ot"]
    active_ot = RegistroDiarioProduccion.query.filter_by(
        public_id=UUID(active_header["public_id"]),
    ).one()
    active_work, active_assignment, active_plan = _target_work(active_ot, "C01")

    superseded_header = create_fabrication_ot_header(
        db.session,
        actor_id=creator.id,
        operation_id=uuid4(),
        data={
            "maquina_id": source.orden_trabajo.maquina_id,
            "fecha_operativa": "2026-09-20",
            "turno": "DIA",
            "maquinista_predeterminado_id": creator.id,
        },
    )["ot"]
    superseded_ot = RegistroDiarioProduccion.query.filter_by(
        public_id=UUID(superseded_header["public_id"]),
    ).one()
    superseded_correction_work, superseded_assignment, superseded_plan = _target_work(
        superseded_ot, "C01",
    )
    latest_header = create_fabrication_ot_header(
        db.session,
        actor_id=creator.id,
        operation_id=uuid4(),
        data={
            "maquina_id": source.orden_trabajo.maquina_id,
            "fecha_operativa": "2026-09-21",
            "turno": "DIA",
            "maquinista_predeterminado_id": creator.id,
        },
    )["ot"]
    superseded_latest_ot = RegistroDiarioProduccion.query.filter_by(
        public_id=UUID(latest_header["public_id"]),
    ).one()
    superseded_latest_work, superseded_latest_assignment, superseded_latest_plan = _target_work(
        superseded_latest_ot, "C01",
    )
    db.session.commit()
    apply_assignment_correction(
        db.session,
        actor_id=creator.id,
        manga_id=manga.public_id,
        operation_id=uuid4(),
        data={
            "destino_trabajo_ot_id": str(active_work.id),
            "destino_asignacion_id": str(active_assignment.id),
            "tramo_objetivo_id": str(manga.tramos_trabajo[0].id),
            "version": manga.version,
            "motivo": "La manga fue vinculada a la OT efectiva del turno.",
        },
    )
    apply_assignment_correction(
        db.session,
        actor_id=creator.id,
        manga_id=second_manga.public_id,
        operation_id=uuid4(),
        data={
            "destino_trabajo_ot_id": str(superseded_correction_work.id),
            "destino_asignacion_id": str(superseded_assignment.id),
            "tramo_objetivo_id": str(second_manga.tramos_trabajo[0].id),
            "version": second_manga.version,
            "motivo": "La manga fue vinculada a la OT efectiva del turno.",
        },
    )
    db.session.add(ScmTramoMangaTrabajo(
        manga=second_manga,
        trabajo=superseded_latest_work,
        asignacion_personal_trabajo=superseded_latest_assignment,
        asignacion_plan_id=superseded_latest_plan.id,
        secuencia=2,
        estado="ACTIVO",
        cantidad_inicio_un=0,
        cantidad_atribuida_un=0,
        created_by_id=creator.id,
    ))
    db.session.commit()
    return {
        "manga": manga,
        "second_manga": second_manga,
        "origin_ot": source.orden_trabajo,
        "effective_ot": active_ot,
        "effective_work": active_work,
        "active_ot": active_ot,
        "active_work": active_work,
        "superseded_correction_ot": superseded_ot,
        "superseded_correction_work": superseded_correction_work,
        "superseded_latest_ot": superseded_latest_ot,
        "superseded_latest_work": superseded_latest_work,
    }


def test_manga_explorer_is_fabrication_only_and_keeps_physical_rows(app, client, scm_config):
    with app.app_context():
        seeded = _seed_observability_graph()
        response = client.get(
            "/api/scm/v1/observabilidad/mangas-fabricacion",
            headers=_headers(seeded["full"]),
        )

        assert response.status_code == 200
        payload = response.get_json()
        assert {item["ot"]["tipo"] for item in payload["items"]} == {"FABRICACION"}
        assert len({item["manga"]["public_id"] for item in payload["items"]}) == len(payload["items"])

        supervision = client.get(
            "/api/scm/v1/observabilidad/mangas",
            headers=_headers(seeded["full"]),
        )
        assert supervision.status_code == 200
        assert "ARMADO" in {item["ot"]["tipo"] for item in supervision.get_json()["items"]}


def test_manga_explorer_final_date_filters_before_limit_and_cursor(app, client, scm_config):
    with app.app_context():
        seeded = _seed_observability_graph()
        corrected = ScmManga.query.filter_by(codigo="M-OT-OBS-FAB-02").one()
        received = ScmManga.query.filter_by(codigo="M-OT-OBS-FAB-03").one()
        corrected_weighing = ScmPesajeManga.query.filter_by(manga_id=corrected.id).one()
        received_weighing = ScmPesajeManga.query.filter_by(manga_id=received.id).one()
        corrected_weighing.fecha_local_pesaje = date(2026, 8, 11)
        received_weighing.fecha_local_pesaje = date(2026, 8, 12)
        db.session.commit()

        first = client.get(
            "/api/scm/v1/observabilidad/mangas-fabricacion?fecha_tipo=PESAJE_FINAL&fecha_desde=2026-08-10&fecha_hasta=2026-08-12&limit=1",
            headers=_headers(seeded["full"]),
        )
        assert first.status_code == 200
        first_payload = first.get_json()
        assert len(first_payload["items"]) == 1
        assert first_payload["items"][0]["manga"]["codigo"] == "M-OT-OBS-FAB-03"
        assert first_payload["page"]["has_more"] is True

        second = client.get(
            f"/api/scm/v1/observabilidad/mangas-fabricacion?fecha_tipo=PESAJE_FINAL&fecha_desde=2026-08-10&fecha_hasta=2026-08-12&limit=1&cursor={first_payload['page']['next_cursor']}",
            headers=_headers(seeded["full"]),
        )
        assert second.status_code == 200
        assert [item["manga"]["codigo"] for item in second.get_json()["items"]] == ["M-OT-OBS-FAB-02"]

        exported = client.get(
            "/api/scm/v1/observabilidad/mangas-fabricacion/export.xlsx?vista=DETALLADO&fecha_tipo=PESAJE_FINAL&fecha_desde=2026-08-10&fecha_hasta=2026-08-12",
            headers=_headers(seeded["full"]),
        )
        workbook = load_workbook(BytesIO(exported.data), read_only=False, data_only=False)
        headers = next(workbook["Detalle"].iter_rows(values_only=True))
        assert "Fecha del pesaje final" in headers
        assert "Criterio de fecha: Fecha del pesaje final" in workbook["Detalle"]["A1"].comment.text
        final_date = headers.index("Fecha del pesaje final")
        final_values = list(workbook["Detalle"].iter_rows(min_row=2, values_only=True))
        assert {row[final_date].date() for row in final_values} == {
            date(2026, 8, 10), date(2026, 8, 12),
        }
        assert all(isinstance(row[final_date], (date, datetime)) for row in final_values)
        assert workbook["Detalle"].cell(row=2, column=final_date + 1).number_format == "dd/mm/yyyy"


def test_manga_explorer_final_date_requires_weight_permission(app, client, scm_config):
    with app.app_context():
        seeded = _seed_observability_graph()
        response = client.get(
            "/api/scm/v1/observabilidad/mangas-fabricacion?fecha_tipo=PESAJE_FINAL",
            headers=_headers(seeded["base"]),
        )
        assert response.status_code == 403
        assert response.get_json()["error"]["code"] == "MANGA_PESAJE_PERMISSION_REQUIRED"


def test_manga_explorer_correction_quantity_source_is_explicit_or_preserved(app, client, scm_config):
    with app.app_context():
        seeded = _seed_observability_graph()
        correction = ScmCorreccionPesajeManga.query.filter_by(estado="APLICADA").one()
        correction.pesaje.fuente_cantidad = "RESPONSABLE_ARMADO"
        db.session.commit()
        response = client.get(
            "/api/scm/v1/observabilidad/mangas-fabricacion",
            headers=_headers(seeded["full"]),
        )
        assert response.status_code == 200
        item = next(
            item for item in response.get_json()["items"]
            if item["manga"]["codigo"] == "M-OT-OBS-FAB-02"
        )
        assert item["manga"]["pesaje"]["fuente_cantidad"] == "RESPONSABLE_ARMADO"
        correction.proposed_json = {"cantidad_confirmada": "96.000"}
        db.session.commit()
        response = client.get(
            "/api/scm/v1/observabilidad/mangas-fabricacion",
            headers=_headers(seeded["full"]),
        )
        item = next(
            item for item in response.get_json()["items"]
            if item["manga"]["codigo"] == "M-OT-OBS-FAB-02"
        )
        assert item["manga"]["pesaje"]["fuente_cantidad"] == "CORRECCION_AUTORIZADA"


def test_manga_explorer_rejects_final_cursor_after_real_manga_annulment_event(app, client, scm_config):
    with app.app_context():
        seeded = _seed_observability_graph()
        first = client.get(
            "/api/scm/v1/observabilidad/mangas-fabricacion?fecha_tipo=PESAJE_FINAL&limit=1",
            headers=_headers(seeded["full"]),
        )
        assert first.status_code == 200
        cursor = first.get_json()["page"]["next_cursor"]
        assert cursor
        manga = ScmManga.query.filter_by(codigo="M-OT-OBS-FAB-02").one()
        manga.anulada_at = datetime.now(timezone.utc)
        manga.estado = "ANULADA"
        db.session.commit()
        stale = client.get(
            f"/api/scm/v1/observabilidad/mangas-fabricacion?fecha_tipo=PESAJE_FINAL&limit=1&cursor={cursor}",
            headers=_headers(seeded["full"]),
        )
        assert stale.status_code == 409
        assert stale.get_json()["error"]["code"] == "STALE_MANGA_EXPLORER_CURSOR"


def test_manga_explorer_rejects_final_cursor_after_applied_correction_resolution(app, client, scm_config):
    with app.app_context():
        seeded = _seed_observability_graph()
        first = client.get(
            "/api/scm/v1/observabilidad/mangas-fabricacion?fecha_tipo=PESAJE_FINAL&limit=1",
            headers=_headers(seeded["full"]),
        )
        assert first.status_code == 200
        cursor = first.get_json()["page"]["next_cursor"]
        correction = ScmCorreccionPesajeManga.query.filter_by(estado="APLICADA").one()
        correction.resolved_at = datetime.now(timezone.utc)
        db.session.commit()
        stale = client.get(
            f"/api/scm/v1/observabilidad/mangas-fabricacion?fecha_tipo=PESAJE_FINAL&limit=1&cursor={cursor}",
            headers=_headers(seeded["full"]),
        )
        assert stale.status_code == 409
        assert stale.get_json()["error"]["code"] == "STALE_MANGA_EXPLORER_CURSOR"


def test_manga_worker_projection_keeps_known_workers_and_incomplete_coverage():
    known = SimpleNamespace(
        id=11, codigo="TRB-11", nombre_corto="Ana", nombre_completo="Ana Planta",
    )
    assignment = SimpleNamespace(trabajador=known)
    first = SimpleNamespace(secuencia=1, created_at=None, asignacion_personal_trabajo=assignment)
    second = SimpleNamespace(secuencia=2, created_at=None, asignacion_personal_trabajo=None)
    manga = SimpleNamespace(
        correccion_asignacion=None,
        tramos_trabajo=[first, second],
        asignacion_personal_trabajo=None,
    )
    projection = _worker_projection(manga)
    assert projection["items"] == [{"id": 11, "codigo": "TRB-11", "nombre": "Ana"}]
    assert projection["cobertura"] == "INCOMPLETA"


def test_manga_explorer_supports_uuid_search_and_export_modes(app, client, scm_config):
    with app.app_context():
        seeded = _seed_observability_graph()
        listed = client.get(
            "/api/scm/v1/observabilidad/mangas-fabricacion?q=00000000-0000-0000-0000-000000000000",
            headers=_headers(seeded["full"]),
        )
        assert listed.status_code == 200
        by_origin_ot = client.get(
            "/api/scm/v1/observabilidad/mangas-fabricacion?q=OT-OBS-FAB",
            headers=_headers(seeded["full"]),
        )
        assert by_origin_ot.status_code == 200
        assert len(by_origin_ot.get_json()["items"]) == 4
        by_of = client.get(
            "/api/scm/v1/observabilidad/mangas-fabricacion?q=OF-OBS-001",
            headers=_headers(seeded["full"]),
        )
        assert by_of.status_code == 200
        assert len(by_of.get_json()["items"]) == 4

        exported = client.get(
            "/api/scm/v1/observabilidad/mangas-fabricacion/export.xlsx?vista=AMBOS&cursor=cursor-de-la-pagina-actual",
            headers=_headers(seeded["full"]),
        )
        assert exported.status_code == 200
        workbook = load_workbook(BytesIO(exported.data), read_only=False, data_only=False)
        assert workbook.sheetnames == ["Resumen", "Detalle", "Auditoría", "Criterios"]
        criteria = list(workbook["Criterios"].iter_rows(values_only=True))
        assert criteria[0] == ("Campo", "Descripción")
        assert criteria[1] == ("Criterio de fecha", "Fecha de OT")
        assert criteria[6][0] == "Campo aplicado"
        assert criteria[7][0] == "Conjunto exportado"
        assert next(workbook["Resumen"].iter_rows(values_only=True))[0] == "Tipo máquina"
        assert next(workbook["Detalle"].iter_rows(values_only=True))[0] == "Tipo máquina"
        assert "Cohorte por OT origen" in workbook["Resumen"]["A1"].comment.text
        assert "no reparto diario de masa" in workbook["Detalle"]["A1"].comment.text
        assert "Maquinistas" in next(workbook["Resumen"].iter_rows(values_only=True))
        assert "Maquinistas" in next(workbook["Detalle"].iter_rows(values_only=True))

        empty_export = client.get(
            "/api/scm/v1/observabilidad/mangas-fabricacion/export.xlsx?vista=AMBOS&q=NO-EXISTE",
            headers=_headers(seeded["full"]),
        )
        empty_workbook = load_workbook(BytesIO(empty_export.data), read_only=False, data_only=False)
        assert empty_workbook.sheetnames == ["Resumen", "Detalle", "Auditoría", "Criterios"]
        assert empty_workbook["Criterios"]["B2"].value == "Fecha de OT"
        assert empty_workbook["Resumen"].max_row == 1
        assert empty_workbook["Detalle"].max_row == 1
        assert empty_workbook["Auditoría"].max_row == 1
        assert list(next(empty_workbook["Resumen"].iter_rows(values_only=True))) == key_fields()
        assert list(next(empty_workbook["Detalle"].iter_rows(values_only=True))) == detail_fields()
        assert list(next(empty_workbook["Auditoría"].iter_rows(values_only=True))) == audit_fields()


def test_manga_explorer_export_rows_deduplicates_workers_and_sums_each_manga_once():
    item_template = {
        "ot": {"fecha_operativa": "2026-08-10", "turno": "DÍA", "codigo": "OT-1"},
        "recurso": {"nombre": "Máquina 1", "codigo": "MAQ-1"},
        "upstream": {"orden": {"codigo": "OF-1"}},
        "maquinistas": [
            {"id": 1, "codigo": "TRB-1", "nombre": "Ana"},
            {"id": 2, "codigo": "TRB-2", "nombre": "Bruno"},
            {"id": 1, "codigo": "TRB-1", "nombre": "Ana"},
        ],
        "maquinistas_cobertura": "COMPLETA",
        "visibilidad": {"pesaje": True},
    }
    first = {
        **item_template,
        "manga": {
            "codigo": "M-1", "public_id": "m-1", "estado_operativo": "CERRADA",
            "color": "Rojo", "articulo": {"nombre": "Pieza", "codigo": "P-1"},
            "pesaje": {"estado": "VIGENTE", "fecha_local_pesaje": "2026-08-10", "peso_fisico_neto_kg": 3.5},
        },
    }
    second = {
        **item_template,
        "manga": {
            "codigo": "M-2", "public_id": "m-2", "estado_operativo": "CERRADA",
            "color": "Rojo", "articulo": {"nombre": "Pieza", "codigo": "P-1"},
            "pesaje": {"estado": "VIGENTE", "fecha_local_pesaje": "2026-08-11", "peso_fisico_neto_kg": 4.5},
        },
    }
    summary, detail = _export_rows({"items": [first, second]})
    assert len(detail) == 2
    assert detail[0]["Maquinistas"] == "Ana, Bruno"
    assert isinstance(detail[0]["Fecha del pesaje final"], date)
    assert detail[0]["Estado fecha final"] == "Fecha final vigente"
    assert list(detail[0]).count("KG final vigente") == 1
    assert len(summary) == 1
    assert summary[0]["Maquinistas"] == "Ana, Bruno"
    assert summary[0]["Mangas"] == 2
    assert summary[0]["Fecha del pesaje final"] is None
    assert summary[0]["Estado fecha final"].startswith("Varias fechas finales:")
    assert summary[0]["KG final vigente"] == 8.0


def test_manga_explorer_export_quantity_requires_observed_un_article():
    base = {
        "ot": {"fecha_operativa": "2026-08-10", "turno": "DÍA", "codigo": "OT-Q"},
        "recurso": {}, "upstream": {"orden": {"codigo": "OF-Q"}},
        "maquinistas": [], "maquinistas_cobertura": "SIN_ASIGNACION",
        "visibilidad": {"pesaje": True},
        "produccion": {},
    }
    kg_item = {
        **base,
        "manga": {
            "codigo": "M-Q-KG", "public_id": "m-q-kg", "estado_operativo": "CERRADA",
            "color": "Azul", "articulo": {"nombre": "Pieza KG", "codigo": "PK", "unidad_inventario": "KG"},
            "pesaje": {
                "estado": "VIGENTE", "fecha_local_pesaje": "2026-08-10",
                "peso_fisico_neto_kg": 3, "cantidad_confirmada": 7,
                "fuente_cantidad": "PLAN_CONFIRMADO_POR_PESAJE",
            },
        },
    }
    un_item = {
        **base,
        "manga": {
            "codigo": "M-Q-UN", "public_id": "m-q-un", "estado_operativo": "CERRADA",
            "color": "Azul", "articulo": {"nombre": "Pieza UN", "codigo": "PU", "unidad_inventario": "UN"},
            "pesaje": {
                "estado": "VIGENTE", "fecha_local_pesaje": "2026-08-10",
                "peso_fisico_neto_kg": 3, "cantidad_confirmada": 7,
                "fuente_cantidad": "PLAN_CONFIRMADO_POR_PESAJE",
            },
        },
    }
    plan_un_item = {
        **un_item,
        "manga": {
            **un_item["manga"],
            "codigo": "M-Q-PLAN-UN",
        },
    }
    _summary, detail = _export_rows({"items": [kg_item, un_item, plan_un_item]})
    by_code = {row["Manga"]: row for row in detail}
    assert by_code["M-Q-KG"]["Cantidad por bulto"] is None
    assert "no acreditada como UN" in by_code["M-Q-KG"]["Observación"]
    assert by_code["M-Q-UN"]["Cantidad por bulto"] is None
    assert "no acreditada como UN" in by_code["M-Q-UN"]["Observación"]
    assert by_code["M-Q-PLAN-UN"]["Cantidad por bulto"] is None
    assert "no acreditada como UN" in by_code["M-Q-PLAN-UN"]["Observación"]
    audit = _audit_rows(detail)
    audit_by_code = {row["Manga"]: row for row in audit}
    assert audit_by_code["M-Q-PLAN-UN"]["Cantidad UN acreditada"] is None
    assert "no acreditada como UN" in audit_by_code["M-Q-PLAN-UN"]["Observación técnica"]


def test_manga_explorer_export_quantity_keeps_observed_un_source():
    item = {
        "ot": {"fecha_operativa": "2026-08-10", "turno": "DÍA", "codigo": "OT-Q"},
        "recurso": {}, "upstream": {"orden": {"codigo": "OF-Q"}},
        "maquinistas": [], "maquinistas_cobertura": "SIN_ASIGNACION",
        "visibilidad": {"pesaje": True}, "produccion": {},
        "manga": {
            "codigo": "M-Q-OBS-UN", "public_id": "m-q-obs-un", "estado_operativo": "CERRADA",
            "color": "Azul", "articulo": {"nombre": "Pieza UN", "codigo": "PU", "unidad_inventario": "UN"},
            "pesaje": {
                "estado": "VIGENTE", "fecha_local_pesaje": "2026-08-10",
                "peso_fisico_neto_kg": 3, "cantidad_confirmada": 7,
                "fuente_cantidad": "RESPONSABLE_ARMADO",
            },
        },
    }
    _summary, detail = _export_rows({"items": [item]})
    assert detail[0]["Cantidad por bulto"] == 7


def test_manga_explorer_summary_does_not_count_annulled_as_pending():
    item = {
        "ot": {"fecha_operativa": "2026-08-10", "turno": "DÍA", "codigo": "OT-A"},
        "recurso": {}, "upstream": {"orden": {"codigo": "OF-A"}},
        "maquinistas": [], "maquinistas_cobertura": "SIN_ASIGNACION",
        "visibilidad": {"pesaje": True},
        "manga": {
            "codigo": "M-A", "public_id": "m-a", "estado_operativo": "ANULADA",
            "color": "Rojo", "articulo": {"nombre": "Pieza", "codigo": "P"},
            "pesaje": None,
        },
    }
    summary, _detail = _export_rows({"items": [item]})
    assert summary[0]["Mangas anuladas"] == 1
    assert summary[0]["Pendientes de pesaje"] == 0


def test_manga_explorer_export_projects_recipe_materials_and_snapshot_cycles_without_cross_output_sum():
    template = {
        "ot": {"fecha_operativa": "2026-08-10", "turno": "DÍA", "codigo": "OT-1"},
        "recurso": {"nombre": "Máquina 1", "codigo": "MAQ-1"},
        "upstream": {"orden": {"codigo": "OF-1"}},
        "maquinistas": [], "maquinistas_cobertura": "SIN_ASIGNACION",
        "visibilidad": {"pesaje": True},
        "produccion": {
            "tipo_maquina": "Inyectora",
            "tipo_materia_prima": "MATERIA_PRIMA",
            "materias_primas": [
                {"codigo": "MP-1", "nombre": "Resina"},
                {"codigo": "MP-2", "nombre": "Carga"},
                {"codigo": "MP-3", "nombre": "Aditivo base"},
            ],
            "receta": {"id": 7, "revision": 4, "hash": "recipe-hash", "status": "OK"},
            "molde": {"codigo": "MOL-1", "nombre": "Molde histórico", "status": "OK"},
        },
    }
    first = {
        **template,
        "manga": {
            "codigo": "M-1", "public_id": "m-1", "estado_operativo": "CERRADA",
            "color": "Rojo", "articulo": {"nombre": "Pieza", "codigo": "P-1"},
            "pesaje": {"estado": "VIGENTE", "fecha_local_pesaje": "2026-08-10", "peso_fisico_neto_kg": 8},
        },
    }
    first["produccion"] = {
        **template["produccion"],
        "salida": {"id": "salida-1", "status": "OK", "cantidad_por_ciclo_snapshot": 4, "peso_unitario_snapshot_g": 1000},
    }
    second = {
        **first,
        "manga": {**first["manga"], "codigo": "M-2", "public_id": "m-2", "pesaje": {**first["manga"]["pesaje"], "peso_fisico_neto_kg": 4}},
        "produccion": {**first["produccion"], "salida": {"id": "salida-2", "status": "OK", "cantidad_por_ciclo_snapshot": 2, "peso_unitario_snapshot_g": 1000}},
    }
    summary, detail = _export_rows({"items": [first, second]})
    assert detail[0]["Material 3"] == "Aditivo base"
    assert detail[0]["Receta ID"] == 7 and detail[0]["Receta revisión"] == 4
    assert detail[0]["Receta hash"] == "recipe-hash"
    assert detail[0]["Ciclos aproximados"] == 2.0
    assert detail[1]["Ciclos aproximados"] == 2.0
    assert len(summary) == 2  # different immutable output bases are not additive
    assert {row["Ciclos aproximados"] for row in summary} == {2.0}
    assert "Material 3" in key_fields(3) and "Material 3" in detail_fields(3)


def test_manga_explorer_recipe_projection_excludes_pigments_and_marks_missing_or_zero_base():
    material = SimpleNamespace(codigo="MP-1", nombre="Resina", clase="MATERIA_PRIMA", materia_prima=SimpleNamespace(tipo="VIRGEN"))
    second = SimpleNamespace(codigo="MP-2", nombre="Recuperado", clase="MATERIA_PRIMA", materia_prima=SimpleNamespace(tipo="SEGUNDA"))
    pigment = SimpleNamespace(codigo="COL-1", nombre="Pigmento", clase="COLORANTE", materia_prima=None)
    recipe = SimpleNamespace(
        id=9, revision=3,
        lineas=[
            SimpleNamespace(id=2, orden=2, tipo_componente="COLORANTE", material=pigment),
            SimpleNamespace(id=1, orden=1, tipo_componente="MATERIA_PRIMA", material=material),
            SimpleNamespace(id=3, orden=3, tipo_componente="MATERIA_PRIMA", material=second),
        ],
    )
    projected = _material_recipe_context(recipe, recipe_id=9, recipe_hash="h")
    assert [item["codigo"] for item in projected["materials"]] == ["MP-1", "MP-2"]
    assert projected["materials"][0]["tipo"] == "VIRGEN"
    complete_recipe = SimpleNamespace(id=11, revision=1, lineas=[
        SimpleNamespace(id=10, orden=1, tipo_componente="MATERIA_PRIMA", material=material, cantidad="0.4", unidad="FRACCION"),
        SimpleNamespace(id=11, orden=2, tipo_componente="MATERIA_PRIMA", material=second, cantidad="0.2", unidad="FRACCION"),
        SimpleNamespace(id=12, orden=3, tipo_componente="MATERIA_PRIMA", material=material, cantidad="0.4", unidad="FRACCION"),
    ])
    complete_projection = _material_recipe_context(complete_recipe, recipe_id=11, recipe_hash="h11")
    assert complete_projection["fraccion_virgen_receta"] == Decimal("0.8")
    incomplete_recipe = SimpleNamespace(id=12, revision=1, lineas=[
        SimpleNamespace(id=13, orden=1, tipo_componente="MATERIA_PRIMA", material=material, cantidad="0.4", unidad="FRACCION"),
        SimpleNamespace(id=14, orden=2, tipo_componente="MATERIA_PRIMA", material=second, cantidad="0.2", unidad="FRACCION"),
    ])
    assert _material_recipe_context(incomplete_recipe, recipe_id=12, recipe_hash="h12")["fraccion_virgen_receta"] is None
    negative_recipe = SimpleNamespace(id=13, revision=1, lineas=[
        SimpleNamespace(id=15, orden=1, tipo_componente="MATERIA_PRIMA", material=material, cantidad="-0.1", unidad="FRACCION"),
        SimpleNamespace(id=16, orden=2, tipo_componente="MATERIA_PRIMA", material=second, cantidad="1.1", unidad="FRACCION"),
    ])
    assert _material_recipe_context(negative_recipe, recipe_id=13, recipe_hash="h13")["fraccion_virgen_receta"] is None
    missing_type = SimpleNamespace(codigo="MP-3", nombre="Sin subtipo", clase="MATERIA_PRIMA", materia_prima=None)
    missing_recipe = SimpleNamespace(id=10, revision=1, lineas=[
        SimpleNamespace(id=4, orden=1, tipo_componente="MATERIA_PRIMA", material=missing_type),
    ])
    missing_projection = _material_recipe_context(missing_recipe, recipe_id=10, recipe_hash="h2")
    assert missing_projection["materials"][0]["tipo"] is None
    assert missing_projection["status"] == "TIPO_MATERIA_PRIMA_AUSENTE"
    assert _material_recipe_context(None, recipe_id=None, recipe_hash=None)["status"] == "RECETA_AUSENTE"
    item = {
        "ot": {"fecha_operativa": "2026-08-10", "turno": "DÍA", "codigo": "OT-1"},
        "recurso": {}, "upstream": {"orden": {"codigo": "OF-1"}},
        "maquinistas": [], "visibilidad": {"pesaje": True},
        "manga": {"codigo": "M-0", "public_id": "m-0", "estado_operativo": "CERRADA", "color": "Rojo", "articulo": {"nombre": "Pieza", "codigo": "P-1"}, "pesaje": {"estado": "VIGENTE", "fecha_local_pesaje": "2026-08-10", "peso_fisico_neto_kg": 0}},
        "produccion": {"receta": projected, "molde": {"status": "OK"}, "salida": {"status": "OK", "cantidad_por_ciclo_snapshot": 0, "peso_unitario_snapshot_g": 100}},
    }
    _summary, detail = _export_rows({"items": [item]})
    assert detail[0]["Ciclos aproximados"] is None
    assert detail[0]["Estado ciclos aproximados"] == "BASE_CICLO_INVALIDA"


def test_manga_explorer_effective_corrida_never_falls_back_to_other_corrida_output():
    article = SimpleNamespace(id=42)
    wrong_output = SimpleNamespace(
        id="other", articulo_scm_id=42, corrida_fabricacion_id="corrida-2",
        cantidad_por_ciclo_snapshot=4, peso_unitario_snapshot_g=100,
    )
    work = SimpleNamespace(
        trabajo_color=SimpleNamespace(
            corrida=SimpleNamespace(id="corrida-1", receta_revision_id=None, receta_hash=None),
            receta_revision_id_snapshot=None, receta_hash_snapshot=None,
            molde_codigo_snapshot=None,
        ),
        orden_operacion=SimpleNamespace(salidas=[wrong_output], fabricacion=None),
    )
    context = _production_context(work, article, recipes={}, molds={})
    assert context["salida"]["status"] == "SALIDA_AUSENTE"
    assert context["salida"]["id"] is None


def test_manga_explorer_export_criteria_normalizes_state_and_matches_rows(app, client, scm_config):
    with app.app_context():
        seeded = _seed_observability_graph()
        manga = ScmManga.query.filter_by(codigo="M-OT-OBS-FAB-03").one()
        manga.estado = "PESADA"
        db.session.commit()
        exported = client.get(
            "/api/scm/v1/observabilidad/mangas-fabricacion/export.xlsx",
            query_string={"vista": "DETALLADO", "estado_manga": "PESADA"},
            headers=_headers(seeded["full"]),
        )
        assert exported.status_code == 200
        workbook = load_workbook(BytesIO(exported.data), read_only=False, data_only=False)
        criteria = list(workbook["Criterios"].iter_rows(values_only=True))
        assert criteria[5] == ("Estado", "Pesada")
        rows = list(workbook["Detalle"].iter_rows(values_only=True))
        state_index = rows[0].index("Estado operativo")
        assert rows[1:]
        assert all(row[state_index] == "Pesada" for row in rows[1:])


@pytest.mark.parametrize("query", [
    "=HYPERLINK('https://example.invalid','x')",
    "+HYPERLINK('https://example.invalid','x')",
    "-HYPERLINK('https://example.invalid','x')",
    "@HYPERLINK('https://example.invalid','x')",
])
def test_manga_explorer_export_criteria_escapes_formula_like_search(query, app, client, scm_config):
    with app.app_context():
        seeded = _seed_observability_graph()
        exported = client.get(
            "/api/scm/v1/observabilidad/mangas-fabricacion/export.xlsx",
            query_string={"vista": "DETALLADO", "q": query},
            headers=_headers(seeded["full"]),
        )
        assert exported.status_code == 200
        workbook = load_workbook(BytesIO(exported.data), read_only=False, data_only=False)
        cell = workbook["Criterios"]["B5"]
        assert cell.value == "'" + query.casefold()
        assert cell.data_type == "s"


def _add_real_export_context():
    """Attach a small ORM-only OF/corrida/recipe graph to the test database."""
    from app.models.producto import ColorBase, ColorProduccion, FamiliaColor

    work = ScmTrabajoOt.query.filter_by(codigo="TC-OBS-AZUL").one()
    manga = ScmManga.query.filter_by(codigo="M-OT-OBS-FAB-03").one()
    fixture_type = TipoMaquina(codigo="XLSX-INY", nombre="HT200A", proceso="INYECCION")
    fixture_machine = Maquina(codigo="XLSX-MAQ", nombre="Prensa fixture", tipo_maquina=fixture_type)
    db.session.add(fixture_machine)
    db.session.flush()
    work.orden_trabajo.maquina_id = fixture_machine.id
    article = manga.lote_articulo.articulo if manga.lote_articulo else ScmArticulo.query.first()
    if article is None:
        article = ScmArticulo(
            codigo="XLSX-PIEZA-1", nombre="Pieza XLSX", clase="PIEZA_COLOR",
            unidad_base="UN", unidad_inventario="KG",
        )
        db.session.add(article)
        db.session.flush()
    mold = Molde.query.first()
    if mold is None:
        mold = Molde(codigo="XLSX-MOLDE", nombre="Molde XLSX", peso_tiro_gr=500)
        db.session.add(mold)
        db.session.flush()
    category = ScmCategoriaRecepcion.query.first()
    if category is None:
        category = ScmCategoriaRecepcion(
            codigo="XLSX-CAT", nombre="Categoría XLSX",
            modalidad_default="VIRGEN_CONFIANZA_PROVEEDOR",
            recepcion_habilitada=True,
        )
        db.session.add(category)
        db.session.flush()
    color = ColorProduccion.query.first()
    if color is None:
        color_base = ColorBase(nombre="XLSX-ROJO")
        familia = FamiliaColor(codigo=987654, nombre="XLSX-SOLIDO")
        db.session.add_all([color_base, familia])
        db.session.flush()
        color = ColorProduccion(color_base_id=color_base.id, familia_color_id=familia.id)
        db.session.add(color)
        db.session.flush()
    assert article is not None
    operation = work.orden_operacion
    fabrication = operation.fabricacion
    if fabrication is None:
        fabrication = ScmOrdenFabricacion(
            orden_operacion=operation, molde_id=mold.codigo,
            snapshot_proceso="INYECCION", fuente_proceso="EXPLICITO",
        )
        db.session.add(fabrication)
        db.session.flush()
    materials = []
    for index, tipo in enumerate(("VIRGEN", "SEGUNDA", "VIRGEN"), start=1):
        scm_material = ScmMaterial(
            codigo=f"XLSX-MP-{index}", nombre=f"Materia {index}", clase="MATERIA_PRIMA",
            categoria_recepcion_id=category.id,
        )
        db.session.add(scm_material)
        db.session.flush()
        materials.append(MateriaPrima(
            nombre=f"Materia {index}", tipo=tipo, scm_material_id=scm_material.id,
        ))
    db.session.add_all(materials)
    db.session.flush()
    recipe = RecetaColorMaestra(
        color_produccion_id=color.id, producto_scope="*", nombre_variante="XLSX-W2",
        revision=4, estado="APROBADA", es_default=False, base_virgen_kg=25,
    )
    db.session.add(recipe)
    db.session.flush()
    fractions = ("0.40", "0.20", "0.40")
    db.session.add_all([
        RecetaColorLinea(receta_id=recipe.id, material_id=materials[index].scm_material_id,
                         tipo_componente="MATERIA_PRIMA", cantidad=fractions[index], unidad="FRACCION", orden=index)
        for index in range(3)
    ])
    corrida = ScmCorridaFabricacion(
        orden_fabricacion_id=operation.id, codigo="XLSX-W2-CORRIDA", secuencia=1,
        color_produccion_id=color.id, receta_revision_id=recipe.id,
        receta_hash="c" * 64, estado="EN_EJECUCION",
    )
    db.session.add(corrida)
    db.session.flush()
    other = ScmCorridaFabricacion(
        orden_fabricacion_id=operation.id, codigo="XLSX-W2-OTHER", secuencia=2,
        color_produccion_id=color.id, receta_revision_id=recipe.id,
        receta_hash="d" * 64, estado="EN_EJECUCION",
    )
    db.session.add(other)
    db.session.flush()
    work.trabajo_color.corrida = corrida
    work.trabajo_color.receta_revision_id_snapshot = recipe.id
    work.trabajo_color.receta_hash_snapshot = "c" * 64
    work.trabajo_color.molde_codigo_snapshot = mold.codigo
    outputs = [
        ScmOrdenOperacionSalida(
            orden_operacion_id=operation.id, corrida_fabricacion_id=corrida.id,
            articulo_scm_id=article.id, cantidad_por_ciclo_snapshot=4,
            peso_unitario_snapshot_g=1000, cantidad_objetivo=100,
        ),
        ScmOrdenOperacionSalida(
            orden_operacion_id=operation.id, corrida_fabricacion_id=other.id,
            articulo_scm_id=article.id, cantidad_por_ciclo_snapshot=99,
            peso_unitario_snapshot_g=999, cantidad_objetivo=100,
        ),
    ]
    db.session.add_all(outputs)
    db.session.flush()
    if manga.lote_articulo is None:
        lot = ScmLoteArticulo(
            codigo="XLSX-W2-LOT", articulo_id=article.id, clase="SALIDA_ORDEN_OPERACION",
            orden_operacion_salida_id=outputs[0].id, cantidad_acreditada=1,
        )
        db.session.add(lot)
        db.session.flush()
        manga.lote_articulo_id = lot.id
    if manga.asignacion_personal_trabajo is None:
        machine_worker = Trabajador.query.filter_by(codigo="TRB-01").one()
        assignment = ScmAsignacionPersonalTrabajoOt(
            trabajo=work,
            trabajador=machine_worker,
            estado="PREVISTA",
            asignada_por_id=machine_worker.id,
        )
        db.session.add(assignment)
        db.session.flush()
        manga.asignacion_personal_trabajo = assignment
    db.session.commit()


def test_manga_explorer_real_export_fixture_contains_new_traceability_fields(app, client, scm_config):
    """Optional evidence hook: writes only when the caller requests a path."""
    target = os.environ.get("MANGA_EXPORT_FIXTURE_PATH")
    with app.app_context():
        seeded = _seed_observability_graph()
        _add_real_export_context()
        response = client.get(
            "/api/scm/v1/observabilidad/mangas-fabricacion/export.xlsx?vista=AMBOS",
            headers=_headers(seeded["full"]),
        )
        assert response.status_code == 200
        workbook = load_workbook(BytesIO(response.data), read_only=True, data_only=False)
        for sheet_name in ("Resumen", "Detalle"):
            headers = list(next(workbook[sheet_name].iter_rows(values_only=True)))
            assert "Tipo máquina" in headers
            assert "Material 1" in headers and "Material 2" in headers
            assert "Material 3" in headers
            assert "Molde" in headers and "Ciclos aproximados" in headers
            rows = list(workbook[sheet_name].iter_rows(values_only=True))
            assert any(row[headers.index("Tipo máquina")] == "INYECCION" for row in rows[1:])
            assert any(row[headers.index("Ciclos aproximados")] is not None for row in rows[1:])
            assert any(row[headers.index("Tipo materia prima")] == "VIRGEN, SEGUNDA" for row in rows[1:])
            assert any(row[headers.index("Fracción virgen (receta)")] == 0.8 for row in rows[1:])
            cycle_column = next(
                cell.column for cell in workbook[sheet_name][1]
                if cell.value == "Ciclos aproximados"
            )
            assert workbook[sheet_name].cell(row=2, column=cycle_column).number_format == "0.0"
        if target:
            Path(target).write_bytes(response.data)


def test_manga_explorer_export_keeps_operational_columns_before_audit(app, client, scm_config):
    with app.app_context():
        seeded = _seed_observability_graph()
        _add_real_export_context()
        response = client.get(
            "/api/scm/v1/observabilidad/mangas-fabricacion/export.xlsx?vista=AMBOS",
            headers=_headers(seeded["full"]),
        )
        assert response.status_code == 200
        workbook = load_workbook(BytesIO(response.data), read_only=False, data_only=False)
        assert workbook.sheetnames == ["Resumen", "Detalle", "Auditoría", "Criterios"]
        summary_headers = list(next(workbook["Resumen"].iter_rows(values_only=True)))
        detail_headers = list(next(workbook["Detalle"].iter_rows(values_only=True)))
        audit_headers = list(next(workbook["Auditoría"].iter_rows(values_only=True)))
        expected_operational = [
            "Tipo máquina", "Máquina", "Fecha de OT", "Maquinistas",
            "Tipo materia prima", "Material 1", "Material 2", "Material 3", "Molde",
            "Pieza-Color", "Ciclos aproximados", "Peso real (kg)", "Turno", "Cantidad por bulto",
            "Bultos pesados", "N°OF", "Fracción virgen (receta)",
        ]
        assert summary_headers[:len(expected_operational)] == expected_operational
        assert detail_headers[:len(expected_operational)] == expected_operational
        assert "Receta hash" not in summary_headers
        assert "Receta hash" not in detail_headers
        assert "Receta hash" in audit_headers
        assert "Estado datos producción" in audit_headers
        assert summary_headers[-1] == "Referencia grupo"
        assert detail_headers[-1] == "Referencia grupo"
        assert "Fracción virgen (receta)" in summary_headers
        assert workbook["Resumen"].cell(row=2, column=summary_headers.index("Fracción virgen (receta)") + 1).number_format == "0.0%"
        assert workbook["Detalle"].auto_filter.ref.startswith("A1:")
        assert workbook["Detalle"].freeze_panes == "A2"
        assert workbook["Auditoría"].max_row > 1
        assert workbook["Resumen"]["A1"].fill.fgColor.rgb.endswith("1F4E78")
        assert workbook["Resumen"]["A1"].font.color.type == "rgb"
        audit_rows = list(workbook["Auditoría"].iter_rows(min_row=2, values_only=True))
        assert any(row[audit_headers.index("Proceso máquina")] == "INYECCION" for row in audit_rows)
        assert any(row[audit_headers.index("Nombre tipo máquina")] == "HT200A" for row in audit_rows)
        assert any(row[detail_headers.index("Maquinistas")] != "Sin maquinista asignado" for row in workbook["Detalle"].iter_rows(min_row=2, values_only=True))


def test_manga_explorer_export_production_projection_queries_are_bulk_bounded(app, client, scm_config):
    with app.app_context():
        seeded = _seed_observability_graph()
        _add_real_export_context()
        _add_scaled_manga_graph(seeded, count=25)
        counts = []
        for query in ("M-OT-OBS-FAB-03", None):
            statements = []

            def capture(_conn, _cursor, statement, _parameters, _context, _executemany):
                statements.append(statement)

            event.listen(db.engine, "before_cursor_execute", capture)
            try:
                params = {"vista": "DETALLADO"}
                if query:
                    params["q"] = query
                response = client.get(
                    "/api/scm/v1/observabilidad/mangas-fabricacion/export.xlsx",
                    query_string=params, headers=_headers(seeded["full"]),
                )
            finally:
                event.remove(db.engine, "before_cursor_execute", capture)
            assert response.status_code == 200
            counts.append(len(statements))
        assert counts[1] <= 120
        assert counts[1] - counts[0] <= 25


def test_manga_explorer_search_uses_effective_ot_after_auditable_corrections(app, client, scm_config):
    with app.app_context():
        first = _add_auditable_correction(app, "PESAJE-MANGA-EXPLORER-CORR-1")
        actor = first["manga"].maquinista_previsto

        effective = client.get(
            f"/api/scm/v1/observabilidad/mangas-fabricacion?q={first['effective_ot'].codigo_ot}",
            headers=_headers(actor),
        )
        assert effective.status_code == 200
        assert ScmCorreccionAsignacionManga.query.count() == 2
        assert ScmTramoMangaTrabajo.query.filter_by(
            secuencia=2, estado="ACTIVO",
        ).count() == 1
        assert ScmTramoMangaTrabajo.query.filter_by(
            manga_id=first["manga"].id, secuencia=2,
        ).count() == 0
        assert ScmTramoMangaTrabajo.query.filter_by(
            manga_id=first["second_manga"].id, secuencia=2,
        ).count() == 1
        effective_ids = {item["manga"]["public_id"] for item in effective.get_json()["items"]}
        assert str(first["manga"].public_id) in effective_ids

        origin = client.get(
            f"/api/scm/v1/observabilidad/mangas-fabricacion?q={first['origin_ot'].codigo_ot}",
            headers=_headers(actor),
        )
        assert origin.status_code == 200
        origin_ids = {item["manga"]["public_id"] for item in origin.get_json()["items"]}
        assert str(first["manga"].public_id) not in origin_ids

        effective_work = client.get(
            f"/api/scm/v1/observabilidad/mangas-fabricacion?q={first['effective_work'].codigo}",
            headers=_headers(actor),
        )
        assert effective_work.status_code == 200
        assert str(first["manga"].public_id) in {
            item["manga"]["public_id"] for item in effective_work.get_json()["items"]
        }
        origin_work = client.get(
            f"/api/scm/v1/observabilidad/mangas-fabricacion?q={first['manga'].trabajo.codigo}",
            headers=_headers(actor),
        )
        assert origin_work.status_code == 200
        assert str(first["manga"].public_id) not in {
            item["manga"]["public_id"] for item in origin_work.get_json()["items"]
        }

        listed = client.get(
            "/api/scm/v1/observabilidad/mangas-fabricacion?limit=100",
            headers=_headers(actor),
        )
        assert listed.status_code == 200
        by_id = {
            item["manga"]["public_id"]: item
            for item in listed.get_json()["items"]
        }
        active_item = by_id[str(first["manga"].public_id)]
        superseded_item = by_id[str(first["second_manga"].public_id)]
        for item in (active_item, superseded_item):
            worker_ids = [worker["id"] for worker in item["maquinistas"]]
            assert worker_ids == list(dict.fromkeys(worker_ids))
            assert all(worker.get("nombre") or worker.get("codigo") for worker in item["maquinistas"])

        active_work = first["active_work"]
        active_ot = first["active_ot"]
        latest_work = first["superseded_latest_work"]
        latest_ot = first["superseded_latest_ot"]
        superseded_ot = first["superseded_correction_ot"]
        for item, ot, work in (
            (active_item, active_ot, active_work),
            (superseded_item, latest_ot, latest_work),
        ):
            assert item["ot"]["public_id"] == str(ot.public_id)
            assert item["ot"]["codigo"] == ot.codigo_ot
            assert item["trabajo"]["codigo"] == work.codigo
            assert item["upstream"]["orden"]["codigo"] == work.orden_operacion.codigo
            assert item["recurso"]["codigo"] == ot.maquina.codigo
            assert item["manga"]["color"] == work.trabajo_color.color_nombre_snapshot
            assert item["trabajo"]["molde_codigo"] == work.trabajo_color.molde_codigo_snapshot

        search_payloads = {}

        def search_ids(term):
            response = client.get(
                "/api/scm/v1/observabilidad/mangas-fabricacion",
                query_string={"q": term, "limit": 100},
                headers=_headers(actor),
            )
            assert response.status_code == 200, (term, response.get_json())
            search_payloads[term] = response.get_json()
            return {
                item["manga"]["public_id"]
                for item in search_payloads[term]["items"]
            }

        active_id = str(first["manga"].public_id)
        superseded_id = str(first["second_manga"].public_id)
        destination_terms = (
            active_ot.codigo_ot,
            active_work.codigo,
            active_work.orden_operacion.codigo,
            active_ot.maquina_codigo_snapshot or active_ot.maquina.codigo,
            active_work.trabajo_color.color_nombre_snapshot,
            active_work.trabajo_color.molde_codigo_snapshot,
        )
        assert all(active_id in search_ids(term) for term in destination_terms)
        assert active_id not in search_ids(first["origin_ot"].codigo_ot)
        assert active_id not in search_ids(first["manga"].trabajo.codigo)

        latest_terms = (
            latest_ot.codigo_ot,
            latest_work.codigo,
            latest_work.orden_operacion.codigo,
            latest_ot.maquina_codigo_snapshot or latest_ot.maquina.codigo,
            latest_work.trabajo_color.color_nombre_snapshot,
            latest_work.trabajo_color.molde_codigo_snapshot,
        )
        for term in latest_terms:
            ids = search_ids(term)
            assert superseded_id in ids, (term, ids, search_payloads[term])
        assert superseded_id not in search_ids(first["origin_ot"].codigo_ot)
        assert superseded_id not in search_ids(first["second_manga"].trabajo.codigo)
        assert superseded_id not in search_ids(superseded_ot.codigo_ot)
        assert superseded_id not in search_ids(first["superseded_correction_work"].codigo)


def test_manga_explorer_search_matches_effective_of_without_resolvable_plan(app, client, scm_config):
    """The visible OF remains searchable when its plan/OP link is absent."""
    with app.app_context():
        seeded = _seed_observability_graph()
        source = ScmManga.query.filter_by(codigo="M-OT-OBS-FAB-02").one()
        source_work = source.trabajo
        operation = ScmOrdenOperacion(
            codigo="OF-NOPLAN-001",
            tipo="FABRICACION",
            origen_demanda="EXCEPCIONAL",
            propuesta_clave="OF-NOPLAN-001",
            estado="EN_EJECUCION",
            created_by_id=seeded["full"].id,
        )
        db.session.add(operation)
        db.session.flush()
        work = ScmTrabajoOt(
            orden_trabajo_id=source.ot_id,
            codigo="TC-NOPLAN-001",
            tipo="COLOR",
            secuencia=3,
            estado="PLANIFICADO",
            orden_operacion_id=operation.id,
            cantidad_objetivo_un=100,
            cantidad_confirmada_un=0,
            created_by_id=seeded["full"].id,
        )
        work.trabajo_color = ScmTrabajoColor(
            corrida_fabricacion_id=source_work.trabajo_color.corrida_fabricacion_id,
            color_nombre_snapshot=source_work.trabajo_color.color_nombre_snapshot,
            molde_codigo_snapshot=source_work.trabajo_color.molde_codigo_snapshot,
        )
        db.session.add(work)
        db.session.flush()
        source.trabajo = work
        db.session.commit()

        response = client.get(
            "/api/scm/v1/observabilidad/mangas-fabricacion",
            query_string={"q": operation.codigo, "limit": 100},
            headers=_headers(seeded["full"]),
        )

        assert response.status_code == 200
        payload = response.get_json()
        assert str(source.public_id) in {
            item["manga"]["public_id"] for item in payload["items"]
        }
        item = next(
            item for item in payload["items"]
            if item["manga"]["public_id"] == str(source.public_id)
        )
        assert item["upstream"]["orden"]["codigo"] == operation.codigo
        assert item["upstream"]["op"] is None
        assert item["upstream"]["plan"] is None


def test_manga_explorer_search_direct_of_survives_unrelated_segment(app, client, scm_config):
    """A segment on another manga must not globally suppress direct OF rows."""
    with app.app_context():
        seeded = _seed_observability_graph()
        direct = ScmManga.query.filter_by(codigo="M-OT-OBS-FAB-02").one()
        segmented = ScmManga.query.filter_by(codigo="M-OT-OBS-FAB-03").one()
        assignment = segmented.asignacion_personal_trabajo
        if assignment is None:
            assignment = ScmAsignacionPersonalTrabajoOt(
                trabajo=segmented.trabajo,
                trabajador=Trabajador.query.filter_by(codigo="TRB-01").one(),
                estado="PREVISTA",
                asignada_por_id=seeded["full"].id,
            )
            db.session.add(assignment)
            db.session.flush()
        db.session.add(ScmTramoMangaTrabajo(
            manga=segmented,
            trabajo=segmented.trabajo,
            asignacion_personal_trabajo=assignment,
            secuencia=1,
            estado="ACTIVO",
            cantidad_inicio_un=0,
            cantidad_atribuida_un=0,
            created_by_id=seeded["full"].id,
        ))
        db.session.commit()

        operation_code = direct.trabajo.orden_operacion.codigo
        response = client.get(
            "/api/scm/v1/observabilidad/mangas-fabricacion",
            query_string={"q": operation_code, "limit": 100},
            headers=_headers(seeded["full"]),
        )
        assert response.status_code == 200
        ids = {item["manga"]["public_id"] for item in response.get_json()["items"]}
        assert str(direct.public_id) in ids


@pytest.mark.parametrize("limit", [1, 25, 100])
def test_manga_explorer_uses_bounded_projection_queries(app, client, scm_config, limit):
    with app.app_context():
        _add_auditable_correction(app, "PESAJE-MANGA-EXPLORER-SCALE")
        seeded = _seed_observability_graph()
        _add_scaled_manga_graph(seeded)
        statements = []

        def capture(_conn, _cursor, statement, _parameters, _context, _executemany):
            statements.append(statement)

        event.listen(db.engine, "before_cursor_execute", capture)
        try:
            response = client.get(
                f"/api/scm/v1/observabilidad/mangas-fabricacion?limit={limit}",
                headers=_headers(seeded["full"]),
            )
        finally:
            event.remove(db.engine, "before_cursor_execute", capture)

        assert response.status_code == 200
        payload = response.get_json()
        assert len(payload["items"]) == limit
        assert payload["page"]["limit"] == limit
        assert payload["page"]["has_more"] is True
        assert ScmManga.query.count() == 127
        assert len(statements) == {1: 84, 25: 90, 100: 90}[limit]
        # The explorer hydrates a page in batches; a full manga-detail query per
        # row would grow linearly with the physical set and is not acceptable.
        # The eager projection is bounded by relation batches, not by the
        # number of physical mangas.  SQLite emits a few additional batches
        # once the page crosses its parameter/chunk boundaries, but it must
        # remain well below a per-row detail query graph.
        assert len(statements) <= 100


def test_manga_explorer_export_uses_canonical_final_and_excludes_annulled_weight(app, client, scm_config):
    with app.app_context():
        _add_auditable_correction(app, "PESAJE-MANGA-EXPLORER-EXPORT")
        seeded = _seed_observability_graph()
        _add_scaled_manga_graph(seeded)
        statements = []

        def capture(_conn, _cursor, statement, _parameters, _context, _executemany):
            statements.append(statement)

        event.listen(db.engine, "before_cursor_execute", capture)
        try:
            exported = client.get(
                "/api/scm/v1/observabilidad/mangas-fabricacion/export.xlsx?vista=DETALLADO",
                headers=_headers(seeded["full"]),
            )
        finally:
            event.remove(db.engine, "before_cursor_execute", capture)
        assert exported.status_code == 200
        assert len(statements) == 94
        workbook = load_workbook(BytesIO(exported.data), read_only=True, data_only=False)
        rows = list(workbook["Detalle"].iter_rows(values_only=True))
        assert len(rows) - 1 == 126
        headers = rows[0]
        by_manga = {row[headers.index("Manga")]: row for row in rows[1:]}
        corrected = by_manga["M-OT-OBS-FAB-02"]
        annulled = by_manga["M-OT-OBS-FAB-04"]
        assert corrected[headers.index("Peso real (kg)")] == 11.5
        assert annulled[headers.index("Peso real (kg)")] is None
        audit_headers = next(workbook["Auditoría"].iter_rows(values_only=True))
        audit_rows = list(workbook["Auditoría"].iter_rows(min_row=2, values_only=True))
        annulled_audit = next(row for row in audit_rows if row[audit_headers.index("Manga")] == "M-OT-OBS-FAB-04")
        assert annulled_audit[audit_headers.index("Estado pesaje")] == "SIN_PESAJE_VIGENTE"


def test_export_without_weighing_permission_is_explicitly_restricted(app, client, scm_config):
    with app.app_context():
        seeded = _seed_observability_graph()
        exported = client.get(
            "/api/scm/v1/observabilidad/mangas-fabricacion/export.xlsx?vista=AMBOS",
            headers=_headers(seeded["base"]),
        )
        assert exported.status_code == 200
        workbook = load_workbook(BytesIO(exported.data), read_only=True, data_only=False)
        summary_headers = next(workbook["Resumen"].iter_rows(values_only=True))
        detail_headers = next(workbook["Detalle"].iter_rows(values_only=True))
        assert "Mangas con pesaje restringido" in summary_headers
        assert "KG abierto" not in detail_headers
        audit_headers = next(workbook["Auditoría"].iter_rows(values_only=True))
        assert "Estado pesaje" in audit_headers
        summary_rows = list(workbook["Resumen"].iter_rows(min_row=2, values_only=True))
        restricted_index = summary_headers.index("Mangas con pesaje restringido")
        open_index = summary_headers.index("Pendientes de pesaje")
        assert any(row[restricted_index] for row in summary_rows)
        assert all(row[open_index] in (None, 0) for row in summary_rows)


def test_manga_explorer_keeps_formula_like_cells_as_text_and_hex_is_canonical():
    assert [_safe_cell(value) for value in ("=SUM(A1)", "+unsafe", "-unsafe", "@unsafe")] == [
        "'=SUM(A1)", "'+unsafe", "'-unsafe", "'@unsafe",
    ]
    assert _safe_cell("plain") == "plain"
    assert _canonical_color_hex(" #aBc123 ") == "#ABC123"
    assert _canonical_color_hex("transparent") is None
    assert _canonical_color_hex("#12345") is None
