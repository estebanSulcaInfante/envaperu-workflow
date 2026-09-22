"""Scenario tests for KG availability and PT manual Kardex projection."""

from datetime import date
from uuid import uuid4
from openpyxl import load_workbook

import pytest

from app.extensions import db
from app.models.scm_articulos import ScmArticulo
from app.models.scm_articulos import ScmArticuloPiezaColor
from app.models.scm_catalogos import ScmCapacidad
from app.models.scm_estructuras import ScmEstructuraComponente, ScmEstructuraRevision
from app.models.scm_inventory import ScmMovimientoInventario, ScmSaldoInventario, ScmUbicacionInventario
from app.models.scm_inventory_kg import ScmSaldoInventarioKg
from app.models.scm_inventory_operations import ScmAlmacen, ScmAlmacenTrabajador
from app.models.scm_auditoria import ScmOperacion
from app.models.trabajador import Trabajador
from app.models.producto import ColorBase, ColorProduccion, FamiliaColor, PiezaColor
from app.models.molde import Pieza
from app.services.scm_kg_pt_availability_service import (
    _hash,
    list_piece_kg_availability,
    list_pt_availability,
    list_pt_manual_movements,
    register_pt_manual_movement,
)
from app.services.scm_kg_pt_availability_export import (
    MAX_MATRIX_COMPONENTS,
    _safe_text,
    generate_pt_availability_xlsx,
)
from app.services.scm_service_support import ScmServiceError


def _actor_with_caps():
    from flask import current_app

    current_app.config["PT_MANUAL_WRITE_ENABLED"] = True
    actor = Trabajador.query.first()
    role = actor.roles[0]
    for code in ("INVENTARIO_VER", "INVENTARIO_AJUSTAR", "INVENTARIO_PT_MOVIMIENTO"):
        capability = ScmCapacidad.query.filter_by(codigo=code).first()
        if capability is None:
            capability = ScmCapacidad(codigo=code, nombre=code)
            db.session.add(capability)
            db.session.flush()
        if capability not in role.capacidades:
            role.capacidades.append(capability)
    db.session.commit()
    return actor


def _article(code, name, article_class):
    item = ScmArticulo(
        codigo=code,
        nombre=name,
        clase=article_class,
        unidad_base="UN",
        unidad_inventario="UN",
    )
    db.session.add(item)
    db.session.flush()
    return item


def test_pt_manual_uses_canonical_un_ledger_and_replays(app):
    with app.app_context():
        actor = _actor_with_caps()
        product = _article("PT-MAN-01", "PT manual de prueba", "PRODUCTO_TERMINADO")
        location = ScmUbicacionInventario(codigo="PT-MESA", nombre="Mesa PT")
        db.session.add(location)
        db.session.commit()

        with pytest.raises(ScmServiceError) as error:
            register_pt_manual_movement(
                db.session,
                actor_id=actor.id,
                operation_id=uuid4(),
                data={
                    "articulo_scm_id": product.id,
                    "ubicacion_id": location.id,
                    "tipo": "SALDO_INICIAL",
                    "cantidad": 2,
                    "fecha_operativa": "2026-09-19",
                    "motivo": "No debe saltar el mecanismo de apertura",
                    "referencia": "APERTURA-INVALIDA",
                },
            )
        assert error.value.code == "PT_MANUAL_MOVEMENT_TYPE_INVALID"

        operation_id = uuid4()
        command = {
            "articulo_scm_id": product.id,
            "ubicacion_id": location.id,
            "tipo": "ENTRADA",
            "cantidad": 5,
            "fecha_operativa": "2026-09-19",
            "motivo": "Entrada PT manual del piloto",
            "referencia": "ACTA-PT-01",
            "version": 1,
        }
        first = register_pt_manual_movement(
            db.session, actor_id=actor.id, operation_id=operation_id, data=command,
        )
        replay = register_pt_manual_movement(
            db.session, actor_id=actor.id, operation_id=operation_id, data=command,
        )
        assert first == replay
        balance = db.session.scalar(db.select(ScmSaldoInventario).where(ScmSaldoInventario.articulo_scm_id == product.id))
        assert balance.cantidad_fisica == 5
        movement = db.session.scalar(db.select(ScmMovimientoInventario).where(ScmMovimientoInventario.saldo_id == balance.id))
        assert movement.tipo == "ENTRADA_MANUAL_PT"
        assert movement.fecha_operativa == date(2026, 9, 19)
        assert movement.referencia == "ACTA-PT-01"

        history = list_pt_manual_movements(db.session, actor_id=actor.id, balance_id=balance.id)
        assert len(history["items"]) == 1
        assert history["items"][0]["saldo_resultante"] == "5.000"


def test_piece_availability_is_empty_honestly_without_kg_projection(app):
    with app.app_context():
        actor = _actor_with_caps()
        payload = list_piece_kg_availability(db.session, actor_id=actor.id)
        assert payload["items"] == []
        assert payload["politica_piloto"] == "SIN_CONTROL_CALIDAD_DESDE_PESAJE"


def test_pt_manual_reference_is_required_bounded_and_has_no_effect_on_rejection(app):
    with app.app_context():
        actor = _actor_with_caps()
        product = _article("PT-REF-01", "PT referencia", "PRODUCTO_TERMINADO")
        location = ScmUbicacionInventario(codigo="PT-REF", nombre="Ubicacion referencia")
        db.session.add(location)
        db.session.commit()
        base = {
            "articulo_scm_id": product.id, "ubicacion_id": location.id,
            "tipo": "ENTRADA", "cantidad": 2, "version": 1,
            "fecha_operativa": "2026-09-19", "motivo": "Prueba referencia",
        }
        for reference in (None, "   ", "x" * 121):
            with pytest.raises(ScmServiceError) as error:
                register_pt_manual_movement(
                    db.session, actor_id=actor.id, operation_id=uuid4(),
                    data={**base, "referencia": reference},
                )
            assert error.value.code == "PT_MANUAL_REFERENCE_REQUIRED"
            assert error.value.status_code == 422
            db.session.rollback()
        assert db.session.scalar(db.select(ScmSaldoInventario).where(
            ScmSaldoInventario.articulo_scm_id == product.id,
        )) is None
        assert db.session.scalar(db.select(ScmOperacion).where(
            ScmOperacion.endpoint == "POST /inventario/pt/movimientos",
        )) is None


def test_pt_manual_replays_completed_legacy_operation_without_reference(app):
    with app.app_context():
        actor = _actor_with_caps()
        product = _article("PT-REF-LEGACY", "PT referencia legacy", "PRODUCTO_TERMINADO")
        location = ScmUbicacionInventario(codigo="PT-REF-LEGACY", nombre="Ubicacion legacy")
        db.session.add(location)
        db.session.flush()
        operation_id = uuid4()
        command = {
            "articulo_scm_id": product.id, "ubicacion_id": location.id,
            "tipo": "ENTRADA_MANUAL_PT", "cantidad": "2.000",
            "fecha_operativa": "2026-09-19", "motivo": "Movimiento anterior",
            "referencia": None, "version": 1,
        }
        response = {"movement": {"id": "legacy", "referencia": None}, "saldo": {"version": 2}}
        db.session.add(ScmOperacion(
            operation_id=operation_id, endpoint="POST /inventario/pt/movimientos",
            actor_id=actor.id,
            request_sha256=_hash({
                "endpoint": "POST /inventario/pt/movimientos",
                "actor_id": actor.id,
                "data": command,
            }),
            response_json=response, estado_http=201,
        ))
        db.session.commit()

        replay = register_pt_manual_movement(
            db.session, actor_id=actor.id, operation_id=operation_id, data={
                "articulo_scm_id": product.id, "ubicacion_id": location.id,
                "tipo": "ENTRADA", "cantidad": 2, "version": 1,
                "fecha_operativa": "2026-09-19", "motivo": "Movimiento anterior",
            },
        )
        assert replay == response
        assert db.session.scalar(db.select(ScmSaldoInventario).where(
            ScmSaldoInventario.articulo_scm_id == product.id,
        )) is None


def test_pt_manual_rejects_stale_version_and_kg_uses_authoritative_saldo(app):
    with app.app_context():
        actor = _actor_with_caps()
        product = _article("PT-VERSION-01", "PT versionado", "PRODUCTO_TERMINADO")
        piece = _article("PC-AUTH-01", "Pieza autoridad", "PIEZA_COLOR")
        location = ScmUbicacionInventario(codigo="PT-VERSION", nombre="Ubicacion version")
        db.session.add(location)
        db.session.flush()
        db.session.add(ScmSaldoInventarioKg(
            articulo_scm_id=piece.id, ubicacion_id=location.id,
            cantidad_fisica_kg=10, cantidad_reservada_kg=2,
            cantidad_no_disponible_kg=1, cantidad_retirada_kg=3,
            atributo_proceso="PROCESO",
        ))
        db.session.commit()
        first = register_pt_manual_movement(
            db.session, actor_id=actor.id, operation_id=uuid4(), data={
                "articulo_scm_id": product.id, "ubicacion_id": location.id,
                "tipo": "ENTRADA", "cantidad": 2, "version": 1,
                "fecha_operativa": "2026-09-19", "motivo": "Alta PT", "referencia": "ALTA-PT-01",
            },
        )
        assert first["saldo"]["version"] == 2
        with pytest.raises(ScmServiceError) as error:
            register_pt_manual_movement(
                db.session, actor_id=actor.id, operation_id=uuid4(), data={
                    "articulo_scm_id": product.id, "ubicacion_id": location.id,
                    "tipo": "SALIDA", "cantidad": 1, "version": 1,
                    "fecha_operativa": "2026-09-19", "motivo": "Version vieja", "referencia": "SALIDA-PT-01",
                },
            )
        assert error.value.code == "VERSION_CONFLICT"
        db.session.rollback()
        payload = list_piece_kg_availability(db.session, actor_id=actor.id, query="PC-AUTH")
        assert payload["items"][0]["kg_medidos"] == "10.000"
        assert payload["items"][0]["kg_retirados"] == "3.000"
        assert payload["items"][0]["kg_disponibles"] == "7.000"
        assert payload["as_of"].endswith("+00:00")


def test_pt_projection_shared_stock_wip_missing_weight_floor_and_pt_query(app):
    with app.app_context():
        actor = _actor_with_caps()
        actor_id = actor.id
        location = ScmUbicacionInventario(codigo="PT-BOM-LOC", nombre="Ubicacion BOM")
        db.session.add(location)
        db.session.flush()
        piece = _article("PC-BOM-01", "Pieza BOM", "PIEZA_COLOR")
        missing_weight = _article("PC-BOM-NO-WEIGHT", "Pieza sin peso", "PIEZA_COLOR")
        wip = _article("WIP-BOM-01", "WIP atomico", "SUBENSAMBLE_WIP")
        pt_one = _article("PT-BOM-ONE", "PT compartido uno", "PRODUCTO_TERMINADO")
        pt_two = _article("PT-BOM-TWO", "PT compartido dos", "PRODUCTO_TERMINADO")
        pt_missing = _article("PT-BOM-MISSING", "PT sin peso", "PRODUCTO_TERMINADO")
        pt_wip = _article("PT-BOM-WIP", "PT WIP atomico", "PRODUCTO_TERMINADO")
        color = db.session.get(PiezaColor, "PC-BOM-01")
        if color is None:
            color = PiezaColor(sku="PC-BOM-01", peso=2000.0)
            db.session.add(color)
            db.session.flush()
        else:
            color.peso = 2000.0
        piece_master = Pieza(codigo="PZ-BOM-01", nombre="Tapa BOM", peso_nominal_gr=2000.0)
        color_base = ColorBase(nombre="Rojo BOM")
        color_family = FamiliaColor(nombre="Solido BOM")
        db.session.add_all([piece_master, color_base, color_family])
        db.session.flush()
        production_color = ColorProduccion(
            color_base_id=color_base.id, familia_color_id=color_family.id,
            hex_referencia="#AA1122",
        )
        db.session.add(production_color)
        db.session.flush()
        color.pieza_id = piece_master.id
        color.color_produccion_id = production_color.id
        piece_link = db.session.query(ScmArticuloPiezaColor).filter_by(articulo_id=piece.id).one_or_none()
        if piece_link is None:
            db.session.add(ScmArticuloPiezaColor(articulo_id=piece.id, pieza_color_sku=color.sku))
        else:
            piece_link.pieza_color_sku = color.sku

        def bom(product, component, qty):
            revision = ScmEstructuraRevision(
                articulo_resultado_id=product.id, numero_revision=1,
                estado="APROBADA", content_hash="a" * 64, creada_por_id=actor_id,
            )
            revision.componentes.append(ScmEstructuraComponente(
                secuencia=1, articulo_componente_id=component.id, cantidad=qty,
            ))
            db.session.add(revision)

        bom(pt_one, piece, 3)
        bom(pt_two, piece, 1)
        bom(pt_missing, missing_weight, 1)
        bom(pt_wip, wip, 1)
        db.session.add(ScmSaldoInventarioKg(
            articulo_scm_id=piece.id, ubicacion_id=location.id,
            cantidad_fisica_kg=10, cantidad_reservada_kg=0,
            cantidad_no_disponible_kg=0, cantidad_retirada_kg=0,
            atributo_proceso="PROCESO",
        ))
        db.session.commit()

        payload = list_pt_availability(
            db.session, actor_id=actor_id, query="PT-BOM-ONE", location="PT-BOM-LOC",
        )
        assert len(payload["items"]) == 1
        item = payload["items"][0]
        assert item["potencial_un_estimado"] == "1.000"
        assert item["potencial_sumable"] is False
        assert item["componentes"][0]["es_limitante"] is True
        assert item["componentes"][0]["faltante_kg"] == "0.000"
        assert item["componentes"][0]["kg_requeridos_por_un_pt"] == "6.000"
        assert item["componentes"][0]["identidad_pieza"] == {
            "pieza_id": piece_master.id,
            "nombre": "Tapa BOM",
            "color_id": production_color.id,
            "color_nombre": production_color.nombre,
            "color_hex": "#AA1122",
        }

        all_items = list_pt_availability(db.session, actor_id=actor_id)
        by_code = {item["pt"]["codigo"]: item for item in all_items["items"]}
        assert by_code[pt_two.codigo]["componentes"][0]["grupo_stock_compartido"] == f"articulo:{piece.id}"
        assert by_code[pt_two.codigo]["potencial_sumable"] is False
        assert by_code[pt_missing.codigo]["potencial_un_estimado"] is None
        assert by_code[pt_missing.codigo]["potencial_motivo"] == "SIN_REFERENCIA_PESO"
        assert by_code[pt_wip.codigo]["componentes"][0]["naturaleza"] == "SUBENSAMBLE_WIP"
        assert len(by_code[pt_wip.codigo]["componentes"]) == 1
        assert by_code[pt_wip.codigo]["componentes"][0]["identidad_pieza"] is None


def test_pt_manual_requires_routine_capability_and_history_honors_location_scope(app):
    with app.app_context():
        actor = _actor_with_caps()
        routine = ScmCapacidad.query.filter_by(codigo="INVENTARIO_PT_MOVIMIENTO").one()
        actor.roles[0].capacidades.remove(routine)
        product = _article("PT-SCOPE-01", "PT scope", "PRODUCTO_TERMINADO")
        kg_piece = _article("PC-SCOPE-01", "Pieza fuera de clase", "PIEZA_COLOR")
        warehouse_allowed = ScmAlmacen(codigo="W2-ALLOWED", nombre="Permitido", tipo="PRODUCTO_TERMINADO")
        warehouse_hidden = ScmAlmacen(codigo="W2-HIDDEN", nombre="Oculto", tipo="PRODUCTO_TERMINADO")
        db.session.add_all([warehouse_allowed, warehouse_hidden])
        db.session.flush()
        allowed_location = ScmUbicacionInventario(codigo="W2-ALLOWED-LOC", nombre="Permitida", almacen_id=warehouse_allowed.id)
        hidden_location = ScmUbicacionInventario(codigo="W2-HIDDEN-LOC", nombre="Oculta", almacen_id=warehouse_hidden.id)
        db.session.add_all([allowed_location, hidden_location])
        db.session.add(ScmAlmacenTrabajador(
            almacen_id=warehouse_allowed.id, trabajador_id=actor.id,
            asignado_por_id=actor.id, clases_articulo_json=["PRODUCTO_TERMINADO"],
        ))
        db.session.flush()
        balance = ScmSaldoInventario(articulo_scm_id=product.id, ubicacion_id=hidden_location.id, cantidad_fisica=1)
        db.session.add_all([
            balance,
            ScmSaldoInventarioKg(
                articulo_scm_id=kg_piece.id, ubicacion_id=allowed_location.id,
                cantidad_fisica_kg=3, atributo_proceso="PROCESO",
            ),
        ])
        db.session.commit()
        with pytest.raises(ScmServiceError) as error:
            register_pt_manual_movement(
                db.session, actor_id=actor.id, operation_id=uuid4(), data={
                    "articulo_scm_id": product.id, "ubicacion_id": allowed_location.id,
                    "tipo": "ENTRADA", "cantidad": 1, "version": 1,
                    "fecha_operativa": "2026-09-19", "motivo": "Sin capacidad rutinaria", "referencia": "SCOPE-PT-01",
                },
            )
        assert error.value.code == "CAPABILITY_REQUIRED"
        db.session.rollback()
        with pytest.raises(ScmServiceError) as error:
            list_pt_manual_movements(db.session, actor_id=actor.id, balance_id=balance.id)
        assert error.value.code == "PT_MANUAL_BALANCE_NOT_FOUND"
        assert list_piece_kg_availability(db.session, actor_id=actor.id)["items"] == []


def test_piece_scope_filters_each_article_class_inside_same_warehouse(app):
    with app.app_context():
        actor = _actor_with_caps()
        warehouse = ScmAlmacen(codigo="W2-MIXED", nombre="Mixto", tipo="MATERIAS_PRIMAS")
        db.session.add(warehouse)
        db.session.flush()
        location = ScmUbicacionInventario(
            codigo="W2-MIXED-LOC", nombre="Mixto", almacen_id=warehouse.id,
        )
        piece = _article("PC-MIXED-01", "Pieza permitida", "PIEZA_COLOR")
        wip = _article("WIP-MIXED-01", "WIP no permitido", "SUBENSAMBLE_WIP")
        db.session.add(location)
        db.session.flush()
        db.session.add(ScmAlmacenTrabajador(
            almacen_id=warehouse.id, trabajador_id=actor.id,
            asignado_por_id=actor.id, clases_articulo_json=["PIEZA_COLOR"],
        ))
        db.session.add_all([
            ScmSaldoInventarioKg(
                articulo_scm_id=piece.id, ubicacion_id=location.id,
                cantidad_fisica_kg=4, atributo_proceso="PROCESO",
            ),
            ScmSaldoInventarioKg(
                articulo_scm_id=wip.id, ubicacion_id=location.id,
                cantidad_fisica_kg=9, atributo_proceso="PROCESO",
            ),
        ])
        db.session.commit()
        payload = list_piece_kg_availability(db.session, actor_id=actor.id)
        codes = {item["articulo"]["codigo"] for item in payload["items"]}
        assert codes == {"PC-MIXED-01"}


def test_piece_availability_includes_production_for_scoped_actor(app):
    with app.app_context():
        app.config["KG_PRODUCTION_LOCATION_CODE"] = "PRODUCCION_KG"
        actor = _actor_with_caps()
        scoped_warehouse = ScmAlmacen(
            codigo="W2-SCOPED", nombre="Almacen acotado", tipo="PIEZAS_WIP",
        )
        production_warehouse = ScmAlmacen(
            codigo="W2-PRODUCTION", nombre="Produccion", tipo="PIEZAS_WIP",
        )
        db.session.add_all([scoped_warehouse, production_warehouse])
        db.session.flush()
        scoped_location = ScmUbicacionInventario(
            codigo="W2-SCOPED-LOC", nombre="Almacen acotado",
            almacen_id=scoped_warehouse.id,
        )
        production_location = ScmUbicacionInventario(
            codigo="PRODUCCION_KG", nombre="Produccion KG",
            almacen_id=production_warehouse.id,
            tipo="PUNTO_PRODUCCION",
            permite_saldo_libre=True,
        )
        piece = _article("PC-PRODUCTION-SCOPE", "Pieza en produccion", "PIEZA_COLOR")
        db.session.add_all([
            scoped_location,
            production_location,
            ScmAlmacenTrabajador(
                almacen_id=scoped_warehouse.id,
                trabajador_id=actor.id,
                asignado_por_id=actor.id,
                clases_articulo_json=["PIEZA_COLOR"],
            ),
        ])
        db.session.flush()
        db.session.add(
            ScmSaldoInventarioKg(
                articulo_scm_id=piece.id,
                ubicacion_id=production_location.id,
                cantidad_fisica_kg=39.5,
                cantidad_reservada_kg=0,
                cantidad_no_disponible_kg=0,
                cantidad_retirada_kg=0,
                atributo_proceso="PROCESO",
            ),
        )
        db.session.commit()

        payload = list_piece_kg_availability(db.session, actor_id=actor.id)

        assert [item["articulo"]["codigo"] for item in payload["items"]] == [
            "PC-PRODUCTION-SCOPE"
        ]
        assert payload["items"][0]["ubicaciones"][0]["codigo"] == "PRODUCCION_KG"
        assert payload["items"][0]["kg_disponibles"] == "39.500"

        assert list_piece_kg_availability(
            db.session, actor_id=actor.id, location="W2-SCOPED-LOC"
        )["items"] == []
        assert list_piece_kg_availability(
            db.session, actor_id=actor.id, location="PRODUCCION_KG"
        )["items"][0]["kg_disponibles"] == "39.500"

        app.config["KG_PRODUCTION_LOCATION_CODE"] = ""
        assert list_piece_kg_availability(db.session, actor_id=actor.id)["items"] == []

        app.config["KG_PRODUCTION_LOCATION_CODE"] = "NO-EXISTE"
        assert list_piece_kg_availability(db.session, actor_id=actor.id)["items"] == []
        app.config["KG_PRODUCTION_LOCATION_CODE"] = "PRODUCCION_KG"

        production_location.activo = False
        db.session.commit()
        assert list_piece_kg_availability(db.session, actor_id=actor.id)["items"] == []
        production_location.activo = True
        production_location.tipo = "ALMACEN"
        db.session.commit()
        assert list_piece_kg_availability(db.session, actor_id=actor.id)["items"] == []
        production_location.tipo = "PUNTO_PRODUCCION"

        assignment = ScmAlmacenTrabajador.query.filter_by(
            almacen_id=scoped_warehouse.id,
            trabajador_id=actor.id,
        ).one()
        assignment.clases_articulo_json = ["SUBENSAMBLE_WIP"]
        db.session.commit()
        assert list_piece_kg_availability(db.session, actor_id=actor.id)["items"] == []


def test_pt_availability_export_has_normalized_and_matrix_bom_sheets(app, monkeypatch):
    with app.app_context():
        actor = _actor_with_caps()
        payload = {
            "items": [{
                "pt": {"id": 1, "codigo": "PT-01", "nombre": "=Producto prueba"},
                "revision_bom": {"numero": 2, "content_hash": "bom-sha"}, "saldo_manual_un": "3.000",
                "potencial_un_estimado": None, "potencial_estado": "NO_CALCULABLE",
                "potencial_motivo": "SIN_REFERENCIA_PESO",
                "componentes": [{
                    "articulo": {"id": 2, "codigo": "WIP-02", "nombre": "WIP armado"},
                    "identidad_pieza": None, "naturaleza": "SUBENSAMBLE_WIP",
                    "cantidad_bom_un": "1.000", "kg_disponibles": "4.500",
                    "peso_unitario_kg": None,
                    "kg_requeridos_por_un_pt": None, "faltante_kg": None,
                    "cobertura_un": None, "estado": "NO_CALCULABLE",
                    "es_limitante": False, "grupo_stock_compartido": "articulo:2",
                    "potencial_sumable": False,
                }],
                "potencial_sumable": False,
            }],
            "politica_piloto": "SIN_CONTROL_CALIDAD_DESDE_PESAJE",
        }
        import app.services.scm_kg_pt_availability_export as export_module
        monkeypatch.setattr(export_module, "list_pt_availability", lambda *args, **kwargs: payload)
        workbook = generate_pt_availability_xlsx(
            db.session, actor_id=actor.id, query="PT-01", location="PT-LOC"
        )
        book = load_workbook(workbook, data_only=False)
        assert book.sheetnames == ["Resumen PT", "Componentes BOM", "Matriz BOM", "Información"]
        assert book["Resumen PT"]["A2"].value == "PT-01"
        assert book["Resumen PT"]["B2"].value == "'=Producto prueba"
        assert book["Resumen PT"]["F2"].value is None
        assert book["Componentes BOM"]["I2"].value == "WIP armado"
        assert book["Componentes BOM"]["J2"].value is None
        assert book["Componentes BOM"]["O2"].number_format == "0.000"
        assert book["Matriz BOM"]["I2"].value == "WIP-02"
        assert book["Matriz BOM"].freeze_panes == "A2"
        assert book["Información"]["B2"].value
        assert book["Información"]["B5"].value == "PT-LOC"


@pytest.mark.parametrize("value", ["=SUM(A1:A2)", "+1", "-2", "@dato"])
def test_pt_availability_export_neutralizes_formula_like_text(value):
    assert _safe_text(value) == f"'{value}"


def test_pt_availability_export_rejects_excessive_matrix(app, monkeypatch):
    with app.app_context():
        actor = _actor_with_caps()
        component = {
            "articulo": {"id": 2, "codigo": "PC-02", "nombre": "Pieza"},
            "identidad_pieza": None, "naturaleza": "PIEZA_COLOR",
        }
        payload = {
            "items": [{
                "pt": {"id": 1, "codigo": "PT-01", "nombre": "Producto"},
                "componentes": [dict(component, articulo={**component["articulo"], "id": index})
                                for index in range(MAX_MATRIX_COMPONENTS + 1)],
            }],
        }
        import app.services.scm_kg_pt_availability_export as export_module
        monkeypatch.setattr(export_module, "list_pt_availability", lambda *args, **kwargs: payload)
        with pytest.raises(ScmServiceError) as error:
            generate_pt_availability_xlsx(db.session, actor_id=actor.id)
        assert error.value.code == "PT_EXPORT_TOO_LARGE"
        assert error.value.status_code == 413


def test_pt_availability_export_endpoint_returns_xlsx(app, monkeypatch):
    with app.app_context():
        actor = _actor_with_caps()
        import app.api.rutas_scm_kg_pt as routes
        from io import BytesIO

        monkeypatch.setattr(routes, "generate_pt_availability_xlsx", lambda *args, **kwargs: BytesIO(b"PK-test"))
        response = app.test_client().get(
            "/api/scm/v1/disponibilidad/productos-terminados/export.xlsx?q=PT-01&ubicacion=PT-LOC",
            headers={"X-Actor-Id": str(actor.id)},
        )
        assert response.status_code == 200
        assert response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert response.headers["Content-Disposition"].startswith("attachment;")
