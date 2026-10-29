from decimal import Decimal
from io import BytesIO
from datetime import date, datetime, timedelta, timezone
import hashlib
import json
from types import SimpleNamespace
from uuid import uuid4

import pytest

from openpyxl import load_workbook

from app.extensions import db
from app.models.scm_auditoria import ScmEvento
from app.models.scm_ot import (
    ScmAsignacionPersonalTrabajoOt,
    ScmAnulacionPesajeManga,
    ScmControlPesoManga,
    ScmCorreccionAsignacionManga,
    ScmAsignacionPlanMangaOt,
    ScmManga,
    ScmLoteArticulo,
    ScmPesajeManga,
    ScmTramoMangaTrabajo,
    ScmTrabajoColor,
    ScmTrabajoOt,
)
from app.models.scm_production_orders import (
    ScmCorridaFabricacion,
    ScmOrdenFabricacion,
    ScmOrdenOperacionSalida,
    ScmOrdenOperacion,
)
from app.models.trabajador import Trabajador
from app.services.scm_production_reports_service import (
    _filters,
    generate_production_history_xlsx,
    _group_history_rows,
    _history_hierarchy,
    _history_rows,
    list_production_history,
    MEASURE_OPTIONS,
    _project_corrected_kg_segments,
    _run_manga_values,
    list_production_progress_tv,
    _segment_conciliates,
    _valid_kg_segments,
    _history_weight_summary,
    _context_group_value,
    list_production_progress,
)
import app.services.scm_production_reports_service as production_reports_service


def test_tv_progress_groups_one_card_per_canonical_run_identity_and_keeps_piece_breakdown(monkeypatch):
    """TV-02 RED: visual labels must not collapse distinct production runs."""
    work_a = SimpleNamespace(id="work-a", codigo="TR-A")
    work_b = SimpleNamespace(id="work-b", codigo="TR-B")
    ot_a = SimpleNamespace(id=101, public_id="ot-a", codigo_ot="OT-101", fecha=date(2026, 10, 10), estado="EN_EJECUCION")
    ot_b = SimpleNamespace(id=102, public_id="ot-b", codigo_ot="OT-102", fecha=date(2026, 10, 10), estado="EN_EJECUCION")
    color_a = SimpleNamespace(id=7, codigo="AZUL", nombre="Azul", hex_referencia="#123456")
    color_b = SimpleNamespace(id=8, codigo="AZUL", nombre="Azul", hex_referencia="#123456")
    piece_a = SimpleNamespace(id=11, nombre="Pieza A")
    piece_b = SimpleNamespace(id=12, nombre="Pieza B")
    article_a = SimpleNamespace(id=201, codigo="PC-A", nombre="Salida A", pieza_color=SimpleNamespace(pieza_color=SimpleNamespace(id=301, sku="PC-A", pieza_rel=piece_a)))
    article_b = SimpleNamespace(id=202, codigo="PC-B", nombre="Salida B", pieza_color=SimpleNamespace(pieza_color=SimpleNamespace(id=302, sku="PC-B", pieza_rel=piece_b)))
    outputs = [
        SimpleNamespace(articulo=article_a, kg_estandar_objetivo=Decimal("2"), cantidad_objetivo=Decimal("10")),
        SimpleNamespace(articulo=article_b, kg_estandar_objetivo=Decimal("3"), cantidad_objetivo=Decimal("10")),
    ]
    runs = []
    for corrida_id, corrida_code, work, ot, color in (
        ("corrida-a", "C-A", work_a, ot_a, color_a),
        ("corrida-b", "C-B", work_b, ot_b, color_b),
    ):
        color_work = SimpleNamespace(
            molde_codigo_snapshot="ML-1", color_id_snapshot=color.id,
            color_nombre_snapshot=color.nombre,
        )
        corrida = SimpleNamespace(
            id=corrida_id, codigo=corrida_code, estado="EN_EJECUCION",
            objetivo_neto_kg=None, color_produccion=color, salidas=outputs,
        )
        runs.append({
            "orden": SimpleNamespace(id="of-1", codigo="OF-1", estado="EN_EJECUCION", molde_id="ML-1"),
            "corrida": corrida,
            "works": [(work, color_work, ot)],
            "contexts": [{"work": work, "color_work": color_work, "ot": ot}],
            "mangas": {}, "moldes": {"ML-1": SimpleNamespace(codigo="ML-1", nombre="Molde 1")},
            "color_name": color.nombre, "color_id": color.id, "color_code": color.codigo,
            "color_hex": color.hex_referencia, "molde": SimpleNamespace(codigo="ML-1", nombre="Molde 1"),
        })
    monkeypatch.setattr(production_reports_service, "load_actor", lambda *_a, **_k: SimpleNamespace(tiene_capacidad=lambda _c: True))
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: runs)

    result = list_production_progress_tv(object(), actor_id=7, filters={})

    assert len(result["items"]) == 2
    assert {item["group_id"] for item in result["items"]} == {
        "of-1|corrida-a|ML-1|7",
        "of-1|corrida-b|ML-1|8",
    }
    assert all(len(item["salidas"]) == 2 for item in result["items"])
    assert all(item["identity"]["molde_id"] == "ML-1" for item in result["items"])


def test_tv_progress_deduplicates_outputs_when_one_run_has_multiple_context_rows(monkeypatch):
    """TV-02 GREEN: repeated OT context rows cannot duplicate physical output."""
    work_a = SimpleNamespace(id="work-a")
    work_b = SimpleNamespace(id="work-b")
    ot = SimpleNamespace(id=101, public_id="ot-a", codigo_ot="OT-101")
    color_work_a = SimpleNamespace(molde_codigo_snapshot="ML-1", color_id_snapshot=7)
    color_work_b = SimpleNamespace(molde_codigo_snapshot="ML-1", color_id_snapshot=7)
    piece = SimpleNamespace(id=11, nombre="Pieza A")
    article = SimpleNamespace(
        id=201, codigo="PC-A", nombre="Salida A",
        pieza_color=SimpleNamespace(pieza_color=SimpleNamespace(id=301, sku="PC-A", pieza_rel=piece)),
    )
    output = SimpleNamespace(articulo=article, kg_estandar_objetivo=Decimal("2"), cantidad_objetivo=Decimal("10"))
    run = {
        "orden": SimpleNamespace(id="of-1", codigo="OF-1", estado="EN_EJECUCION", molde_id="ML-1"),
        "corrida": SimpleNamespace(
            id="corrida-a", codigo="C-A", estado="EN_EJECUCION", objetivo_neto_kg=None,
            color_produccion=SimpleNamespace(id=7, codigo="AZUL", nombre="Azul", hex_referencia="#123456"),
            salidas=[output],
        ),
        "works": [(work_a, color_work_a, ot), (work_b, color_work_b, ot)],
        "contexts": [
            {"work": work_a, "color_work": color_work_a, "ot": ot},
            {"work": work_b, "color_work": color_work_b, "ot": ot},
        ],
        "mangas": {}, "moldes": {"ML-1": SimpleNamespace(codigo="ML-1", nombre="Molde 1")},
        "color_name": "Azul", "color_id": 7, "color_code": "AZUL", "color_hex": "#123456",
        "molde": SimpleNamespace(codigo="ML-1", nombre="Molde 1"),
    }
    monkeypatch.setattr(production_reports_service, "load_actor", lambda *_a, **_k: SimpleNamespace(tiene_capacidad=lambda _c: True))
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: [run])

    result = list_production_progress_tv(object(), actor_id=7, filters={})

    assert len(result["items"]) == 1
    assert len(result["items"][0]["salidas"]) == 1


def test_tv_progress_marks_missing_canonical_identity_without_fake_percentage(monkeypatch):
    """TV-02 RED/GREEN: missing OT/mold/color identity stays explicit."""
    work = SimpleNamespace(id="work-a")
    piece = SimpleNamespace(id=11, nombre="Pieza A")
    article = SimpleNamespace(
        id=201, codigo="PC-A", nombre="Salida A",
        pieza_color=SimpleNamespace(pieza_color=SimpleNamespace(id=301, sku="PC-A", pieza_rel=piece)),
    )
    run = {
        "orden": SimpleNamespace(id="of-1", codigo="OF-1", estado="EN_EJECUCION", molde_id=None),
        "corrida": SimpleNamespace(
            id="corrida-a", codigo="C-A", estado="EN_EJECUCION", objetivo_neto_kg=None,
            color_produccion=SimpleNamespace(id=None, codigo="AZUL", nombre="Transparente", hex_referencia=None),
            salidas=[SimpleNamespace(articulo=article, kg_estandar_objetivo=Decimal("2"), cantidad_objetivo=Decimal("10"))],
        ),
        "works": [(work, None, None)],
        "contexts": [{"work": work, "color_work": SimpleNamespace(molde_codigo_snapshot=None, color_id_snapshot=None), "ot": None}],
        "mangas": {}, "moldes": {}, "color_name": "Transparente", "color_id": None, "color_hex": None,
        "molde": None,
    }
    monkeypatch.setattr(production_reports_service, "load_actor", lambda *_a, **_k: SimpleNamespace(tiene_capacidad=lambda _c: True))
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: [run])

    row = list_production_progress_tv(object(), actor_id=7, filters={})["items"][0]

    assert row["identity"]["identity_status"] == "INCOMPLETA"
    assert row["color_hex"] is None
    assert row["salidas"][0]["porcentaje"] is None
    assert row["salidas"][0]["estado_avance"] == "INCOMPLETO"


def test_tv_progress_redacts_mold_identity_without_of_permission(monkeypatch):
    monkeypatch.setenv("SECRET_KEY", "tv-02-server-secret-for-tests")
    work = SimpleNamespace(id="work-authority")
    ot = SimpleNamespace(id=101, public_id="ot-authority", codigo_ot="OT-101")
    color_work = SimpleNamespace(molde_codigo_snapshot="ML-SECRET", color_id_snapshot=7)
    run = {
        "orden": SimpleNamespace(id="of-authority", codigo="OF-AUTH", estado="EN_EJECUCION"),
        "corrida": SimpleNamespace(
            id="corrida-authority", codigo="C-AUTH", estado="EN_EJECUCION",
            color_produccion=SimpleNamespace(id=7, nombre="Azul", hex_referencia="#123456"),
            salidas=[],
        ),
        "works": [(work, color_work, ot)],
        "contexts": [{"work": work, "color_work": color_work, "ot": ot}],
        "moldes": {"ML-SECRET": SimpleNamespace(codigo="ML-SECRET", nombre="Molde secreto")},
        "color_name": "Azul", "color_id": 7, "color_hex": "#123456",
    }
    actor = SimpleNamespace(tiene_capacidad=lambda capability: capability != "OF_VER")
    monkeypatch.setattr(production_reports_service, "load_actor", lambda *_a, **_k: actor)
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: [run])

    hidden = list_production_progress_tv(object(), actor_id=7, filters={})["items"][0]

    assert "molde_visible" not in hidden
    assert hidden["group_id"].startswith("tvgrp-")
    assert "molde_id" not in hidden
    assert "molde" not in hidden
    assert "molde_id" not in hidden["identity"]
    assert "molde_nombre" not in hidden["identity"]
    assert all("MOLDE" not in issue for issue in hidden["identity"]["identity_issues"])
    assert "ML-SECRET" not in hidden["group_id"]
    simple_sha = hashlib.sha256(
        "of-authority|ot-authority|corrida-authority|ML-SECRET|7".encode("utf-8")
    ).hexdigest()[:24]
    assert hidden["group_id"] != f"tvgrp-{simple_sha}"
    hidden_again = list_production_progress_tv(object(), actor_id=7, filters={})["items"][0]
    assert hidden_again["group_id"] == hidden["group_id"]
    color_work.molde_codigo_snapshot = None
    hidden_incomplete = list_production_progress_tv(object(), actor_id=7, filters={})["items"][0]
    assert "molde_visible" not in hidden_incomplete
    assert hidden_incomplete["identity"]["identity_status"] == "INCOMPLETA"
    assert all("MOLDE" not in issue for issue in hidden_incomplete["identity"]["identity_issues"])
    public_without_mold = {
        key: ({inner_key: inner_value for inner_key, inner_value in value.items() if inner_key != "group_id"}
              if key == "identity" else value)
        for key, value in hidden_incomplete.items() if key != "group_id"
    }
    public_with_mold = {
        key: ({inner_key: inner_value for inner_key, inner_value in value.items() if inner_key != "group_id"}
              if key == "identity" else value)
        for key, value in hidden.items() if key != "group_id"
    }
    assert public_without_mold["identity"]["identity_status"] == "INCOMPLETA"
    assert public_with_mold["identity"]["identity_status"] == "COMPLETA"
    assert "molde" not in json.dumps(hidden_incomplete).lower()
    color_work.molde_codigo_snapshot = "ML-SECRET"

    monkeypatch.delenv("SECRET_KEY")
    with pytest.raises(RuntimeError, match="TV_GROUP_ID_SECRET_NOT_CONFIGURED"):
        list_production_progress_tv(object(), actor_id=7, filters={})

    monkeypatch.setattr(
        production_reports_service,
        "load_actor",
        lambda *_a, **_k: SimpleNamespace(tiene_capacidad=lambda _capability: True),
    )
    visible = list_production_progress_tv(object(), actor_id=7, filters={})["items"][0]
    assert visible["molde_visible"] is True
    assert visible["molde_id"] == "ML-SECRET"
    assert visible["molde"]["nombre"] == "Molde secreto"
    assert visible["identity"]["molde_id"] == "ML-SECRET"


def test_tv_progress_separates_incomplete_identity_by_work_source_marker(monkeypatch):
    work_a = SimpleNamespace(id="work-missing-a")
    work_b = SimpleNamespace(id="work-missing-b")
    ot = SimpleNamespace(id=101, public_id="ot-shared", codigo_ot="OT-SHARED")
    corrida = SimpleNamespace(
        id="corrida-shared", codigo="C-SHARED", estado="EN_EJECUCION",
        color_produccion=SimpleNamespace(id=None, nombre="Azul", hex_referencia=None),
        salidas=[],
    )
    run = {
        "orden": SimpleNamespace(id="of-shared", codigo="OF-SHARED", estado="EN_EJECUCION"),
        "corrida": corrida,
        "works": [(work_a, None, ot), (work_b, None, ot)],
        "contexts": [
            {"work": work_a, "color_work": SimpleNamespace(molde_codigo_snapshot=None, color_id_snapshot=None), "ot": ot},
            {"work": work_b, "color_work": SimpleNamespace(molde_codigo_snapshot=None, color_id_snapshot=None), "ot": ot},
        ],
        "moldes": {}, "color_name": "Azul", "color_id": None, "color_hex": None,
    }
    monkeypatch.setattr(
        production_reports_service,
        "load_actor",
        lambda *_a, **_k: SimpleNamespace(tiene_capacidad=lambda _capability: True),
    )
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: [run])

    items = list_production_progress_tv(object(), actor_id=7, filters={})["items"]

    assert len(items) == 2
    assert {item["group_id"] for item in items} == {
        "incompleto:of-shared|corrida-shared|?|?|work-missing-a",
        "incompleto:of-shared|corrida-shared|?|?|work-missing-b",
    }
    assert all(item["identity"]["identity_status"] == "INCOMPLETA" for item in items)


def test_tv_progress_keeps_unit_compatible_outputs_and_real_overproduction(monkeypatch):
    work = SimpleNamespace(id="work-tv", codigo="TR-TV")
    kg_article = SimpleNamespace(id=1, codigo="PC-KG", nombre="Pieza roja", unidad_inventario="KG", pieza_color=None)
    un_article = SimpleNamespace(id=2, codigo="PT-UN", nombre="Caja azul", unidad_inventario="UN", pieza_color=None)
    kg_output = SimpleNamespace(articulo=kg_article, kg_estandar_objetivo=Decimal("5"), cantidad_objetivo=Decimal("50"))
    un_output = SimpleNamespace(articulo=un_article, kg_estandar_objetivo=None, cantidad_objetivo=Decimal("10"))
    corrida = SimpleNamespace(id="corrida-tv", codigo="C-TV", objetivo_neto_kg=Decimal("30"), salidas=[kg_output, un_output])
    manga_kg = SimpleNamespace(
        id=1, estado="PESADA", trabajo=work, correccion_asignacion=None,
        lote_articulo=SimpleNamespace(articulo=kg_article), _report_segments=[],
        _report_final_kg=Decimal("6"), _report_open_kg=None,
        _report_quantity_un=Decimal("40"), _report_weight_corrected=False,
    )
    manga_un = SimpleNamespace(
        id=2, estado="PESADA", trabajo=work, correccion_asignacion=None,
        lote_articulo=SimpleNamespace(articulo=un_article), _report_segments=[],
        _report_final_kg=Decimal("8"), _report_open_kg=None,
        _report_quantity_un=Decimal("12"), _report_weight_corrected=False,
    )
    run = {
        "orden": SimpleNamespace(id="of-tv", codigo="OF-TV", estado="EN_EJECUCION"),
        "corrida": corrida, "works": [(work, None, None)], "mangas": {1: manga_kg, 2: manga_un},
        "color_name": "Rojo", "color_hex": "#AA0000",
    }
    monkeypatch.setattr(production_reports_service, "load_actor", lambda *_a, **_k: SimpleNamespace(tiene_capacidad=lambda _c: True))
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: [run])

    result = list_production_progress_tv(object(), actor_id=7, filters={})

    kg, units = result["items"][0]["salidas"]
    assert (kg["unidad"], kg["meta"], kg["tipo_meta"], kg["avance"], kg["pendiente"], kg["porcentaje"], kg["estado_avance"]) == (
        "KG", 5.0, "ESTANDAR", 6.0, 0, 120.0, "SOBRE_REFERENCIA",
    )
    assert units["unidad"] == "UN"
    assert (units["meta"], units["tipo_meta"], units["avance"], units["pendiente"], units["porcentaje"]) == (10.0, "OBJETIVO", 12.0, 0, 120.0)
    assert kg["peso_fisico_kg"] == 6.0
    assert units["peso_fisico_kg"] == 8.0
    assert kg["meta"] != 30.0  # no copy of the run-level kg net objective to a coproduct.


def test_tv_progress_does_not_project_ambiguous_segment_kg_to_each_run():
    from app.services.scm_production_reports_service import _progress_tv_actual

    work_a = SimpleNamespace(id="work-tv-a")
    work_b = SimpleNamespace(id="work-tv-b")
    article = SimpleNamespace(id=10)
    first, second = _segment(1, "0", "4", "4"), _segment(2, "4", "9", "5")
    first.trabajo, second.trabajo = work_a, work_b
    manga = SimpleNamespace(
        id=11, estado="PESADA", trabajo=work_a, correccion_asignacion=None,
        lote_articulo=SimpleNamespace(articulo=article), _report_segments=[first, second],
        _report_final_kg=Decimal("9"), _report_open_kg=None, _report_weight_corrected=False,
    )
    a = {"works": [(work_a, None, None)], "mangas": {11: manga}}
    b = {"works": [(work_b, None, None)], "mangas": {11: manga}}
    assert _progress_tv_actual(a, 10, "KG") == (Decimal("4"), False)
    assert _progress_tv_actual(b, 10, "KG") == (Decimal("5"), False)


def test_tv_progress_projects_net_correction_before_checking_kg_segments():
    from app.services.scm_production_reports_service import _progress_tv_actual

    work_a, work_b = SimpleNamespace(id="corr-a"), SimpleNamespace(id="corr-b")
    article = SimpleNamespace(id=12)
    first, second = _segment(1, "0", "4", "4"), _segment(2, "4", "9", "5")
    first.trabajo, second.trabajo = work_a, work_b
    manga = SimpleNamespace(
        id=12, estado="PESADA", trabajo=work_a, correccion_asignacion=None,
        lote_articulo=SimpleNamespace(articulo=article), _report_segments=[first, second],
        _report_final_kg=Decimal("10"), _report_open_kg=None, _report_weight_corrected=True,
    )
    a = {"works": [(work_a, None, None)], "mangas": {12: manga}}
    b = {"works": [(work_b, None, None)], "mangas": {12: manga}}
    assert _progress_tv_actual(a, 12, "KG") == (Decimal("4"), False)
    assert _progress_tv_actual(b, 12, "KG") == (Decimal("6"), False)


def test_tv_progress_rejects_broken_kg_ledger_and_only_falls_back_to_valid_un_ledger():
    from app.services.scm_production_reports_service import _progress_tv_actual

    work_a = SimpleNamespace(id="un-a")
    article = SimpleNamespace(id=13)
    broken, gap = _segment(1, "0", "4", "4"), _segment(2, "5", "9", "4")
    broken.trabajo, gap.trabajo = work_a, work_a
    broken.cantidad_inicio_un, broken.cantidad_fin_un, broken.cantidad_atribuida_un = Decimal("0"), Decimal("5"), Decimal("5")
    gap.cantidad_inicio_un, gap.cantidad_fin_un, gap.cantidad_atribuida_un = Decimal("6"), Decimal("10"), Decimal("4")
    manga_kg = SimpleNamespace(
        id=13, estado="PESADA", trabajo=work_a, correccion_asignacion=None,
        lote_articulo=SimpleNamespace(articulo=article), _report_segments=[broken, gap],
        _report_final_kg=Decimal("9"), _report_open_kg=None, _report_weight_corrected=False,
        _report_quantity_un=Decimal("10"),
    )
    run = {"works": [(work_a, None, None)], "mangas": {13: manga_kg}}
    assert _progress_tv_actual(run, 13, "KG") == (None, True)
    assert _progress_tv_actual(run, 13, "UN") == (None, True)

    gap.cantidad_inicio_un, gap.cantidad_fin_un, gap.cantidad_atribuida_un = Decimal("5"), Decimal("10"), Decimal("5")
    broken.cantidad_inicio_kg = gap.cantidad_inicio_kg = None
    broken.cantidad_fin_kg = gap.cantidad_fin_kg = None
    broken.cantidad_atribuida_kg = gap.cantidad_atribuida_kg = None
    broken.calidad_evidencia_kg = gap.calidad_evidencia_kg = None
    # With no KG ledger and one effective corrida owner, validated UN spans
    # prove ownership of the physical net exactly once.
    assert _progress_tv_actual(run, 13, "KG") == (Decimal("9"), False)


def test_tv_progress_exposes_real_physical_kg_when_no_meta_is_available(monkeypatch):
    work = SimpleNamespace(id="work-no-meta")
    article = SimpleNamespace(id=90, codigo="PC-NO-META", nombre="Sin objetivo", pieza_color=None)
    output = SimpleNamespace(articulo=article, kg_estandar_objetivo=None, cantidad_objetivo=None)
    corrida = SimpleNamespace(id="run-no-meta", codigo="C-NO-META", objetivo_neto_kg=None, estado="EN_EJECUCION", salidas=[output])
    manga = SimpleNamespace(
        id=90, estado="PESADA", trabajo=work, correccion_asignacion=None,
        lote_articulo=SimpleNamespace(articulo=article), _report_segments=[],
        _report_final_kg=Decimal("2.0"), _report_open_kg=None,
        _report_quantity_un=Decimal("0"), _report_weight_corrected=False,
    )
    run = {
        "orden": SimpleNamespace(id="of-no-meta", codigo="OF-NO-META", estado="EN_EJECUCION"),
        "corrida": corrida, "works": [(work, None, None)], "mangas": {90: manga},
        "color_name": "Verde", "color_hex": "#008800",
    }
    monkeypatch.setattr(production_reports_service, "load_actor", lambda *_a, **_k: SimpleNamespace(tiene_capacidad=lambda _c: True))
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: [run])
    item = list_production_progress_tv(object(), actor_id=7, filters={})["items"][0]["salidas"][0]
    assert item["unidad"] is None and item["meta"] is None
    assert item["estado_avance"] == "SIN_META"
    assert item["peso_fisico_kg"] == 2.0
    assert item["avance"] is None and item["porcentaje"] is None


def test_tv_progress_honors_historical_un_evidence_when_master_now_says_kg(monkeypatch):
    work = SimpleNamespace(id="work-of74", codigo="TR-74")
    # The current catalogue now says KG, but the approved output has no kg
    # target and the effective pesaje proves a corrected UN quantity.
    article = SimpleNamespace(id=74, codigo="PC-OF74", nombre="Pieza histórica", unidad_inventario="KG", pieza_color=None)
    output = SimpleNamespace(articulo=article, kg_estandar_objetivo=None, cantidad_objetivo=Decimal("20"))
    corrida = SimpleNamespace(id="corrida-of74", codigo="C-OF74", objetivo_neto_kg=None, estado="EN_EJECUCION", salidas=[output])
    manga = SimpleNamespace(
        id=74, estado="PESADA", trabajo=work, correccion_asignacion=None,
        lote_articulo=SimpleNamespace(articulo=article), _report_segments=[],
        _report_final_kg=Decimal("1.800"), _report_open_kg=None,
        _report_quantity_un=Decimal("18"), _report_weight_corrected=True,
    )
    run = {
        "orden": SimpleNamespace(id="of-74", codigo="OF-000074", estado="EN_EJECUCION"),
        "corrida": corrida, "works": [(work, None, None)], "mangas": {74: manga},
        "color_name": "Azul", "color_hex": "#0000AA",
    }
    monkeypatch.setattr(production_reports_service, "load_actor", lambda *_a, **_k: SimpleNamespace(tiene_capacidad=lambda _c: True))
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: [run])

    item = list_production_progress_tv(object(), actor_id=7, filters={})["items"][0]["salidas"][0]
    assert item["unidad"] == "UN"
    assert item["tipo_meta"] == "OBJETIVO"
    assert (item["meta"], item["avance"], item["pendiente"], item["porcentaje"]) == (20.0, 18.0, 2.0, 90.0)
    assert item["peso_fisico_kg"] == 1.8


def test_tv_progress_zero_goal_has_no_comparison_or_fake_zero_percent(monkeypatch):
    work = SimpleNamespace(id="work-zero")
    article = SimpleNamespace(id=80, codigo="PC-ZERO", nombre="Sin referencia", unidad_inventario="KG", pieza_color=None)
    output = SimpleNamespace(articulo=article, kg_estandar_objetivo=Decimal("0"), cantidad_objetivo=Decimal("10"))
    corrida = SimpleNamespace(id="run-zero", codigo="C-ZERO", objetivo_neto_kg=None, estado="EN_EJECUCION", salidas=[output])
    run = {
        "orden": SimpleNamespace(id="of-zero", codigo="OF-ZERO", estado="EN_EJECUCION"),
        "corrida": corrida, "works": [(work, None, None)], "mangas": {},
        "color_name": "Sin color", "color_hex": None,
    }
    monkeypatch.setattr(production_reports_service, "load_actor", lambda *_a, **_k: SimpleNamespace(tiene_capacidad=lambda _c: True))
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: [run])
    item = list_production_progress_tv(object(), actor_id=7, filters={})["items"][0]["salidas"][0]
    assert item["estado_avance"] == "SIN_META"
    assert item["meta"] is None
    assert item["porcentaje"] is None
    assert item["avance"] is None and item["pendiente"] is None
from app.services.scm_service_support import ScmServiceError


def _segment(sequence, start, end, attributed, quality="MEDIDA_DIRECTA", attributed_un=None):
    return SimpleNamespace(
        id=sequence,
        secuencia=sequence,
        estado="CERRADO",
        cantidad_inicio_kg=Decimal(start),
        cantidad_fin_kg=Decimal(end),
        cantidad_atribuida_kg=Decimal(attributed),
        cantidad_atribuida_un=attributed_un,
        cerrada_at=None,
        calidad_evidencia_kg=quality,
    )


def _tv_aggregate_run(*, corrida_id="run-aggregate", works=None, contexts=None,
                      mangas=None, color_id=7, color_name="Azul", output_meta="5"):
    works = works or [SimpleNamespace(id="work-a")]
    contexts = contexts or [
        {"work": works[0], "color_work": SimpleNamespace(
            molde_codigo_snapshot="ML-1", color_id_snapshot=color_id,
        ), "ot": SimpleNamespace(id=1, public_id="ot-a", codigo_ot="OT-A")}
    ]
    piece = SimpleNamespace(id=11, nombre="Pieza azul")
    article = SimpleNamespace(
        id=201, codigo="PC-A", nombre="Salida azul",
        pieza_color=SimpleNamespace(pieza_color=SimpleNamespace(
            id=301, sku="PC-A", pieza_rel=piece,
        )),
    )
    output = SimpleNamespace(
        id=401, articulo=article, kg_estandar_objetivo=Decimal(output_meta),
        cantidad_objetivo=Decimal("10"),
    )
    corrida = SimpleNamespace(
        id=corrida_id, codigo=f"C-{corrida_id}", estado="EN_EJECUCION",
        objetivo_neto_kg=None,
        color_produccion=SimpleNamespace(
            id=color_id, codigo=color_name.upper(), nombre=color_name,
            hex_referencia="#123456",
        ), salidas=[output],
    )
    return {
        "orden": SimpleNamespace(
            id="of-aggregate", codigo="OF-AGG", estado="EN_EJECUCION",
            molde_id="ML-1",
        ),
        "corrida": corrida,
        "works": [(context["work"], context.get("color_work"), context.get("ot")) for context in contexts],
        "contexts": contexts,
        "mangas": mangas or {},
        "moldes": {"ML-1": SimpleNamespace(codigo="ML-1", nombre="Molde 1")},
        "color_name": color_name, "color_id": color_id,
        "color_code": color_name.upper(), "color_hex": "#123456",
        "molde": SimpleNamespace(codigo="ML-1", nombre="Molde 1"),
    }


def _tv_manga(manga_id, work, kg, *, segments=None):
    article = SimpleNamespace(id=201)
    return SimpleNamespace(
        id=manga_id, estado="PESADA", trabajo=work,
        correccion_asignacion=None,
        lote_articulo=SimpleNamespace(articulo=article),
        _report_segments=segments or [], _report_final_kg=Decimal(str(kg)),
        _report_open_kg=None, _report_quantity_un=None,
        _report_weight_corrected=False,
    )


def test_tv_progress_excludes_released_without_work_but_keeps_legacy_active(monkeypatch):
    runs = []
    for state in ("LIBERADA", "EN_EJECUCION", "PAUSADA", "CANCELADA"):
        run = _tv_aggregate_run(corrida_id=f"run-{state}")
        run["orden"] = SimpleNamespace(
            id=f"of-{state}", codigo=f"OF-{state}", estado=state, molde_id="ML-1",
        )
        # Explicitly clear after the helper: its `works or [...]` default
        # intentionally supplies work for other report scenarios.
        if state in {"LIBERADA", "EN_EJECUCION"}:
            run["works"] = []
            run["contexts"] = []
        runs.append(run)
    monkeypatch.setattr(
        production_reports_service, "load_actor",
        lambda *_a, **_k: SimpleNamespace(tiene_capacidad=lambda _c: True),
    )
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: runs)

    items = list_production_progress_tv(object(), actor_id=7, filters={})["items"]

    assert [item["estado"] for item in items] == ["EN_EJECUCION"]
    assert items[0]["of_id"] == "of-EN_EJECUCION"


def test_tv_progress_consolidates_two_ots_into_one_output_target(monkeypatch):
    work_a, work_b = SimpleNamespace(id="work-a"), SimpleNamespace(id="work-b")
    contexts = [
        {"work": work_a, "color_work": SimpleNamespace(molde_codigo_snapshot="ML-1", color_id_snapshot=7),
         "ot": SimpleNamespace(id=1, public_id="ot-a", codigo_ot="OT-A")},
        {"work": work_b, "color_work": SimpleNamespace(molde_codigo_snapshot="ML-1", color_id_snapshot=7),
         "ot": SimpleNamespace(id=2, public_id="ot-b", codigo_ot="OT-B")},
    ]
    mangas = {1: _tv_manga(1, work_a, "2"), 2: _tv_manga(2, work_b, "3")}
    run = _tv_aggregate_run(works=[work_a, work_b], contexts=contexts, mangas=mangas)
    monkeypatch.setattr(production_reports_service, "load_actor", lambda *_a, **_k: SimpleNamespace(tiene_capacidad=lambda _c: True))
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: [run])

    item = list_production_progress_tv(object(), actor_id=7, filters={})["items"]

    assert len(item) == 1
    output = item[0]["salidas"]
    assert len(output) == 1
    assert (output[0]["meta"], output[0]["avance"], output[0]["porcentaje"]) == (5.0, 5.0, 100.0)
    assert item[0]["ot_ids"] == ["ot-a", "ot-b"]
    assert {entry["codigo"] for entry in item[0]["ots"]} == {"OT-A", "OT-B"}


def test_tv_progress_does_not_double_count_manga_shared_by_contexts(monkeypatch):
    work_a, work_b = SimpleNamespace(id="work-a"), SimpleNamespace(id="work-b")
    contexts = [
        {"work": work_a, "color_work": SimpleNamespace(molde_codigo_snapshot="ML-1", color_id_snapshot=7),
         "ot": SimpleNamespace(id=1, public_id="ot-a", codigo_ot="OT-A")},
        {"work": work_b, "color_work": SimpleNamespace(molde_codigo_snapshot="ML-1", color_id_snapshot=7),
         "ot": SimpleNamespace(id=2, public_id="ot-b", codigo_ot="OT-B")},
    ]
    # The same physical manga is present once in the normalized run and is
    # owned by only work A, even though the run has two OT contexts.
    run = _tv_aggregate_run(works=[work_a, work_b], contexts=contexts,
                            mangas={1: _tv_manga(1, work_a, "2")})
    monkeypatch.setattr(production_reports_service, "load_actor", lambda *_a, **_k: SimpleNamespace(tiene_capacidad=lambda _c: True))
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: [run])

    output = list_production_progress_tv(object(), actor_id=7, filters={})["items"][0]["salidas"][0]

    assert output["avance"] == 2.0
    assert output["porcentaje"] == 40.0


def test_tv_progress_keeps_different_runs_separate_even_with_same_color_name(monkeypatch):
    runs = [
        _tv_aggregate_run(corrida_id="run-a", color_name="Azul"),
        _tv_aggregate_run(corrida_id="run-b", color_name="Azul"),
    ]
    monkeypatch.setattr(production_reports_service, "load_actor", lambda *_a, **_k: SimpleNamespace(tiene_capacidad=lambda _c: True))
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: runs)

    items = list_production_progress_tv(object(), actor_id=7, filters={})["items"]

    assert len(items) == 2
    assert {item["corrida_id"] for item in items} == {"run-a", "run-b"}


def test_tv_progress_allows_missing_ot_when_core_identity_and_evidence_are_attributable(monkeypatch):
    work = SimpleNamespace(id="work-no-ot")
    context = {"work": work, "color_work": SimpleNamespace(
        molde_codigo_snapshot="ML-1", color_id_snapshot=7,
    ), "ot": None}
    run = _tv_aggregate_run(works=[work], contexts=[context],
                            mangas={1: _tv_manga(1, work, "5")})
    monkeypatch.setattr(production_reports_service, "load_actor", lambda *_a, **_k: SimpleNamespace(tiene_capacidad=lambda _c: True))
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: [run])

    item = list_production_progress_tv(object(), actor_id=7, filters={})["items"][0]
    output = item["salidas"][0]

    assert output["porcentaje"] == 100.0
    assert "OT_ID_AUSENTE" in item["provenance"]["issues"]
    assert item["identity"]["identity_status"] == "COMPLETA"


def test_tv_progress_keeps_percentage_null_when_kg_is_not_attributable(monkeypatch):
    work = SimpleNamespace(id="work-owned")
    foreign_work = SimpleNamespace(id="work-foreign")
    segment = _segment(1, "0", "5", "5")
    segment.trabajo = foreign_work
    run = _tv_aggregate_run(works=[work], mangas={
        1: _tv_manga(1, work, "5", segments=[segment]),
    })
    monkeypatch.setattr(production_reports_service, "load_actor", lambda *_a, **_k: SimpleNamespace(tiene_capacidad=lambda _c: True))
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: [run])

    output = list_production_progress_tv(object(), actor_id=7, filters={})["items"][0]["salidas"][0]

    assert output["avance"] is None
    assert output["porcentaje"] is None
    assert output["estado_avance"] == "SIN_PESAJES"


def test_tv_progress_preserves_meta_but_hides_weighted_progress_without_permission(monkeypatch):
    work = SimpleNamespace(id="work-restricted")
    run = _tv_aggregate_run(works=[work], mangas={1: _tv_manga(1, work, "5")})
    actor = SimpleNamespace(tiene_capacidad=lambda capability: capability != "MANGA_PESAJE_VER")
    monkeypatch.setattr(production_reports_service, "load_actor", lambda *_a, **_k: actor)
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: [run])

    output = list_production_progress_tv(object(), actor_id=7, filters={})["items"][0]["salidas"][0]

    assert output["meta"] == 5.0
    assert output["avance"] is None
    assert output["porcentaje"] is None
    assert output["estado_avance"] == "SIN_PERMISO_PESAJE"


def test_tv_actual_ignores_known_foreign_mass_but_marks_own_plus_ownerless_incomplete():
    work = SimpleNamespace(id="work-owned")
    foreign = SimpleNamespace(id="work-foreign")
    ownerless = _segment(2, "3", "5", "2")
    ownerless.trabajo = None
    own = _segment(1, "0", "3", "3")
    own.trabajo = work
    foreign_only = _tv_manga(1, work, "5", segments=[_segment(1, "0", "5", "5")])
    foreign_only._report_segments[0].trabajo = foreign
    own_plus_unknown = _tv_manga(2, work, "5", segments=[own, ownerless])
    run = {"works": [(work, None, None)], "mangas": {1: foreign_only, 2: own_plus_unknown}}

    from app.services.scm_production_reports_service import _progress_tv_actual

    assert _progress_tv_actual({"works": run["works"], "mangas": {1: foreign_only}}, 201, "KG") == (None, False)
    # The attributable 3 kg remains visible while the ownerless remainder
    # keeps the result explicitly partial.
    assert _progress_tv_actual({"works": run["works"], "mangas": {2: own_plus_unknown}}, 201, "KG") == (Decimal("3"), True)


def test_tv_actual_marks_no_segments_ownerless_physical_mass_incomplete():
    work = SimpleNamespace(id="work-owned")
    ownerless = _tv_manga(1, None, "2")
    ownerless.estado = "EN_LLENADO"
    ownerless._report_final_kg = None
    ownerless._report_open_kg = Decimal("2")
    from app.services.scm_production_reports_service import _progress_tv_actual

    assert _progress_tv_actual({"works": [(work, None, None)], "mangas": {1: ownerless}}, 201, "KG") == (None, True)
    breakdown = production_reports_service._progress_tv_breakdown(
        {"works": [(work, None, None)], "mangas": {1: ownerless}}, 201, "KG"
    )
    assert breakdown["fisico_observado"] == Decimal("2")
    assert breakdown["medido"] is None
    assert breakdown["incomplete"] is True


def test_tv_actual_adds_latest_open_control_once_and_exposes_breakdown():
    work = SimpleNamespace(id="work-open")
    article = SimpleNamespace(id=201)
    final = SimpleNamespace(
        id=1, estado="PESADA", trabajo=work, correccion_asignacion=None,
        lote_articulo=SimpleNamespace(articulo=article), _report_segments=[],
        _report_final_kg=Decimal("90.6"), _report_open_kg=None,
        _report_quantity_un=None, _report_weight_corrected=False,
    )
    opened = SimpleNamespace(
        id=2, estado="EN_LLENADO", trabajo=work, correccion_asignacion=None,
        lote_articulo=SimpleNamespace(articulo=article), _report_segments=[],
        _report_final_kg=None, _report_open_kg=Decimal("1.1"),
        _report_quantity_un=None, _report_weight_corrected=False,
    )
    pending = SimpleNamespace(
        id=3, estado="EN_LLENADO", trabajo=work, correccion_asignacion=None,
        lote_articulo=SimpleNamespace(articulo=article), _report_segments=[],
        _report_final_kg=None, _report_open_kg=None, _report_quantity_un=None,
        _report_weight_corrected=False,
    )
    run = {"works": [(work, None, None)], "mangas": {1: final, 2: opened, 3: pending}}

    breakdown = production_reports_service._progress_tv_breakdown(run, 201, "KG")

    assert breakdown == {
        "finalizados": Decimal("90.6"),
        "en_proceso": Decimal("1.1"),
        "medido": Decimal("91.7"),
        "fisico_observado": Decimal("91.7"),
        "mangas_total": 3,
        "mangas_conocidas": 2,
        "pendientes_sin_peso": 1,
        "pendientes_sin_cantidad_un": 0,
        "incomplete": True,
        "motivos_atribucion": [],
    }


def test_tv_final_replaces_open_control_even_with_multiple_control_records():
    work = SimpleNamespace(id="work-replacement-tv")
    article = SimpleNamespace(id=201)
    manga = SimpleNamespace(
        id=1, estado="PESADA", trabajo=work, correccion_asignacion=None,
        lote_articulo=SimpleNamespace(articulo=article), _report_segments=[],
        _report_final_kg=Decimal("90.6"), _report_open_kg=Decimal("1.1"),
        _report_quantity_un=None, _report_weight_corrected=False,
        _report_controls=[SimpleNamespace(estado="ANULADO"), SimpleNamespace(estado="CERRADO")],
        _report_reopened=True,
    )

    breakdown = production_reports_service._progress_tv_breakdown(
        {"works": [(work, None, None)], "mangas": {1: manga}}, 201, "KG"
    )

    assert breakdown["finalizados"] == Decimal("90.6")
    assert breakdown["en_proceso"] == Decimal("0")
    assert breakdown["medido"] == Decimal("90.6")


def test_tv_breakdown_excludes_foreign_known_and_pending_mangas_from_scope_counts():
    work = SimpleNamespace(id="work-owned")
    foreign = SimpleNamespace(id="work-foreign")
    owned = _tv_manga(1, work, "2")
    foreign_known = _tv_manga(2, foreign, "9")
    foreign_pending = _tv_manga(3, foreign, "2")
    foreign_pending._report_final_kg = None
    foreign_pending._report_open_kg = None

    breakdown = production_reports_service._progress_tv_breakdown(
        {"works": [(work, None, None)], "mangas": {
            1: owned, 2: foreign_known, 3: foreign_pending,
        }},
        201,
        "KG",
    )

    assert breakdown["medido"] == Decimal("2")
    assert breakdown["fisico_observado"] == Decimal("2")
    assert breakdown["mangas_total"] == 1
    assert breakdown["mangas_conocidas"] == 1
    assert breakdown["pendientes_sin_peso"] == 0


def test_standard_and_tv_share_incompatible_mixed_ledger_projection(monkeypatch):
    work = SimpleNamespace(id="work-owned")
    foreign = SimpleNamespace(id="work-foreign")
    own = _segment(1, "0", "4", "4")
    own.trabajo = work
    foreign_segment = _segment(2, "4", "9", "5")
    foreign_segment.trabajo = foreign
    ownerless = _segment(3, "9", "10", "1")
    ownerless.trabajo = None
    manga = SimpleNamespace(
        id=88, estado="PESADA", trabajo=work, correccion_asignacion=None,
        _report_final_kg=Decimal("10"), _report_open_kg=None,
        _report_segments=[own, foreign_segment, ownerless],
        _report_all_segments=[own, foreign_segment, ownerless],
        _report_weight_corrected=False, _report_quantity_un=None,
        lote_articulo=SimpleNamespace(articulo=SimpleNamespace(id=201)),
    )
    standard_run = {
        "corrida": SimpleNamespace(objetivo_neto_kg=Decimal("10")),
        "mangas": {88: manga}, "contexts": [{"work": work}],
    }
    standard = _run_manga_values(standard_run)
    tv = production_reports_service._progress_tv_breakdown(
        {"works": [(work, None, None)], "mangas": {88: manga}}, 201, "KG"
    )

    assert standard[:5] == (Decimal("4"), Decimal("0"), Decimal("4"), 1, 0)
    standard_run["corrida"] = SimpleNamespace(
        id="run-mixed", codigo="C-MIXED", objetivo_neto_kg=Decimal("10"),
        color_produccion=None, salidas=[],
    )
    standard_run.update({
        "orden": SimpleNamespace(id="of-mixed", codigo="OF-MIXED", estado="EN_EJECUCION"),
        "ot": None, "color_name": "Azul", "color_hex": None, "molde": None,
    })
    monkeypatch.setattr(
        production_reports_service, "load_actor",
        lambda *_a, **_k: SimpleNamespace(tiene_capacidad=lambda _c: True),
    )
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: [standard_run])
    item = list_production_progress(object(), actor_id=1, filters={})["items"][0]
    assert item["kg_medidos_efectivos"] == 4.0
    assert item["coverage"]["estado"] == "INCOMPLETA"
    assert tv["finalizados"] == Decimal("4")
    assert tv["medido"] == Decimal("4")
    assert tv["incomplete"] is True


@pytest.mark.parametrize("shape, expected_physical", [("incompatible", 9.0), ("active", 5.0)])
def test_standard_and_tv_keep_own_span_scope_when_effective_owner_moves_foreign(
    monkeypatch, shape, expected_physical,
):
    """A prior own span keeps the manga visible after a foreign handoff.

    The current foreign owner must not turn the current final/control into
    attributable KG.  Both serializers retain the physical diagnostic and
    incomplete coverage, while a foreign-only ledger remains out of scope.
    """
    work = SimpleNamespace(id="work-handoff-own")
    foreign = SimpleNamespace(id="work-handoff-foreign")
    own = _segment(1, "0", "4", "4")
    own.trabajo = work
    if shape == "incompatible":
        foreign_segment = _segment(2, "4", "9", "4", quality="PENDIENTE")
        manga_state = "PESADA"
        final_kg, open_kg = Decimal("9"), None
    else:
        foreign_segment = _segment(2, "4", "4", "0")
        foreign_segment.estado = "ACTIVO"
        foreign_segment.cantidad_fin_kg = None
        foreign_segment.cantidad_atribuida_kg = None
        manga_state = "EN_LLENADO"
        final_kg, open_kg = None, Decimal("5")
    foreign_segment.trabajo = foreign
    manga = SimpleNamespace(
        id=91, estado=manga_state, trabajo=foreign, correccion_asignacion=None,
        _report_final_kg=final_kg, _report_open_kg=open_kg,
        _report_segments=[], _report_all_segments=[own, foreign_segment],
        _report_weight_corrected=False, _report_quantity_un=None,
        lote_articulo=SimpleNamespace(articulo=SimpleNamespace(id=201)),
    )
    run = _tv_aggregate_run(works=[work], mangas={91: manga}, output_meta="10")
    run["corrida"].objetivo_neto_kg = Decimal("10")
    run["ot"] = None
    run["ots"] = []

    monkeypatch.setattr(
        production_reports_service, "load_actor",
        lambda *_a, **_k: SimpleNamespace(tiene_capacidad=lambda _c: True),
    )
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: [run])

    standard = list_production_progress(object(), actor_id=1, filters={})["items"][0]
    assert standard["kg_finalizados_efectivos"] is None
    assert standard["kg_medidos_en_abiertas"] is None
    assert standard["kg_medidos_efectivos"] is None
    assert standard["avance_medido_kg"] is None
    assert standard["peso_fisico_kg"] == expected_physical
    assert standard["mangas"] == {"total": 1, "conocidas": 0}
    assert standard["coverage"]["estado"] == "INCOMPLETA"
    assert standard["motivos_atribucion"] == ["ATRIBUCION_INCOMPLETA"]

    tv = list_production_progress_tv(object(), actor_id=1, filters={})["items"][0]["salidas"][0]
    assert tv["avance"] is None
    assert tv["finalizados_kg"] is None
    assert tv["en_proceso_kg"] is None
    assert tv["avance_medido_kg"] is None
    assert tv["peso_fisico_kg"] == expected_physical
    assert tv["mangas_total"] == 1
    assert tv["estado_avance"] == "INCOMPLETO"
    assert "ATRIBUCION_INCOMPLETA" in tv["motivos_atribucion"]


def test_standard_and_tv_keep_ownerless_ot_with_segments_in_scope(monkeypatch):
    work = SimpleNamespace(id="work-ownerless-segments")
    ot = SimpleNamespace(id="ot-ownerless-segments", public_id="ot-ownerless-segments", codigo_ot="OT-OWNERLESS")
    ownerless_segment = _segment(1, "0", "5", "5")
    ownerless_segment.trabajo = None
    manga = SimpleNamespace(
        id=92, estado="PESADA", trabajo=None, ot_id=ot.id,
        correccion_asignacion=None, _report_final_kg=Decimal("5"), _report_open_kg=None,
        _report_segments=[], _report_all_segments=[ownerless_segment],
        _report_weight_corrected=False, _report_quantity_un=None,
        lote_articulo=SimpleNamespace(articulo=SimpleNamespace(id=201)),
    )
    run = _tv_aggregate_run(
        works=[work],
        contexts=[{"work": work, "color_work": None, "ot": ot}],
        mangas={92: manga}, output_meta="10",
    )
    run["corrida"].objetivo_neto_kg = Decimal("10")
    run["ot"] = ot
    run["ots"] = [ot]
    monkeypatch.setattr(
        production_reports_service, "load_actor",
        lambda *_a, **_k: SimpleNamespace(tiene_capacidad=lambda _c: True),
    )
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: [run])

    standard = list_production_progress(object(), actor_id=1, filters={})["items"][0]
    assert standard["mangas"] == {"total": 1, "conocidas": 0}
    assert standard["peso_fisico_kg"] == 5.0
    assert standard["kg_medidos_efectivos"] is None
    assert standard["coverage"]["estado"] == "INCOMPLETA"
    assert standard["motivos_atribucion"] == ["ATRIBUCION_INCOMPLETA"]

    tv = list_production_progress_tv(object(), actor_id=1, filters={})["items"][0]["salidas"][0]
    assert tv["mangas_total"] == 1
    assert tv["peso_fisico_kg"] == 5.0
    assert tv["avance"] is None
    assert tv["estado_avance"] == "INCOMPLETO"
    assert "ATRIBUCION_INCOMPLETA" in tv["motivos_atribucion"]


def test_standard_valid_own_foreign_split_has_physical_without_false_attribution_reason(monkeypatch):
    work = SimpleNamespace(id="work-split-own")
    foreign = SimpleNamespace(id="work-split-foreign")
    own = _segment(1, "0", "4", "4")
    own.trabajo = work
    foreign_segment = _segment(2, "4", "9", "5")
    foreign_segment.trabajo = foreign
    manga = SimpleNamespace(
        id=87, estado="PESADA", trabajo=work, correccion_asignacion=None,
        _report_final_kg=Decimal("9"), _report_open_kg=None,
        _report_segments=[own, foreign_segment], _report_all_segments=[own, foreign_segment],
        _report_weight_corrected=False, _report_quantity_un=None,
    )
    run = {
        "corrida": SimpleNamespace(
            id="run-split", codigo="C-SPLIT", objetivo_neto_kg=Decimal("10"),
            color_produccion=None, salidas=[],
        ),
        "orden": SimpleNamespace(id="of-split", codigo="OF-SPLIT", estado="EN_EJECUCION"),
        "ot": None, "color_name": "Melón", "color_hex": None, "molde": None,
        "mangas": {87: manga}, "contexts": [{"work": work}],
    }
    monkeypatch.setattr(production_reports_service, "load_actor", lambda *_a, **_k: SimpleNamespace(tiene_capacidad=lambda _c: True))
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: [run])

    item = list_production_progress(object(), actor_id=1, filters={})["items"][0]

    assert item["kg_medidos_efectivos"] == 4.0
    assert item["peso_fisico_kg"] == 9.0
    assert item["mangas"] == {"total": 1, "conocidas": 1}
    assert item["coverage"]["estado"] == "COMPLETA"
    assert item["motivos_atribucion"] == []


def test_standard_foreign_only_has_no_physical_or_attribution_reason(monkeypatch):
    work = SimpleNamespace(id="work-foreign-only")
    foreign = SimpleNamespace(id="work-foreign-source")
    foreign_segment = _segment(1, "0", "9", "9")
    # A current foreign owner plus an ownerless segment is still outside this
    # corrida; unknown ownership must not leak the OT's physical fact here.
    foreign_segment.trabajo = None
    manga = SimpleNamespace(
        id=86, estado="PESADA", trabajo=foreign, correccion_asignacion=None,
        _report_final_kg=Decimal("9"), _report_open_kg=None,
        _report_segments=[foreign_segment], _report_all_segments=[foreign_segment],
        _report_weight_corrected=False, _report_quantity_un=None,
    )
    run = {
        "corrida": SimpleNamespace(
            id="run-foreign-only", codigo="C-FOREIGN-ONLY", objetivo_neto_kg=Decimal("10"),
            color_produccion=None, salidas=[],
        ),
        "orden": SimpleNamespace(id="of-foreign-only", codigo="OF-FOREIGN-ONLY", estado="EN_EJECUCION"),
        "ot": None, "color_name": "Rojo", "color_hex": None, "molde": None,
        "mangas": {86: manga}, "contexts": [{"work": work}],
    }
    result = _run_manga_values(run)

    assert result[:5] == (Decimal("0"), Decimal("0"), None, 0, 0)
    assert run["_report_physical_kg"] is None
    assert run["_report_attribution_reasons"] == []


def test_tv_serializer_uses_un_quantity_counters_and_keeps_kg_diagnostic_separate(monkeypatch):
    work = SimpleNamespace(id="work-un-serializer")
    foreign = SimpleNamespace(id="work-un-serializer-foreign")
    ot = SimpleNamespace(id="ot-un-serializer", public_id="ot-un-serializer", codigo_ot="OT-UN-SERIALIZER")
    own = _tv_manga(1, work, "0")
    own._report_final_kg = None
    own._report_quantity_un = Decimal("4")
    foreign_manga = _tv_manga(2, foreign, "0")
    foreign_manga._report_final_kg = None
    foreign_manga._report_quantity_un = Decimal("9")
    ownerless = _tv_manga(3, None, "0")
    ownerless._report_final_kg = None
    ownerless._report_quantity_un = Decimal("7")
    ownerless.ot_id = ot.id
    run = _tv_aggregate_run(
        works=[work],
        contexts=[{"work": work, "color_work": SimpleNamespace(molde_codigo_snapshot="ML-1", color_id_snapshot=7), "ot": ot}],
        mangas={1: own, 2: foreign_manga, 3: ownerless}, output_meta="0",
    )
    output = run["corrida"].salidas[0]
    output.kg_estandar_objetivo = None
    output.cantidad_objetivo = Decimal("10")
    run["corrida"].objetivo_neto_kg = None
    monkeypatch.setattr(production_reports_service, "load_actor", lambda *_a, **_k: SimpleNamespace(tiene_capacidad=lambda _c: True))
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: [run])

    serialized = list_production_progress_tv(object(), actor_id=1, filters={})["items"][0]["salidas"][0]

    assert serialized["unidad"] == "UN"
    assert serialized["avance"] == 4.0
    assert serialized["mangas_total"] == 2
    assert serialized["mangas_conocidas"] == 2
    assert serialized["pendientes_sin_peso"] == 0
    assert serialized["pendientes_sin_cantidad_un"] == 0
    assert serialized["estado_avance"] == "INCOMPLETO"


def test_standard_and_tv_do_not_fallback_incompatible_closed_ledger_to_whole_manga(monkeypatch):
    work = SimpleNamespace(id="work-invalid")
    invalid = _segment(1, "0", "5", "4", quality="PENDIENTE")
    invalid.trabajo = work
    manga = SimpleNamespace(
        id=89, estado="PESADA", trabajo=work, correccion_asignacion=None,
        _report_final_kg=Decimal("5"), _report_open_kg=None,
        _report_segments=[], _report_all_segments=[invalid],
        _report_weight_corrected=False, _report_quantity_un=None,
        lote_articulo=SimpleNamespace(articulo=SimpleNamespace(id=201)),
    )

    standard = _run_manga_values({
        "corrida": SimpleNamespace(objetivo_neto_kg=Decimal("5")),
        "mangas": {89: manga}, "contexts": [{"work": work}],
    })
    tv = production_reports_service._progress_tv_breakdown(
        {"works": [(work, None, None)], "mangas": {89: manga}}, 201, "KG"
    )

    assert standard[:5] == (Decimal("0"), Decimal("0"), None, 1, 0)
    assert tv["medido"] is None
    assert tv["incomplete"] is True

    monkeypatch.setattr(
        production_reports_service,
        "load_actor",
        lambda *_a, **_k: SimpleNamespace(tiene_capacidad=lambda _capability: True),
    )
    standard_run = {
        "corrida": SimpleNamespace(
            id="run-invalid", codigo="C-INVALID", objetivo_neto_kg=Decimal("5"),
            color_produccion=None, salidas=[],
        ),
        "orden": SimpleNamespace(id="of-invalid", codigo="OF-INVALID", estado="EN_EJECUCION"),
        "ot": None, "color_name": "Azul", "color_hex": None, "molde": None,
        "mangas": {89: manga}, "contexts": [{"work": work}],
    }
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: [standard_run])
    item = list_production_progress(object(), actor_id=1, filters={})["items"][0]
    assert item["kg_finalizados_efectivos"] is None
    assert item["kg_medidos_en_abiertas"] is None
    assert item["kg_medidos_efectivos"] is None
    assert item["subtotal_conocido_kg"] is None
    assert item["peso_fisico_kg"] == 5.0
    assert item["coverage"]["motivos"] == ["ATRIBUCION_INCOMPLETA"]


@pytest.mark.parametrize("state", [
    "PESADA", "ETIQUETADA_FINAL", "PENDIENTE_RECEPCION_ALMACEN", "RECIBIDA", "CONTINUIDAD_PENDIENTE",
])
def test_tv_does_not_reuse_open_control_from_terminal_manga_states(state):
    work = SimpleNamespace(id="work-terminal")
    manga = _tv_manga(1, work, "2")
    manga.estado = state
    manga._report_final_kg = None
    manga._report_open_kg = Decimal("2")

    breakdown = production_reports_service._progress_tv_breakdown(
        {"works": [(work, None, None)], "mangas": {1: manga}}, 201, "KG"
    )

    assert breakdown["medido"] is None
    assert breakdown["en_proceso"] is None
    assert breakdown["fisico_observado"] is None


def test_tv_un_keeps_kg_control_as_physical_observation_without_converting_to_units():
    work = SimpleNamespace(id="work-un-control")
    manga = _tv_manga(1, work, "2")
    manga.estado = "EN_LLENADO"
    manga._report_final_kg = None
    manga._report_open_kg = Decimal("2")
    manga._report_quantity_un = None

    breakdown = production_reports_service._progress_tv_breakdown(
        {"works": [(work, None, None)], "mangas": {1: manga}}, 201, "UN"
    )

    assert breakdown["medido"] is None
    assert breakdown["fisico_observado"] == Decimal("2")
    assert breakdown["incomplete"] is True
    assert "KG_SIN_TRAMOS_UN" in breakdown["motivos_atribucion"]


def test_tv_un_counters_use_quantity_and_exclude_foreign_without_weight_conversion():
    work = SimpleNamespace(id="work-un-own")
    foreign = SimpleNamespace(id="work-un-foreign")
    ot = SimpleNamespace(id="ot-un", public_id="ot-un", codigo_ot="OT-UN")

    own = _tv_manga(1, work, "0")
    own._report_final_kg = None
    own._report_quantity_un = Decimal("4")
    foreign_manga = _tv_manga(2, foreign, "0")
    foreign_manga._report_final_kg = None
    foreign_manga._report_quantity_un = Decimal("9")
    ownerless = _tv_manga(3, None, "0")
    ownerless._report_final_kg = None
    ownerless._report_quantity_un = Decimal("7")
    ownerless.ot_id = ot.id
    run = {"works": [(work, None, ot)], "mangas": {1: own, 2: foreign_manga, 3: ownerless}}

    breakdown = production_reports_service._progress_tv_breakdown(run, 201, "UN")

    assert breakdown["medido"] == Decimal("4")
    assert breakdown["mangas_total"] == 2
    assert breakdown["mangas_conocidas"] == 2
    assert breakdown["pendientes_sin_peso"] == 0
    assert breakdown["pendientes_sin_cantidad_un"] == 0
    assert breakdown["incomplete"] is True
    assert breakdown["motivos_atribucion"] == ["ATRIBUCION_INCOMPLETA"]


def test_tv_actual_keeps_open_control_with_active_unclosed_runtime_segment(monkeypatch):
    """Runtime shape: open manga has active _report_all_segments only."""
    work = SimpleNamespace(id="work-of122")
    article = SimpleNamespace(id=201)
    final = SimpleNamespace(
        id=1, estado="PESADA", trabajo=work, correccion_asignacion=None,
        lote_articulo=SimpleNamespace(articulo=article), _report_segments=[],
        _report_all_segments=[], _report_final_kg=Decimal("90.6"),
        _report_open_kg=None, _report_quantity_un=None, _report_weight_corrected=False,
    )
    active = _segment(1, "0", "0", "0")
    active.estado = "ACTIVO"
    active.cantidad_fin_kg = None
    active.cantidad_atribuida_kg = Decimal("0")
    active.trabajo = work
    opened = SimpleNamespace(
        id=2, estado="EN_LLENADO", trabajo=work, correccion_asignacion=None,
        lote_articulo=SimpleNamespace(articulo=article), _report_segments=[],
        _report_all_segments=[active], _report_final_kg=None,
        _report_open_kg=Decimal("1.1"), _report_quantity_un=None,
        _report_weight_corrected=False,
    )
    pending = SimpleNamespace(
        id=3, estado="EN_LLENADO", trabajo=work, correccion_asignacion=None,
        lote_articulo=SimpleNamespace(articulo=article), _report_segments=[],
        _report_all_segments=[], _report_final_kg=None, _report_open_kg=None,
        _report_quantity_un=None, _report_weight_corrected=False,
    )
    run = _tv_aggregate_run(works=[work], mangas={1: final, 2: opened, 3: pending}, output_meta="250")

    assert production_reports_service._progress_tv_actual(run, 201, "KG") == (Decimal("91.7"), True)
    monkeypatch.setattr(production_reports_service, "load_actor", lambda *_a, **_k: SimpleNamespace(tiene_capacidad=lambda _c: True))
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: [run])
    output = list_production_progress_tv(object(), actor_id=7, filters={})["items"][0]["salidas"][0]
    assert output["avance"] == pytest.approx(91.7)
    assert output["porcentaje"] == pytest.approx(36.68)
    assert output["avance_medido_kg"] == pytest.approx(91.7)
    assert output["finalizados_kg"] == pytest.approx(90.6)
    assert output["en_proceso_kg"] == pytest.approx(1.1)
    assert output["mangas_total"] == 3
    assert output["mangas_conocidas"] == 2
    assert output["pendientes_sin_peso"] == 1


def test_ownerless_ot_manga_without_evidence_is_pending_in_both_serializers(app, client, monkeypatch):
    work = SimpleNamespace(id="work-owner")
    known = _tv_manga(1, work, "3")
    pending = _tv_manga(2, None, "3")
    pending._report_final_kg = None
    pending._report_open_kg = None
    run = _tv_aggregate_run(works=[work], mangas={1: known, 2: pending}, output_meta="5")
    ot = run["contexts"][0]["ot"]
    run["ots"] = [ot]
    run["ot"] = ot
    pending.ot_id = ot.id
    run["corrida"].objetivo_neto_kg = Decimal("5")

    monkeypatch.setattr(
        production_reports_service, "load_actor",
        lambda *_a, **_k: SimpleNamespace(tiene_capacidad=lambda _c: True),
    )
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: [run])
    with app.app_context():
        app.config["SECRET_KEY"] = "ownerless-pending-test-secret"
        standard = client.get(
            "/api/scm/v1/observabilidad/avance-of",
            headers={"X-Actor-Id": "7"},
        ).get_json()["items"][0]
        tv = client.get(
            "/api/scm/v1/observabilidad/avance-of-tv",
            headers={"X-Actor-Id": "7"},
        ).get_json()["items"][0]["salidas"][0]

    assert standard["kg_medidos_efectivos"] == 3.0
    assert standard["mangas"] == {"total": 2, "conocidas": 1}
    assert standard["coverage"]["estado"] == "INCOMPLETA"
    assert tv["avance"] == 3.0
    assert tv["estado_avance"] == "INCOMPLETO"
    assert tv["mangas_total"] == 2
    assert tv["mangas_conocidas"] == 1
    assert tv["pendientes_sin_peso"] == 1


def test_standard_progress_keeps_partial_measured_and_percentage_separate_from_coverage(monkeypatch):
    work = SimpleNamespace(id="work-partial")
    known = SimpleNamespace(
        id=1, estado="PESADA", trabajo=work, _report_final_kg=Decimal("90.6"),
        _report_open_kg=None, _report_segments=[], _report_weight_corrected=False,
    )
    pending = SimpleNamespace(
        id=2, estado="EN_LLENADO", trabajo=work, _report_final_kg=None,
        _report_open_kg=None, _report_segments=[], _report_weight_corrected=False,
    )
    opened = SimpleNamespace(
        id=3, estado="EN_LLENADO", trabajo=work, _report_final_kg=None,
        _report_open_kg=Decimal("1.1"), _report_segments=[], _report_weight_corrected=False,
    )
    run = {
        "corrida": SimpleNamespace(id="run-partial", codigo="C-PARTIAL", objetivo_neto_kg=Decimal("250"), color_produccion=None, salidas=[]),
        "orden": SimpleNamespace(id="of-partial", codigo="OF-122", estado="EN_EJECUCION"),
        "ot": None, "color_name": "Melón", "color_hex": None, "molde": None,
        "mangas": {1: known, 2: pending, 3: opened}, "contexts": [{"work": work}],
    }
    monkeypatch.setattr(production_reports_service, "load_actor", lambda *_a, **_k: SimpleNamespace(tiene_capacidad=lambda _c: True))
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: [run])
    item = list_production_progress(object(), actor_id=1, filters={})["items"][0]

    assert item["kg_finalizados_efectivos"] == 90.6
    assert item["kg_medidos_en_abiertas"] == 1.1
    assert item["kg_medidos_efectivos"] == 91.7
    assert item["avance_medido_kg"] == 91.7
    assert item["porcentaje"] == pytest.approx(36.68)
    assert item["coverage"]["estado"] == "INCOMPLETA"
    assert "SUBTOTAL_PARCIAL" in item["coverage"]["motivos"]


def test_standard_pending_only_does_not_turn_missing_measurement_into_zero(monkeypatch):
    work = SimpleNamespace(id="work-pending-only")
    pending = SimpleNamespace(
        id=1, estado="EN_LLENADO", trabajo=work, _report_final_kg=None,
        _report_open_kg=None, _report_segments=[], _report_weight_corrected=False,
    )
    run = {
        "corrida": SimpleNamespace(
            id="run-pending-only", codigo="C-PENDING-ONLY", objetivo_neto_kg=Decimal("10"),
            color_produccion=None, salidas=[],
        ),
        "orden": SimpleNamespace(id="of-pending-only", codigo="OF-PENDING-ONLY", estado="EN_EJECUCION"),
        "ot": None, "color_name": "Azul", "color_hex": None, "molde": None,
        "mangas": {1: pending}, "contexts": [{"work": work}],
    }
    monkeypatch.setattr(
        production_reports_service,
        "load_actor",
        lambda *_a, **_k: SimpleNamespace(tiene_capacidad=lambda _capability: True),
    )
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: [run])

    item = list_production_progress(object(), actor_id=1, filters={})["items"][0]

    assert item["kg_medidos_efectivos"] is None
    assert item["subtotal_conocido_kg"] is None
    assert item["coverage"]["motivos"] == ["PESO_MANGA_SIN_EVIDENCIA"]


def test_final_vigente_replaces_control_on_same_manga():
    work = SimpleNamespace(id="work-replacement")
    manga = SimpleNamespace(
        id=1, estado="PESADA", trabajo=work, _report_final_kg=Decimal("90.6"),
        _report_open_kg=Decimal("1.1"), _report_segments=[], _report_weight_corrected=False,
    )
    run = {
        "corrida": SimpleNamespace(objetivo_neto_kg=Decimal("250")),
        "mangas": {1: manga}, "contexts": [{"work": work}],
    }

    final, opened, measured, total, known, _objective = _run_manga_values(run)

    assert (final, opened, measured, total, known) == (
        Decimal("90.6"), Decimal("0"), Decimal("90.6"), 1, 1,
    )


def test_tv_group_id_and_ot_order_do_not_depend_on_context_order(monkeypatch):
    work_a, work_b = SimpleNamespace(id="work-a"), SimpleNamespace(id="work-b")
    context_a = {"work": work_a, "color_work": SimpleNamespace(molde_codigo_snapshot="ML-1", color_id_snapshot=7),
                 "ot": SimpleNamespace(id=2, public_id="ot-b", codigo_ot="OT-B")}
    context_b = {"work": work_b, "color_work": SimpleNamespace(molde_codigo_snapshot="ML-1", color_id_snapshot=7),
                 "ot": SimpleNamespace(id=1, public_id="ot-a", codigo_ot="OT-A")}
    runs = [
        _tv_aggregate_run(works=[work_a, work_b], contexts=[context_a, context_b]),
        _tv_aggregate_run(works=[work_a, work_b], contexts=[context_b, context_a]),
    ]
    monkeypatch.setattr(production_reports_service, "load_actor", lambda *_a, **_k: SimpleNamespace(tiene_capacidad=lambda _c: True))
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: runs[:1])
    first = list_production_progress_tv(object(), actor_id=7, filters={})["items"][0]
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: runs[1:])
    second = list_production_progress_tv(object(), actor_id=7, filters={})["items"][0]

    assert first["group_id"] == second["group_id"] == "of-aggregate|run-aggregate|ML-1|7"
    assert first["ots"] == second["ots"] == [
        {"id": "ot-a", "codigo": "OT-A"}, {"id": "ot-b", "codigo": "OT-B"},
    ]
    assert first["ot_id"] is second["ot_id"] is None
    assert first["ot"] is second["ot"] is None
    assert first["identity"]["ot_id"] is second["identity"]["ot_id"] is None


def test_tv_mold_missing_has_same_cardinality_and_state_with_or_without_of_permission(monkeypatch):
    monkeypatch.setenv("SECRET_KEY", "tv-02-server-secret-for-tests")
    work = SimpleNamespace(id="work-mold-missing")
    context = {"work": work, "color_work": SimpleNamespace(molde_codigo_snapshot=None, color_id_snapshot=7),
               "ot": SimpleNamespace(id=1, public_id="ot-a", codigo_ot="OT-A")}
    run = _tv_aggregate_run(works=[work], contexts=[context], mangas={1: _tv_manga(1, work, "2")})
    run["orden"].molde_id = None
    run["moldes"] = {}
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_a, **_k: [run])
    visible_actor = SimpleNamespace(tiene_capacidad=lambda _capability: True)
    hidden_actor = SimpleNamespace(tiene_capacidad=lambda capability: capability != "OF_VER")
    monkeypatch.setattr(production_reports_service, "load_actor", lambda *_a, **_k: visible_actor)
    visible = list_production_progress_tv(object(), actor_id=7, filters={})
    monkeypatch.setattr(production_reports_service, "load_actor", lambda *_a, **_k: hidden_actor)
    hidden = list_production_progress_tv(object(), actor_id=7, filters={})

    assert len(visible["items"]) == len(hidden["items"]) == 1
    visible_row = visible["items"][0]["salidas"][0]
    hidden_row = hidden["items"][0]["salidas"][0]
    assert (visible_row["estado_avance"], visible_row["avance"], visible_row["porcentaje"]) == (
        hidden_row["estado_avance"], hidden_row["avance"], hidden_row["porcentaje"],
    ) == ("INCOMPLETO", None, None)
    assert visible["items"][0]["identity"]["identity_status"] == hidden["items"][0]["identity"]["identity_status"] == "INCOMPLETA"


def test_kg_segments_require_zero_continuity_and_exact_net():
    segments = [_segment(1, "0", "4", "4"), _segment(2, "4", "9", "5")]
    assert _segment_conciliates(segments, Decimal("9")) is True
    assert _segment_conciliates(segments, Decimal("8")) is False
    assert _segment_conciliates(
        [_segment(1, "0", "4", "3"), _segment(2, "4", "9", "5")],
        Decimal("9"),
    ) is False


def test_history_projects_corrected_net_only_into_last_kg_segment():
    from datetime import date

    work = SimpleNamespace(id="work-1", codigo="TR-1")
    ot = SimpleNamespace(
        fecha=date(2026, 9, 1), codigo_ot="OT-1", estado="CERRADA",
        maquina_nombre_snapshot="M1", maquina_codigo_snapshot=None,
        responsable=None, orden_operacion=None,
    )
    color_work = SimpleNamespace(peso_neto_snapshot_g=100)
    segments = [
        _segment(1, "0", "5", "5", attributed_un=50),
        _segment(2, "5", "12", "7", attributed_un=70),
    ]
    for segment in segments:
        segment.trabajo = work
    manga = SimpleNamespace(
        id=77, _report_final_kg=Decimal("11.900"), _report_segments=segments,
        _report_weight_corrected=True,
        peso_unitario_snapshot_g=100, cantidad_confirmada_un=120,
        cantidad_asignada_un=120, articulo_codigo_snapshot="A",
        articulo_nombre_snapshot="Artículo", correccion_asignacion=None,
        trabajo=work,
    )
    run = {
        "corrida": SimpleNamespace(id="run-1", codigo="C-1", objetivo_neto_kg=12),
        "orden": SimpleNamespace(codigo="OF-1", estado="CERRADA"),
        "ot": ot, "work": work, "color_work": color_work,
        "color_name": "Rojo", "resource": "M1", "responsible": None,
        "contexts": [{"work": work, "color_work": color_work, "ot": ot}],
        "mangas": {77: manga},
    }

    rows = _history_rows([run], ["DIA"], None)

    assert [row["PESO_KG"] for row in rows] == [5, 6.9]
    assert sum(row["SUBTOTAL_CONOCIDO_KG"] for row in rows) == 11.9
    assert [row["_manga_id"] for row in rows] == [77, 77]
    assert {row["DIA"] for row in rows} == {"2026-09-01"}
    assert [segment.cantidad_fin_kg for segment in segments] == [Decimal("5"), Decimal("12")]


@pytest.mark.parametrize("net, expected", [("12.500", [5, 7.5]), ("5.000", [None])])
def test_history_correction_projection_handles_positive_delta_and_crossed_frontier(net, expected):
    from datetime import date

    work = SimpleNamespace(id="work-1", codigo="TR-1")
    ot = SimpleNamespace(
        fecha=date(2026, 9, 1), codigo_ot="OT-1", estado="CERRADA",
        maquina_nombre_snapshot="M1", maquina_codigo_snapshot=None,
        responsable=None, orden_operacion=None,
    )
    segments = [_segment(1, "0", "5", "5"), _segment(2, "5", "12", "7")]
    for segment in segments:
        segment.trabajo = work
    manga = SimpleNamespace(
        id=77, _report_final_kg=Decimal(net), _report_segments=segments,
        _report_weight_corrected=True,
        peso_unitario_snapshot_g=100, cantidad_confirmada_un=120,
        cantidad_asignada_un=120, articulo_codigo_snapshot="A",
        articulo_nombre_snapshot="Artículo", correccion_asignacion=None,
        trabajo=work,
    )
    run = {
        "corrida": SimpleNamespace(id="run-1", codigo="C-1", objetivo_neto_kg=12),
        "orden": SimpleNamespace(codigo="OF-1", estado="CERRADA"),
        "ot": ot, "work": work, "color_work": SimpleNamespace(peso_neto_snapshot_g=100),
        "color_name": "Rojo", "resource": "M1", "responsible": None,
        "contexts": [{"work": work, "color_work": SimpleNamespace(peso_neto_snapshot_g=100), "ot": ot}],
        "mangas": {77: manga},
    }

    rows = _history_rows([run], ["DIA"], None)

    assert [row["PESO_KG"] for row in rows] == expected


@pytest.mark.parametrize("net, expected", [("11.900", "11.900"), ("12.500", "12.500")])
def test_corrected_net_projection_supports_one_segment_without_mutating_history(net, expected):
    segments = [_segment(1, "0", "12", "12")]

    projected = _project_corrected_kg_segments(segments, Decimal(net))

    assert len(projected) == 1
    assert projected[0].cantidad_fin_kg == Decimal(expected)
    assert projected[0].cantidad_atribuida_kg == Decimal(expected)
    assert projected[0].calidad_evidencia_kg == "CONCILIADA"
    assert segments[0].cantidad_fin_kg == Decimal("12")
    assert segments[0].cantidad_atribuida_kg == Decimal("12")


def test_corrected_net_projection_is_idempotent_and_keeps_un_axis_separate():
    segments = [_segment(1, "0", "12", "12")]
    first = _project_corrected_kg_segments(segments, Decimal("11.900"))
    second = _project_corrected_kg_segments(first, Decimal("11.900"))
    un_segment = SimpleNamespace(
        secuencia=1, cantidad_inicio_kg=None, cantidad_fin_kg=None,
        cantidad_atribuida_kg=None, calidad_evidencia_kg=None,
    )

    assert [item.cantidad_fin_kg for item in second] == [Decimal("11.900")]
    assert [item.cantidad_atribuida_kg for item in second] == [Decimal("11.900")]
    assert _project_corrected_kg_segments([un_segment], Decimal("11.900")) == [un_segment]


def test_un_and_default_kg_segments_are_not_evidence():
    manga = SimpleNamespace(
        tramos_trabajo=[
            _segment(1, "0", "4", "4", quality="MEDIDA_DIRECTA"),
            _segment(2, "4", "9", "5", quality="PENDIENTE"),
            _segment(3, "9", "10", "1", quality="MEDIDA_DIRECTA"),
        ]
    )
    assert [segment.secuencia for segment in _valid_kg_segments(manga, {})] == [1, 3]


def test_explicit_control_close_is_kg_evidence():
    manga = SimpleNamespace(tramos_trabajo=[_segment(1, "0", "9", "9", quality="MEDIDA_DIRECTA_CIERRE_CONTROL")])
    assert [segment.secuencia for segment in _valid_kg_segments(manga, {})] == [1]


def test_objective_without_mangas_does_not_claim_complete_weight_coverage():
    corrida = SimpleNamespace(objetivo_neto_kg=Decimal("10"))
    run = {"corrida": corrida, "mangas": {}}
    _final, _open, measured, total, known, _objective = _run_manga_values(run)
    assert (measured, total, known) == (None, 0, 0)


def _progress_manga(manga_id, *, state="PENDIENTE_RECEPCION_ALMACEN", final=Decimal("6"), open_kg=None, work=None):
    return SimpleNamespace(
        id=manga_id, estado=state, trabajo=work,
        _report_final_kg=final, _report_open_kg=open_kg,
        _report_segments=[], _report_weight_corrected=False,
    )


def test_progress_excludes_cancelled_mangas_from_coverage_and_totals(monkeypatch):
    work = SimpleNamespace(id="work-of78")
    active = {
        index: _progress_manga(index, final=(Decimal("55.7") if index == 68 else Decimal("6")), work=work)
        for index in range(1, 69)
    }
    cancelled = {
        100 + index: _progress_manga(100 + index, state="ANULADA", final=Decimal("10"), work=work)
        for index in range(1, 10)
    }
    run = {
        "corrida": SimpleNamespace(objetivo_neto_kg=Decimal("1000")),
        "mangas": {**active, **cancelled},
        "contexts": [{"work": work}],
    }

    final, opened, measured, total, known, objective = _run_manga_values(run)

    assert (final, opened, measured, total, known, objective) == (
        Decimal("457.7"), Decimal("0"), Decimal("457.7"), 68, 68, Decimal("1000")
    )
    run.update({
        "corrida": SimpleNamespace(
            id="run-of78", codigo="C-OF78", objetivo_neto_kg=Decimal("1000"),
            color_produccion=None, salidas=[],
        ),
        "orden": SimpleNamespace(codigo="OF-000078", estado="ABIERTA"),
        "ot": None, "color_name": "Rojo", "molde": None,
    })
    monkeypatch.setattr(production_reports_service, "load_actor", lambda *_args, **_kwargs: SimpleNamespace(tiene_capacidad=lambda _capability: True))
    monkeypatch.setattr(production_reports_service, "_load_rows", lambda *_args, **_kwargs: [run])
    item = list_production_progress(object(), actor_id=1, filters={})["items"][0]
    assert item["kg_finalizados_efectivos"] == 457.7
    assert item["porcentaje"] == pytest.approx(45.77)
    assert item["coverage"]["estado"] == "COMPLETA"
    assert item["mangas"] == {"total": 68, "conocidas": 68}


def test_progress_excludes_cancelled_manga_even_when_old_weight_is_present():
    work = SimpleNamespace(id="work-cancelled")
    run = {
        "corrida": SimpleNamespace(objetivo_neto_kg=Decimal("100")),
        "mangas": {1: _progress_manga(1, state="ANULADA", final=Decimal("90"), work=work)},
        "contexts": [{"work": work}],
    }

    final, opened, measured, total, known, _objective = _run_manga_values(run)

    assert (final, opened, measured, total, known) == (Decimal("0"), Decimal("0"), None, 0, 0)


def test_progress_with_only_cancelled_mangas_has_no_fabricated_evidence():
    work = SimpleNamespace(id="work-only-cancelled")
    run = {
        "corrida": SimpleNamespace(objetivo_neto_kg=Decimal("100")),
        "mangas": {1: _progress_manga(1, state="ANULADA", final=None, open_kg=None, work=work)},
        "contexts": [{"work": work}],
    }

    final, opened, measured, total, known, _objective = _run_manga_values(run)

    assert (final, opened, measured, total, known) == (Decimal("0"), Decimal("0"), None, 0, 0)


def test_progress_active_manga_without_evidence_remains_incomplete():
    work = SimpleNamespace(id="work-unknown")
    run = {
        "corrida": SimpleNamespace(objetivo_neto_kg=Decimal("100")),
        "mangas": {1: _progress_manga(1, final=None, open_kg=None, work=work)},
        "contexts": [{"work": work}],
    }

    final, opened, measured, total, known, _objective = _run_manga_values(run)

    assert (final, opened, measured, total, known) == (Decimal("0"), Decimal("0"), None, 1, 0)


def test_progress_exposes_canonical_color_hex_or_null(monkeypatch):
    actor = SimpleNamespace(tiene_capacidad=lambda _capability: True)
    corrida = SimpleNamespace(
        id=1,
        codigo="C-HEX",
        objetivo_neto_kg=Decimal("10"),
        color_produccion=SimpleNamespace(nombre="Rojo", hex_referencia="#AABBCC"),
        salidas=[],
    )
    no_hex_corrida = SimpleNamespace(
        id=2,
        codigo="C-NO-HEX",
        objetivo_neto_kg=Decimal("10"),
        color_produccion=SimpleNamespace(nombre="Sin referencia", hex_referencia=None),
        salidas=[],
    )
    base = {
        "orden": SimpleNamespace(codigo="OF-HEX", estado="ABIERTA"),
        "ot": None,
        "color_name": "Rojo",
        "molde": None,
        "mangas": {},
        "contexts": [],
    }
    monkeypatch.setattr(production_reports_service, "load_actor", lambda *_args, **_kwargs: actor)
    monkeypatch.setattr(
        production_reports_service,
        "_load_rows",
        lambda *_args, **_kwargs: [
            {**base, "corrida": corrida},
            {**base, "corrida": no_hex_corrida},
        ],
    )

    payload = list_production_progress(object(), actor_id=1, filters={})

    assert [item["color_hex"] for item in payload["items"]] == ["#AABBCC", None]


def test_progress_attributes_one_manga_once_per_effective_work_segment():
    work_a = SimpleNamespace(id="work-a")
    work_b = SimpleNamespace(id="work-b")
    first = _segment(1, "0", "4", "4")
    second = _segment(2, "4", "9", "5")
    first.trabajo = work_a
    second.trabajo = work_b
    manga = SimpleNamespace(
        id=77,
        _report_final_kg=Decimal("9"),
        _report_open_kg=None,
        _report_segments=[first, second],
        _report_weight_corrected=False,
    )
    run_a = {"corrida": SimpleNamespace(objetivo_neto_kg=9), "mangas": {77: manga}, "contexts": [{"work": work_a}]}
    run_b = {"corrida": SimpleNamespace(objetivo_neto_kg=9), "mangas": {77: manga}, "contexts": [{"work": work_b}]}

    assert _run_manga_values(run_a)[:5] == (Decimal("4"), Decimal("0"), Decimal("4"), 1, 1)
    assert _run_manga_values(run_b)[:5] == (Decimal("5"), Decimal("0"), Decimal("5"), 1, 1)


def test_open_control_is_attributed_only_when_kg_segments_conciliate():
    work_a = SimpleNamespace(id="work-a")
    work_b = SimpleNamespace(id="work-b")
    first = _segment(1, "0", "4", "4")
    second = _segment(2, "4", "6", "2")
    first.trabajo = work_a
    second.trabajo = work_b
    manga = SimpleNamespace(
        id=78, _report_final_kg=None, _report_open_kg=Decimal("6"),
        _report_segments=[first, second], _report_weight_corrected=False,
        trabajo=work_b,
    )
    run_a = {"corrida": SimpleNamespace(objetivo_neto_kg=9), "mangas": {78: manga}, "contexts": [{"work": work_a}]}
    run_b = {"corrida": SimpleNamespace(objetivo_neto_kg=9), "mangas": {78: manga}, "contexts": [{"work": work_b}]}

    assert _run_manga_values(run_a)[:5] == (Decimal("0"), Decimal("4"), Decimal("4"), 1, 1)
    assert _run_manga_values(run_b)[:5] == (Decimal("0"), Decimal("2"), Decimal("2"), 1, 1)

    manga._report_open_kg = Decimal("5")
    assert _run_manga_values(run_a)[:5] == (Decimal("0"), Decimal("0"), None, 1, 0)
    assert _run_manga_values(run_b)[:5] == (Decimal("0"), Decimal("0"), None, 1, 0)


def test_identity_dimensions_keep_canonical_key_and_readable_metadata():
    row = {
        "MOLDE": "ML-001", "MOLDE_NOMBRE": "Molde multipieza", "MOLDE_CODIGO": "ML-001",
        "PIEZA": "PZ-BASE", "PIEZA_NOMBRE": "Pieza base", "PIEZA_CODIGO": "PZ-BASE",
        "PESO_KG": 2, "SUBTOTAL_CONOCIDO_KG": 2, "SUBTOTAL_TEORICO_KG": None,
        "MANGAS": 1, "P_UNITARIO_WEIGHT": None, "P_UNITARIO_QTY": None,
        "P_TEORICO_KG": None, "_known": True, "_manga_id": 1,
    }
    items = _group_history_rows([row], {"groups": ["MOLDE", "PIEZA"], "measures": list(MEASURE_OPTIONS)})

    assert items[0]["MOLDE"] == "ML-001"
    assert items[0]["MOLDE_NOMBRE"] == "Molde multipieza"
    assert items[0]["PIEZA"] == "PZ-BASE"
    assert items[0]["PIEZA_NOMBRE"] == "Pieza base"


def test_history_uses_effective_work_mold_snapshot_over_of_mold():
    context = {
        "work": SimpleNamespace(id="work-a"),
        "color_work": SimpleNamespace(molde_codigo_snapshot="ML-SNAPSHOT"),
        "ot": SimpleNamespace(
            fecha=date(2026, 9, 1), codigo_ot="OT-1", estado="CERRADA",
            maquina_nombre_snapshot="M1", maquina_codigo_snapshot=None, responsable=None,
        ),
    }
    run = {
        "of": SimpleNamespace(molde_id="ML-OF"),
        "moldes": {
            "ML-SNAPSHOT": SimpleNamespace(codigo="ML-SNAPSHOT", nombre="Molde snapshot"),
            "ML-OF": SimpleNamespace(codigo="ML-OF", nombre="Molde OF"),
        },
        "corrida": SimpleNamespace(codigo="C-1"),
        "orden": SimpleNamespace(codigo="OF-1"),
        "color_name": "ROJO",
    }
    assert _context_group_value(run, context, "MOLDE") == "ML-SNAPSHOT"


def test_history_color_hex_is_canonical_and_mixed_groups_are_neutral():
    base = {
        "OF": "OF-1", "COLOR": "Rojo", "COLOR_CODIGO": 10,
        "COLOR_HEX": "#AABBCC", "_COLOR_ID": 1,
        "PESO_KG": 2, "SUBTOTAL_CONOCIDO_KG": 2,
        "SUBTOTAL_TEORICO_KG": None, "MANGAS": 1,
        "P_UNITARIO_WEIGHT": None, "P_UNITARIO_QTY": None,
        "P_TEORICO_KG": None, "_known": True, "_manga_id": 1,
    }
    normalized = {"groups": ["OF"], "measures": list(MEASURE_OPTIONS)}

    single = _group_history_rows([base], normalized)[0]
    assert single["COLOR_HEX"] == "#AABBCC"
    assert single["COLOR_CODIGO"] == 10

    mixed = _group_history_rows([
        base,
        {**base, "COLOR": "Azul", "COLOR_CODIGO": 20, "COLOR_HEX": "#112233", "_COLOR_ID": 2, "_manga_id": 2},
    ], normalized)[0]
    assert mixed["COLOR_HEX"] is None
    assert mixed["COLOR_CODIGO"] is None

    hierarchy, _summary = _history_hierarchy(
        [base, {**base, "COLOR": "Azul", "COLOR_CODIGO": 20, "COLOR_HEX": "#112233", "_COLOR_ID": 2, "_manga_id": 2}],
        ["OF"],
        list(MEASURE_OPTIONS),
    )
    assert hierarchy[0]["item"]["COLOR_HEX"] is None

    work = SimpleNamespace(id="work-hex", codigo="TR-HEX")
    ot = SimpleNamespace(
        fecha=date(2026, 9, 1), codigo_ot="OT-HEX", estado="CERRADA",
        maquina_nombre_snapshot="M1", maquina_codigo_snapshot=None,
        responsable=None,
    )
    segment = _segment(101, "0", "2", "2")
    segment.trabajo = work
    manga = SimpleNamespace(
        id=101, _report_final_kg=Decimal("2"), _report_segments=[segment],
        _report_weight_corrected=False, peso_unitario_snapshot_g=100,
        cantidad_confirmada_un=20, cantidad_asignada_un=20,
        articulo_codigo_snapshot="ART-HEX", articulo_nombre_snapshot="Artículo hex",
        _report_identity={"pieza_codigo": "PZ-HEX", "pieza_nombre": "Pieza hex"},
    )
    run = {
        "corrida": SimpleNamespace(id="run-hex", codigo="C-HEX", objetivo_neto_kg=2),
        "orden": SimpleNamespace(codigo="OF-HEX", estado="CERRADA"),
        "ot": ot, "work": work, "color_work": SimpleNamespace(peso_neto_snapshot_g=100),
        "color_name": "Rojo", "color_id": 1, "color_code": 10,
        "color_hex": "#AABBCC", "resource": "M1", "responsible": None,
        "contexts": [{"work": work, "color_work": SimpleNamespace(peso_neto_snapshot_g=100), "ot": ot}],
        "mangas": {101: manga},
    }
    rows = _history_rows([run], ["OF", "COLOR"], None)
    assert rows[0]["COLOR_HEX"] == "#AABBCC"


def test_history_subtotal_deduplicates_manga_across_kg_segments_and_groups():
    rows = []
    for day, ot, kg in (("2026-09-01", "OT-1", 4), ("2026-09-02", "OT-2", 5)):
        rows.append({
            "DIA": day, "MES": "2026-09", "OF": "OF-1", "CORRIDA": "C-1",
            "COLOR": "Rojo", "OT": ot, "RECURSO": "M1", "RESPONSABLE": "R",
            "ARTICULO": "A", "ARTICULO_NOMBRE": "Artículo legible", "PESO_KG": kg, "SUBTOTAL_CONOCIDO_KG": kg,
            "SUBTOTAL_TEORICO_KG": None, "MANGAS": 1, "P_UNITARIO_G": 100,
            "P_UNITARIO_WEIGHT": 100, "P_UNITARIO_QTY": 1, "P_TEORICO_KG": 1,
            "_known": True, "_manga_id": 77,
        })
    for groups in (("OF",), ("DIA",), ("DIA", "OT")):
        items = _group_history_rows(rows, {"groups": list(groups), "measures": list(MEASURE_OPTIONS)})
        assert len(items) == (1 if groups == ("OF",) else 2)
        assert sum(item["PESO_KG"] for item in items) == 9
        assert sorted(item["PESO_KG"] for item in items) == ([9] if groups == ("OF",) else [4, 5])

    article = _group_history_rows(rows, {"groups": ["ARTICULO"], "measures": list(MEASURE_OPTIONS)})[0]
    assert article["ARTICULO"] == "A"
    assert article["ARTICULO_CODIGO"] == "A"
    assert article["ARTICULO_NOMBRE"] == "Artículo legible"


def test_theoretical_measure_stays_null_without_un_evidence():
    row = {
        "DIA": "2026-09-01", "PESO_KG": 9, "SUBTOTAL_CONOCIDO_KG": 9,
        "SUBTOTAL_TEORICO_KG": 2, "P_TEORICO_KG": None, "MANGAS": 1,
        "P_UNITARIO_WEIGHT": None, "P_UNITARIO_QTY": None, "P_UNITARIO_G": None,
        "coverage": "INCOMPLETA", "_known": False, "_manga_id": 11,
    }
    item = _group_history_rows([row], {"groups": ["DIA"], "measures": list(MEASURE_OPTIONS)})[0]
    assert item["P_TEORICO_KG"] is None
    assert item["SUBTOTAL_TEORICO_KG"] == 2


def test_context_filter_does_not_turn_excluded_conciliated_segments_into_fallback():
    from datetime import date

    work = SimpleNamespace(id="work-1", codigo="TR-1")
    ot = SimpleNamespace(
        fecha=date(2026, 9, 1), codigo_ot="OT-1", estado="CERRADA",
        maquina_nombre_snapshot="M1", maquina_codigo_snapshot=None,
        responsable=None, orden_operacion=None,
    )
    color_work = SimpleNamespace(peso_neto_snapshot_g=100)
    manga = SimpleNamespace(
        id=77, _report_final_kg=Decimal("9"),
        _report_segments=[
            _segment(1, "0", "4", "4", attributed_un=40),
            _segment(2, "4", "9", "5", attributed_un=50),
        ], peso_unitario_snapshot_g=100, cantidad_confirmada_un=90,
        cantidad_asignada_un=90, articulo_codigo_snapshot="A",
        articulo_nombre_snapshot="Artículo buscable",
        correccion_asignacion=None,
    )
    manga._report_segments[0].trabajo = work
    manga._report_segments[1].trabajo = work
    run = {
        "corrida": SimpleNamespace(id="run-1", codigo="C-1", objetivo_neto_kg=9),
        "orden": SimpleNamespace(codigo="OF-1", estado="CERRADA"),
        "ot": ot, "work": work, "color_work": color_work,
        "color_name": "Rojo", "resource": "M1", "responsible": None,
        "contexts": [{"work": work, "color_work": color_work, "ot": ot}],
        "mangas": {77: manga},
    }
    from app.services.scm_production_reports_service import _history_rows
    for raw in ({"q": "OT-2"}, {"estado_ot": "ANULADA"}):
        filters = _filters({"fecha_desde": "2026-09-01", "fecha_hasta": "2026-09-01", **raw})
        assert _history_rows([run], ["DIA"], filters) == []

    for raw in ({"q": "buscable"}, {"articulo": "Artículo buscable"}, {"articulo": "A"}):
        filters = _filters({"fecha_desde": "2026-09-01", "fecha_hasta": "2026-09-01", **raw})
        assert len(_history_rows([run], ["ARTICULO"], filters)) == 2


def test_report_filters_reject_unknown_group_and_measure_without_fallback():
    for key, value, code in (
        ("agrupaciones", "NO_EXISTE", "INVALID_OBSERVABILITY_GROUP"),
        ("medidas", "NO_EXISTE", "INVALID_OBSERVABILITY_MEASURE"),
    ):
        try:
            _filters({"fecha_desde": "2026-09-01", "fecha_hasta": "2026-09-02", key: value})
        except ScmServiceError as error:
            assert error.status_code == 400
            assert error.code == code
        else:
            raise AssertionError("el filtro inválido no debe usar fallback")


def test_report_filters_distinguish_omitted_from_explicit_empty_values():
    omitted = _filters({"fecha_desde": "2026-09-01", "fecha_hasta": "2026-09-02"})
    total = _filters({"fecha_desde": "2026-09-01", "fecha_hasta": "2026-09-02", "agrupaciones": ""})
    assert omitted["groups"] == ["DIA"]
    assert total["groups"] == []
    try:
        _filters({"fecha_desde": "2026-09-01", "fecha_hasta": "2026-09-02", "medidas": ""})
    except ScmServiceError as error:
        assert error.code == "INVALID_OBSERVABILITY_MEASURE"
        assert error.status_code == 400
    else:
        raise AssertionError("medidas vacías deben ser inválidas explícitamente")


def test_report_filters_reject_empty_tokens_inside_nonempty_lists():
    for key, value, code in (
        ("agrupaciones", "DIA,", "INVALID_OBSERVABILITY_GROUP"),
        ("agrupaciones", "DIA,,OT", "INVALID_OBSERVABILITY_GROUP"),
        ("medidas", "PESO_KG,", "INVALID_OBSERVABILITY_MEASURE"),
    ):
        try:
            _filters({"fecha_desde": "2026-09-01", "fecha_hasta": "2026-09-02", key: value})
        except ScmServiceError as error:
            assert error.code == code
            assert error.status_code == 400
        else:
            raise AssertionError("los tokens vacíos no deben desaparecer silenciosamente")


def test_xlsx_weight_summary_keeps_partial_known_subtotal_without_false_zero():
    total, coverage = _history_weight_summary(
        [{"PESO_KG": Decimal("4.5")}, {"PESO_KG": None}], ["PESO_KG"]
    )
    assert total == Decimal("4.5")
    assert coverage == "Parcial: 1 fila(s) sin peso"
    assert _history_weight_summary([{"PESO_KG": None}], ["MANGAS"]) == ("No solicitado", None)


def test_history_hierarchy_counts_a_manga_once_and_weights_unit_average_from_evidence():
    rows = [
        {
            "DIA": "2026-09-01", "OT": "OT-1", "OF": "OF-1", "CORRIDA": "C-1",
            "COLOR": "Rojo", "RECURSO": "M1", "RESPONSABLE": "R", "ARTICULO": "A",
            "ARTICULO_NOMBRE": "Artículo", "PESO_KG": 4, "SUBTOTAL_CONOCIDO_KG": 4,
            "SUBTOTAL_TEORICO_KG": None, "MANGAS": 1, "P_UNITARIO_WEIGHT": 1000,
            "P_UNITARIO_QTY": 10, "P_UNITARIO_G": 100, "P_TEORICO_KG": None,
            "_known": True, "_manga_id": 77,
        },
        {
            "DIA": "2026-09-02", "OT": "OT-2", "OF": "OF-1", "CORRIDA": "C-1",
            "COLOR": "Rojo", "RECURSO": "M1", "RESPONSABLE": "R", "ARTICULO": "A",
            "ARTICULO_NOMBRE": "Artículo", "PESO_KG": 5, "SUBTOTAL_CONOCIDO_KG": 5,
            "SUBTOTAL_TEORICO_KG": None, "MANGAS": 1, "P_UNITARIO_WEIGHT": 18000,
            "P_UNITARIO_QTY": 90, "P_UNITARIO_G": 200, "P_TEORICO_KG": None,
            "_known": True, "_manga_id": 77,
        },
    ]

    hierarchy, summary = _history_hierarchy(rows, ["DIA", "OT"], list(MEASURE_OPTIONS))

    assert len(hierarchy) == 2
    assert sum(node["item"]["MANGAS"] for node in hierarchy) == 2
    assert summary["MANGAS"] == 1
    assert summary["PESO_KG"] == 9.0
    assert summary["P_UNITARIO_G"] == 190.0
    assert all(node["children"] for node in hierarchy)
    assert hierarchy[0]["id"] != hierarchy[1]["id"]


def test_history_real_rows_keep_weighted_unit_average_independent_of_grain(monkeypatch):
    from datetime import date
    from app.services import scm_production_reports_service as reports

    def context(work, day, ot_code, unit):
        ot = SimpleNamespace(
            fecha=date.fromisoformat(day), codigo_ot=ot_code, estado="CERRADA",
            maquina_nombre_snapshot="M1", maquina_codigo_snapshot=None,
            responsable=None,
        )
        return {"work": work, "color_work": SimpleNamespace(peso_neto_snapshot_g=Decimal(unit)), "ot": ot}

    work_a1 = SimpleNamespace(id="work-a1", codigo="TR-A1")
    work_a2 = SimpleNamespace(id="work-a2", codigo="TR-A2")
    work_b = SimpleNamespace(id="work-b", codigo="TR-B")
    contexts = [
        context(work_a1, "2026-09-01", "OT-1", "100"),
        context(work_a2, "2026-09-02", "OT-2", "100"),
        context(work_b, "2026-09-01", "OT-3", "200"),
    ]

    segments_a = [
        _segment(1, "0", "4", "4", attributed_un=40),
        _segment(2, "4", "9", "5", attributed_un=50),
    ]
    segments_a[0].trabajo = work_a1
    segments_a[1].trabajo = work_a2
    segment_b = _segment(3, "0", "2", "2", attributed_un=10)
    segment_b.trabajo = work_b
    manga_a = SimpleNamespace(
        id=77, _report_final_kg=Decimal("9"), _report_segments=segments_a,
        peso_unitario_snapshot_g=Decimal("100"), cantidad_confirmada_un=90,
        cantidad_asignada_un=90, articulo_codigo_snapshot="A",
        articulo_nombre_snapshot="Artículo A", correccion_asignacion=None,
    )
    manga_b = SimpleNamespace(
        id=88, _report_final_kg=Decimal("2"), _report_segments=[segment_b],
        peso_unitario_snapshot_g=Decimal("200"), cantidad_confirmada_un=10,
        cantidad_asignada_un=10, articulo_codigo_snapshot="B",
        articulo_nombre_snapshot="Artículo B", correccion_asignacion=None,
    )
    run = {
        "corrida": SimpleNamespace(id="run-1", codigo="C-1", objetivo_neto_kg=11),
        "orden": SimpleNamespace(codigo="OF-1", estado="CERRADA"),
        "ot": contexts[0]["ot"], "work": work_a1, "color_work": contexts[0]["color_work"],
        "color_name": "Rojo", "resource": "M1", "responsible": None,
        "contexts": contexts, "mangas": {77: manga_a, 88: manga_b},
    }

    measures = list(MEASURE_OPTIONS)
    for groups, raw_groups in (([], ""), (["DIA", "OT"], "DIA,OT")):
        rows = _history_rows([run], groups, None)
        flat = _group_history_rows(rows, {"groups": groups, "measures": measures})[0]
        _hierarchy, summary = _history_hierarchy(rows, groups, measures)
        if not groups:
            assert flat["P_UNITARIO_G"] == 110.0
        assert summary["P_UNITARIO_G"] == 110.0

        monkeypatch.setattr(reports, "load_actor", lambda *_args, **_kwargs: SimpleNamespace(tiene_capacidad=lambda _capability: True))
        monkeypatch.setattr(reports, "_load_rows", lambda _session, _filters: [run])
        payload = list_production_history(
            object(), actor_id=1,
            filters={
                "fecha_desde": "2026-09-01", "fecha_hasta": "2026-09-02",
                "agrupaciones": raw_groups, "medidas": "P_UNITARIO_G",
                "incluir_jerarquia": "1",
            },
        )
        if not groups:
            assert payload["items"][0]["P_UNITARIO_G"] == 110.0
            workbook = load_workbook(
                BytesIO(generate_production_history_xlsx(
                    object(), actor_id=1,
                    filters={
                        "fecha_desde": "2026-09-01", "fecha_hasta": "2026-09-02",
                        "agrupaciones": "", "medidas": "P_UNITARIO_G",
                    },
                ).getvalue()),
                read_only=True,
                data_only=True,
            )
            headers = next(workbook["Datos"].iter_rows(values_only=True))
            values = next(workbook["Datos"].iter_rows(min_row=2, values_only=True))
            assert headers[0] == "P_UNITARIO_G"
            assert values[0] == 110.0
        assert payload["resumen"]["P_UNITARIO_G"] == 110.0


def test_history_hierarchy_empty_is_null_summary_and_no_nodes():
    assert _history_hierarchy([], ["DIA"], list(MEASURE_OPTIONS)) == ([], None)


def test_history_hierarchy_keeps_unattributed_manga_out_of_count_and_weight():
    rows = [{
        "DIA": "2026-09-01", "PESO_KG": None, "SUBTOTAL_CONOCIDO_KG": 4,
        "SUBTOTAL_TEORICO_KG": None, "MANGAS": 0, "P_UNITARIO_WEIGHT": None,
        "P_UNITARIO_QTY": None, "P_TEORICO_KG": None, "_known": False,
        "_manga_id": 77,
    }]

    hierarchy, summary = _history_hierarchy(rows, ["DIA"], list(MEASURE_OPTIONS))

    assert hierarchy[0]["item"]["MANGAS"] == 0
    assert hierarchy[0]["item"]["PESO_KG"] is None
    assert summary["MANGAS"] == 0
    assert summary["PESO_KG"] is None


def test_history_xlsx_uses_flat_items_without_hierarchy_parents(monkeypatch):
    from app.services import scm_production_reports_service as reports

    def flat_payload(_session, *, actor_id, filters=None):
        return {
            "grouped_by": ["OF"],
            "measures": ["PESO_KG"],
            "items": [
                {"OF": "OF-1", "PESO_KG": 4.0, "SUBTOTAL_CONOCIDO_KG": 4.0, "SUBTOTAL_TEORICO_KG": None, "coverage": "COMPLETA"},
                {"OF": "OF-2", "PESO_KG": 5.0, "SUBTOTAL_CONOCIDO_KG": 5.0, "SUBTOTAL_TEORICO_KG": None, "coverage": "COMPLETA"},
            ],
        }

    monkeypatch.setattr(reports, "load_actor", lambda *_args, **_kwargs: SimpleNamespace(tiene_capacidad=lambda _capability: True))
    monkeypatch.setattr(reports, "list_production_history", flat_payload)
    workbook = load_workbook(
        BytesIO(generate_production_history_xlsx(object(), actor_id=1, filters={"incluir_jerarquia": "1"}).getvalue()),
        read_only=True,
        data_only=True,
    )

    assert workbook["Datos"].max_row == 3
    assert [row[0].value for row in workbook["Datos"].iter_rows(min_row=2)] == ["OF-1", "OF-2"]


def test_history_hierarchy_is_opt_in_and_legacy_payload_stays_flat(app, client, scm_config):
    from test_scm_production_observability import _seed_observability_graph

    with app.app_context():
        seeded = _seed_observability_graph()
        params = {"fecha_desde": "2026-01-01", "fecha_hasta": "2026-12-31", "agrupaciones": "DIA"}
        legacy = client.get(
            "/api/scm/v1/observabilidad/produccion-historica",
            query_string=params,
            headers={"X-Actor-Id": str(seeded["full"].id)},
        )
        hierarchical = client.get(
            "/api/scm/v1/observabilidad/produccion-historica",
            query_string={**params, "incluir_jerarquia": "1"},
            headers={"X-Actor-Id": str(seeded["full"].id)},
        )

        assert legacy.status_code == hierarchical.status_code == 200
        assert "jerarquia" not in legacy.get_json()
        assert "resumen" not in legacy.get_json()
        assert isinstance(hierarchical.get_json()["jerarquia"], list)
        assert hierarchical.get_json()["resumen"] is None


def test_history_requires_operational_date_range(app, client, scm_config):
    from test_scm_production_observability import _seed_observability_graph

    with app.app_context():
        seeded = _seed_observability_graph()
        response = client.get(
            "/api/scm/v1/observabilidad/produccion-historica",
            headers={"X-Actor-Id": str(seeded["full"].id)},
        )
        assert response.status_code == 400
        assert "fecha_desde" in response.get_json()["error"]["message"]


def test_history_hierarchy_restricted_actor_returns_neutral_empty_payload(app, client, scm_config):
    from test_scm_production_observability import _seed_observability_graph

    with app.app_context():
        seeded = _seed_observability_graph()
        response = client.get(
            "/api/scm/v1/observabilidad/produccion-historica",
            query_string={
                "desde": "2026-09-01",
                "hasta": "2026-09-28",
                "incluir_jerarquia": "1",
            },
            headers={"X-Actor-Id": str(seeded["base"].id)},
        )

        payload = response.get_json()
        assert response.status_code == 200
        assert payload["visibilidad"]["pesaje"] is False
        assert payload["items"] == []
        assert payload["jerarquia"] == []
        assert payload["resumen"] is None
        assert payload["subtotal_conocido_kg"] == 0


def test_history_export_requires_weighing_capability_and_returns_workbook(app, client, scm_config):
    from test_scm_production_observability import _seed_observability_graph

    with app.app_context():
        seeded = _seed_observability_graph()
        params = {"fecha_desde": "2026-08-01", "fecha_hasta": "2026-08-31", "agrupaciones": "ARTICULO"}
        denied = client.get(
            "/api/scm/v1/observabilidad/produccion-historica/export.xlsx",
            query_string=params,
            headers={"X-Actor-Id": str(seeded["base"].id)},
        )
        assert denied.status_code == 403
        exported = client.get(
            "/api/scm/v1/observabilidad/produccion-historica/export.xlsx",
            query_string=params,
            headers={"X-Actor-Id": str(seeded["full"].id)},
        )
        assert exported.status_code == 200
        assert exported.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        workbook = load_workbook(BytesIO(exported.data), read_only=True, data_only=True)
        assert workbook.sheetnames == ["Resumen", "Datos"]
        headers = next(workbook["Datos"].iter_rows(values_only=True))
        assert headers[:2] == ("ARTICULO_NOMBRE", "ARTICULO_CODIGO")
        assert "SUBTOTAL_CONOCIDO_KG" in headers
        assert "PESO_KG" in headers
        mangas_only = client.get(
            "/api/scm/v1/observabilidad/produccion-historica/export.xlsx",
            query_string={**params, "medidas": "MANGAS"},
            headers={"X-Actor-Id": str(seeded["full"].id)},
        )
        mangas_book = load_workbook(BytesIO(mangas_only.data), read_only=True, data_only=True)
        assert mangas_only.status_code == 200
        assert mangas_book["Resumen"]["B2"].value == "No solicitado"
        mangas_headers = next(mangas_book["Datos"].iter_rows(values_only=True))
        assert mangas_headers == ("ARTICULO_NOMBRE", "ARTICULO_CODIGO", "MANGAS", "coverage")
        mangas_json = client.get(
            "/api/scm/v1/observabilidad/produccion-historica",
            query_string={**params, "medidas": "MANGAS", "incluir_jerarquia": "1"},
            headers={"X-Actor-Id": str(seeded["full"].id)},
        ).get_json()
        assert "subtotal_conocido_kg" not in mangas_json
        def assert_no_unrequested_weight(item):
            assert not {"PESO_KG", "P_UNITARIO_G", "P_TEORICO_KG", "SUBTOTAL_CONOCIDO_KG", "SUBTOTAL_TEORICO_KG"}.intersection(item)
        for item in mangas_json["items"]:
            assert_no_unrequested_weight(item)
        if mangas_json["resumen"]:
            assert_no_unrequested_weight(mangas_json["resumen"])
        for group in mangas_json["jerarquia"]:
            assert_no_unrequested_weight(group["item"])


def test_progress_http_keeps_positive_objective_without_mangas_incomplete(
    app, client, scm_config
):
    from test_scm_production_observability import _seed_observability_graph

    with app.app_context():
        seeded = _seed_observability_graph()
        order = ScmOrdenOperacion.query.filter_by(codigo="OF-OBS-001").one()
        fabricacion = ScmOrdenFabricacion(orden_operacion_id=order.id)
        db.session.add(fabricacion)
        db.session.flush()
        corrida = ScmCorridaFabricacion(
            orden_fabricacion_id=fabricacion.orden_operacion_id,
            codigo="C-OBS-SIN-MANGAS",
            secuencia=3,
            objetivo_neto_kg=Decimal("10"),
            estado="EN_EJECUCION",
        )
        db.session.add(corrida)
        db.session.commit()

        response = client.get(
            "/api/scm/v1/observabilidad/avance-of",
            headers={"X-Actor-Id": str(seeded["full"].id)},
        )

        assert response.status_code == 200
        item = next(
            item for item in response.get_json()["items"]
            if item["corrida_id"] == str(corrida.id)
        )
        assert item["objetivo_neto_kg"] == 10.0
        assert item["mangas"] == {"total": 0, "conocidas": 0}
        assert item["kg_medidos_efectivos"] is None
        assert item["coverage"]["estado"] == "INCOMPLETA"
        assert item["porcentaje"] is None


def test_progress_http_requires_ot_visibility_capability(app, client, scm_config):
    from test_scm_production_observability import _seed_observability_graph

    with app.app_context():
        seeded = _seed_observability_graph()
        response = client.get(
            "/api/scm/v1/observabilidad/avance-of",
            headers={"X-Actor-Id": str(seeded["denied"].id)},
        )

        assert response.status_code == 403
        assert response.get_json()["error"]["details"] == {"capability": "OT_VER"}


def test_tv_progress_http_requires_existing_ot_capability_without_writing(app, client, scm_config):
    from test_scm_production_observability import _seed_observability_graph

    with app.app_context():
        seeded = _seed_observability_graph()
        response = client.get(
            "/api/scm/v1/observabilidad/avance-of-tv",
            headers={"X-Actor-Id": str(seeded["denied"].id)},
        )
        assert response.status_code == 403
        assert response.get_json()["error"]["details"] == {"capability": "OT_VER"}
        assert not db.session.new
        assert not db.session.dirty


def test_tv_progress_http_uses_real_corrected_weighing_and_output_relation(app, client, scm_config):
    from test_scm_production_observability import _seed_observability_graph
    from app.models.scm_articulos import ScmArticulo
    from app.models.molde import Pieza
    from app.models.producto import PiezaColor
    from app.models.registro import RegistroDiarioProduccion
    from uuid import uuid4

    with app.app_context():
        app.config["SECRET_KEY"] = "tv-02-server-secret-for-tests"
        seeded = _seed_observability_graph()
        order = ScmOrdenOperacion.query.filter_by(codigo="OF-OBS-001").one()
        fabrication = ScmOrdenFabricacion(orden_operacion_id=order.id)
        db.session.add(fabrication)
        db.session.flush()
        corrida = ScmCorridaFabricacion(
            orden_fabricacion_id=order.id, codigo="C-OF-TV-REAL", secuencia=11,
            objetivo_neto_kg=Decimal("10"), estado="EN_EJECUCION",
        )
        piece = Pieza(codigo="PZ-TV-REAL", nombre="Cuerpo común", peso_nominal_gr=100)
        db.session.add_all([corrida, piece])
        db.session.flush()
        variant = PiezaColor(sku=f"PC-TV-{uuid4().hex[:12]}".upper(), pieza_id=piece.id, piezas=piece.nombre)
        db.session.add(variant)
        db.session.flush()
        from app.services.scm_article_service import _ensure_piece_article
        _ensure_piece_article(db.session, variant)
        article = ScmArticulo.query.filter_by(codigo=variant.sku).one()
        output = ScmOrdenOperacionSalida(
            orden_operacion_id=order.id, corrida_fabricacion_id=corrida.id,
            articulo_scm_id=article.id, cantidad_objetivo=100,
            kg_estandar_objetivo=10, excedente_objetivo=0,
        )
        db.session.add(output)
        db.session.flush()
        lot = ScmLoteArticulo(
            codigo="LOTE-TV-REAL", articulo_id=article.id,
            clase="SALIDA_ORDEN_OPERACION", orden_operacion_salida_id=output.id,
            cantidad_acreditada=1,
        )
        db.session.add(lot)
        blue = ScmTrabajoOt.query.filter_by(codigo="TC-OBS-AZUL").one()
        color_work = ScmTrabajoColor.query.filter_by(trabajo_ot_id=blue.id).one()
        color_work.corrida_fabricacion_id = corrida.id
        # TV-02 now requires a complete canonical card identity.  The seeded
        # graph predates the snapshot fields, so make this fixture explicit
        # instead of treating visual names as identity.
        color_work.molde_codigo_snapshot = "ML-OBS"
        color_work.color_id_snapshot = 1
        weighed_manga = ScmManga.query.filter_by(codigo="M-OT-OBS-FAB-02").one()
        weighed_manga.lote_articulo = lot
        weighing = ScmPesajeManga.query.filter_by(manga_id=weighed_manga.id).one()
        event_time = datetime(2026, 8, 10, 6, 0, tzinfo=timezone.utc)
        assignment = ScmAsignacionPersonalTrabajoOt(
            trabajo_ot_id=blue.id, trabajador_id=seeded["full"].id,
            estado="CERRADA", finalizada_at=event_time,
            asignada_por_id=seeded["full"].id,
            finalizada_por_id=seeded["full"].id,
        )
        db.session.add(assignment)
        db.session.flush()
        plan_assignment = ScmAsignacionPlanMangaOt.query.filter_by(
            plan_linea_id=weighed_manga.plan_linea_id,
            trabajo_ot_id=blue.id,
        ).first()
        if plan_assignment is None:
            plan_assignment = ScmAsignacionPlanMangaOt(
                plan_linea_id=weighed_manga.plan_linea_id,
                ot_id=weighed_manga.ot_id,
                trabajo_ot_id=blue.id,
                cantidad_asignada_un=95,
                mangas_asignadas=1,
                asignada_por_id=seeded["full"].id,
            )
            db.session.add(plan_assignment)
            db.session.flush()
        segment = ScmTramoMangaTrabajo(
            manga_id=weighed_manga.id, trabajo_ot_id=blue.id,
            asignacion_plan_id=plan_assignment.id,
            asignacion_personal_trabajo_id=assignment.id, secuencia=1,
            estado="CERRADO", cantidad_inicio_un=0, cantidad_fin_un=95,
            cantidad_atribuida_un=95, cantidad_inicio_kg=0,
            cantidad_fin_kg=Decimal("11.500"), cantidad_atribuida_kg=Decimal("11.500"),
            calidad_evidencia_kg="MEDIDA_DIRECTA_CIERRE_CONTROL",
            cerrada_at=event_time,
            motivo_cierre="PESAJE_FINAL",
            created_by_id=seeded["full"].id,
        )
        db.session.add(segment)
        db.session.flush()
        controls = []
        for index, net in enumerate((Decimal("10.000"), Decimal("11.000")), start=1):
            controls.append(ScmControlPesoManga(
                manga_id=weighed_manga.id, tramo_id=segment.id,
                operation_id=uuid4(), source_system="TEST",
                station_id="ST-OBS", capture_id=uuid4(), tipo="AVANCE_KG",
                peso_bruto_kg=net + Decimal("0.100"), tara_kg=Decimal("0.100"),
                peso_neto_kg=net, aporte_desde_control_anterior_kg=0,
                tara_fuente="TIPO_MANGA", unidad_evidencia="KG",
                calidad_evidencia="MEDIDA_DIRECTA", motivo=f"control {index}",
                pesado_at=event_time + timedelta(minutes=index),
                fecha_local_pesaje=date(2026, 8, 10), pesado_por_id=seeded["full"].id,
            ))
        db.session.add_all(controls)
        db.session.flush()
        from app.models.scm_catalogos import ScmCapacidad
        from app.models.trabajador import RolOperativo
        from app.services.scm_weighing_service import reopen_manga_after_accidental_close
        reopen_role = RolOperativo(
            codigo=f"R-REOPEN-{uuid4().hex[:8]}", nombre="Reapertura de prueba",
            capacidades=[ScmCapacidad.query.filter_by(codigo="MANGA_REABRIR").one()],
        )
        db.session.add(reopen_role)
        seeded["full"].roles.append(reopen_role)
        # A canonical close can finish in the terminal receiving state.  The
        # direct-close segment and append-only event must project the final
        # even though no open control is allowed in that state.
        weighed_manga.estado = "PENDIENTE_RECEPCION_ALMACEN"
        db.session.add(ScmEvento(
            aggregate_type="MANGA", aggregate_id=str(weighed_manga.id),
            tipo="KG_MANGA_CLOSED_FROM_LAST_CONTROL", actor_id=seeded["full"].id,
            actor_snapshot={"codigo": "OBS-FULL"}, motivo="cierre de control",
            before_json={}, after_json={"control_fuente": {"id": str(controls[-1].public_id)}},
            occurred_at=event_time + timedelta(minutes=2),
        ))
        db.session.commit()

        terminal_rows = production_reports_service._load_rows(
            db.session, production_reports_service._filters({}, require_dates=False),
            order_states={"EN_EJECUCION"}, exclude_annulled_runs=True,
        )
        terminal_manga = next(
            row["mangas"][weighed_manga.id]
            for row in terminal_rows
            if row["corrida"].id == corrida.id
        )
        assert terminal_manga._report_final_kg == Decimal("11.500")
        assert terminal_manga._report_open_kg is None
        assert terminal_manga._report_closure_event is not None

        # Reopen through the production service: it marks the persisted
        # segment ACTIVO but leaves its historical KG endpoint intact.
        reopened_response = reopen_manga_after_accidental_close(
            db.session,
            actor_id=seeded["full"].id,
            manga_id=weighed_manga.public_id,
            operation_id=uuid4(),
            data={
                "version": weighed_manga.version,
                "tipo_reapertura": "CONTINUAR_LLENADO",
                "motivo": "Continuar llenado con evidencia",
            },
        )
        assert reopened_response["manga"]["estado"] == "EN_LLENADO"

        reopened_rows = production_reports_service._load_rows(
            db.session, production_reports_service._filters({}, require_dates=False),
            order_states={"EN_EJECUCION"}, exclude_annulled_runs=True,
        )
        reopened_manga = next(
            row["mangas"][weighed_manga.id]
            for row in reopened_rows
            if row["corrida"].id == corrida.id
        )
        assert reopened_manga._report_final_kg is None
        assert reopened_manga._report_open_kg == Decimal("11.000")
        assert len(reopened_manga._report_controls) == 2
        assert reopened_manga._report_closure_event is None
        assert reopened_manga._report_all_segments[0].estado == "ACTIVO"
        assert reopened_manga._report_all_segments[0].cantidad_fin_kg == Decimal("11.500")

        reopened_tv_payload = client.get(
            "/api/scm/v1/observabilidad/avance-of-tv",
            headers={"X-Actor-Id": str(seeded["full"].id)},
        ).get_json()
        reopened_tv_order = next(item for item in reopened_tv_payload["items"] if item["of"] == "OF-OBS-001")
        reopened_tv_row = next(
            item for item in reopened_tv_order["salidas"]
            if item["corrida_id"] == str(corrida.id)
        )
        assert reopened_tv_row["finalizados_kg"] == 0
        assert reopened_tv_row["en_proceso_kg"] == pytest.approx(11.0)
        assert reopened_tv_row["avance"] == pytest.approx(11.0)

        reopened_standard_payload = client.get(
            "/api/scm/v1/observabilidad/avance-of",
            headers={"X-Actor-Id": str(seeded["full"].id)},
        ).get_json()
        reopened_standard_item = next(
            item for item in reopened_standard_payload["items"]
            if item["corrida_id"] == str(corrida.id)
        )
        assert reopened_standard_item["kg_finalizados_efectivos"] == pytest.approx(8.0)
        assert reopened_standard_item["kg_medidos_en_abiertas"] == pytest.approx(11.0)
        assert reopened_standard_item["kg_medidos_efectivos"] == pytest.approx(19.0)

        new_final = ScmPesajeManga(
            manga_id=weighed_manga.id, operation_id=uuid4(), source_system="TEST",
            station_id="ST-OBS", capture_id=uuid4(), peso_bruto_kg=Decimal("11.600"),
            tara_kg=Decimal("0.100"), peso_fisico_neto_kg=Decimal("11.500"),
            tara_fuente="TIPO_MANGA", cantidad_confirmada=Decimal("95"),
            fuente_cantidad="PLAN_CONFIRMADO_POR_PESAJE", estado="VIGENTE",
            kg_produccion_ot=Decimal("11.500"), pesada_at=event_time + timedelta(minutes=4),
            fecha_local_pesaje=date(2026, 8, 10), dias_desfase_operativo=0,
            pesado_por_id=seeded["full"].id, snapshots_json={}, created_at=event_time + timedelta(minutes=4),
        )
        db.session.add(new_final)
        weighed_manga.estado = "PENDIENTE_RECEPCION_ALMACEN"
        db.session.commit()

        response = client.get(
            "/api/scm/v1/observabilidad/avance-of-tv",
            headers={"X-Actor-Id": str(seeded["full"].id)},
        )

        assert response.status_code == 200
        payload = response.get_json()
        order_item = next(item for item in payload["items"] if item["of"] == "OF-OBS-001")
        row = next(item for item in order_item["salidas"] if item["corrida_id"] == str(corrida.id))
        assert (row["unidad"], row["tipo_meta"], row["meta"]) == ("KG", "NETA", 10)
        assert row["pieza_id"] == piece.id
        assert row["avance"] == pytest.approx(11.5)
        assert row["estado_avance"] == "SOBREPRODUCCION"
        standard_payload = client.get(
            "/api/scm/v1/observabilidad/avance-of",
            headers={"X-Actor-Id": str(seeded["full"].id)},
        ).get_json()
        standard_item = next(
            item for item in standard_payload["items"]
            if item["corrida_id"] == str(corrida.id)
        )
        assert standard_item["kg_finalizados_efectivos"] == pytest.approx(19.5)
        assert standard_item["kg_medidos_en_abiertas"] == 0
        assert standard_item["kg_medidos_efectivos"] == pytest.approx(19.5)
        assert standard_item["peso_fisico_kg"] == pytest.approx(19.5)
        assert standard_item["motivos_atribucion"] == []
        loaded = production_reports_service._load_rows(
            db.session, production_reports_service._filters({}, require_dates=False),
            order_states={"EN_EJECUCION"}, exclude_annulled_runs=True,
        )
        loaded_manga = next(
            row["mangas"][weighed_manga.id]
            for row in loaded
            if row["corrida"].id == corrida.id
        )
        assert loaded_manga._report_final_kg == Decimal("11.500")
        assert loaded_manga._report_open_kg is None
        assert len(loaded_manga._report_controls) == 2
        assert loaded_manga._report_closure_event is None
        assert len(loaded_manga.reaperturas) == 1
        assert db.session.get(ScmManga, weighed_manga.id).estado == "PENDIENTE_RECEPCION_ALMACEN"

        annulled_final = ScmAnulacionPesajeManga(
            pesaje_id=new_final.id,
            motivo="Anulación real de prueba",
            anulada_por_id=seeded["full"].id,
            anulada_at=event_time + timedelta(minutes=5),
            operation_id=uuid4(),
            cantidad_devuelta_plan_un=95,
            ot_reabierta=False,
        )
        db.session.add(annulled_final)
        db.session.commit()
        annulled_rows = production_reports_service._load_rows(
            db.session, production_reports_service._filters({}, require_dates=False),
            order_states={"EN_EJECUCION"}, exclude_annulled_runs=True,
        )
        annulled_manga = next(
            row["mangas"][weighed_manga.id]
            for row in annulled_rows
            if row["corrida"].id == corrida.id
        )
        assert annulled_manga._report_final_kg is None
        assert annulled_manga._report_open_kg is None


def test_progress_http_serializes_corrected_kg_before_ledger_reconciliation(app, client, scm_config):
    """ORM regression: corrected 11.9 reconciles a persisted 12.0 ledger."""
    from test_scm_production_observability import _seed_observability_graph
    from app.models.scm_articulos import ScmArticulo
    from app.models.molde import Pieza
    from app.models.producto import PiezaColor
    from uuid import uuid4

    with app.app_context():
        app.config["SECRET_KEY"] = "standard-corrected-ledger-secret"
        seeded = _seed_observability_graph()
        order = ScmOrdenOperacion.query.filter_by(codigo="OF-OBS-001").one()
        fabrication = ScmOrdenFabricacion(orden_operacion_id=order.id)
        db.session.add(fabrication)
        db.session.flush()
        corrida = ScmCorridaFabricacion(
            orden_fabricacion_id=order.id, codigo="C-STD-CORRECTED", secuencia=14,
            objetivo_neto_kg=Decimal("10"), estado="EN_EJECUCION",
        )
        piece = Pieza(codigo="PZ-STD-CORRECTED", nombre="Pieza corregida", peso_nominal_gr=100)
        db.session.add_all([corrida, piece])
        db.session.flush()
        variant = PiezaColor(
            sku=f"PC-STD-CORRECTED-{uuid4().hex[:10]}".upper(),
            pieza_id=piece.id, piezas=piece.nombre,
        )
        db.session.add(variant)
        db.session.flush()
        from app.services.scm_article_service import _ensure_piece_article
        _ensure_piece_article(db.session, variant)
        article = ScmArticulo.query.filter_by(codigo=variant.sku).one()
        output = ScmOrdenOperacionSalida(
            orden_operacion_id=order.id, corrida_fabricacion_id=corrida.id,
            articulo_scm_id=article.id, cantidad_objetivo=100,
            kg_estandar_objetivo=10, excedente_objetivo=0,
        )
        db.session.add(output)
        db.session.flush()
        lot = ScmLoteArticulo(
            codigo="LOTE-STD-CORRECTED", articulo_id=article.id,
            clase="SALIDA_ORDEN_OPERACION", orden_operacion_salida_id=output.id,
            cantidad_acreditada=1,
        )
        db.session.add(lot)
        blue = ScmTrabajoOt.query.filter_by(codigo="TC-OBS-AZUL").one()
        color_work = ScmTrabajoColor.query.filter_by(trabajo_ot_id=blue.id).one()
        color_work.corrida_fabricacion_id = corrida.id
        color_work.molde_codigo_snapshot = "ML-OBS"
        color_work.color_id_snapshot = 1
        weighed_manga = ScmManga.query.filter_by(codigo="M-OT-OBS-FAB-02").one()
        weighed_manga.lote_articulo = lot
        weighing = ScmPesajeManga.query.filter_by(manga_id=weighed_manga.id).one()
        from app.models.scm_ot import ScmCorreccionPesajeManga
        correction = ScmCorreccionPesajeManga.query.filter_by(pesaje_id=weighing.id).one()
        correction.result_projection_json = {
            **(correction.result_projection_json or {}),
            "peso_fisico_neto_kg": "11.900",
            "kg_produccion_ot": "11.900",
        }
        event_time = datetime(2026, 8, 10, 6, 0, tzinfo=timezone.utc)
        assignment = ScmAsignacionPersonalTrabajoOt(
            trabajo_ot_id=blue.id, trabajador_id=seeded["full"].id,
            estado="CERRADA", finalizada_at=event_time,
            asignada_por_id=seeded["full"].id, finalizada_por_id=seeded["full"].id,
        )
        db.session.add(assignment)
        db.session.flush()
        plan_assignment = ScmAsignacionPlanMangaOt.query.filter_by(
            plan_linea_id=weighed_manga.plan_linea_id,
            trabajo_ot_id=blue.id,
        ).first()
        if plan_assignment is None:
            plan_assignment = ScmAsignacionPlanMangaOt(
                plan_linea_id=weighed_manga.plan_linea_id,
                ot_id=weighed_manga.ot_id, trabajo_ot_id=blue.id,
                cantidad_asignada_un=95, mangas_asignadas=1,
                asignada_por_id=seeded["full"].id,
            )
            db.session.add(plan_assignment)
            db.session.flush()
        segment = ScmTramoMangaTrabajo(
            manga_id=weighed_manga.id, trabajo_ot_id=blue.id,
            asignacion_plan_id=plan_assignment.id,
            asignacion_personal_trabajo_id=assignment.id, secuencia=1,
            estado="CERRADO", cantidad_inicio_un=0, cantidad_fin_un=95,
            cantidad_atribuida_un=95, cantidad_inicio_kg=0,
            cantidad_fin_kg=Decimal("12.000"), cantidad_atribuida_kg=Decimal("12.000"),
            calidad_evidencia_kg="MEDIDA_DIRECTA", cerrada_at=event_time,
            created_by_id=seeded["full"].id,
        )
        db.session.add(segment)
        db.session.flush()
        control = ScmControlPesoManga(
            manga_id=weighed_manga.id, tramo_id=segment.id,
            operation_id=uuid4(), source_system="TEST", station_id="ST-OBS",
            capture_id=uuid4(), tipo="AVANCE_KG", peso_bruto_kg=Decimal("11.100"),
            tara_kg=Decimal("0.100"), peso_neto_kg=Decimal("11.000"),
            aporte_desde_control_anterior_kg=0, tara_fuente="TIPO_MANGA",
            unidad_evidencia="KG", calidad_evidencia="MEDIDA_DIRECTA",
            motivo="control corregido", pesado_at=event_time,
            fecha_local_pesaje=date(2026, 8, 10), pesado_por_id=seeded["full"].id,
        )
        db.session.add(control)
        db.session.commit()

        rows = production_reports_service._load_rows(
            db.session, production_reports_service._filters({}, require_dates=False),
            order_states={"EN_EJECUCION"}, exclude_annulled_runs=True,
        )
        run = next(row for row in rows if row["corrida"].id == corrida.id)
        loaded_manga = run["mangas"][weighed_manga.id]
        assert loaded_manga._report_final_kg == Decimal("11.900")
        assert loaded_manga._report_segments[-1].cantidad_fin_kg == Decimal("12.000")
        isolated = {**run, "mangas": {weighed_manga.id: loaded_manga}}
        standard_values = _run_manga_values(isolated)
        assert standard_values[:3] == (Decimal("11.900"), Decimal("0"), Decimal("11.900"))
        tv_values = production_reports_service._progress_tv_breakdown(isolated, article.id, "KG")
        assert tv_values["finalizados"] == Decimal("11.900")
        assert tv_values["medido"] == Decimal("11.900")

        standard_payload = client.get(
            "/api/scm/v1/observabilidad/avance-of",
            headers={"X-Actor-Id": str(seeded["full"].id)},
        ).get_json()
        standard = next(item for item in standard_payload["items"] if item["corrida_id"] == str(corrida.id))
        # The seeded OT has one separate 8 kg manga in the standard OT scope;
        # the isolated assertion proves the corrected manga contributes 11.9.
        assert standard["kg_medidos_efectivos"] == pytest.approx(19.9)
        assert standard["peso_fisico_kg"] == pytest.approx(19.9)
        assert standard["kg_medidos_efectivos"] != pytest.approx(20.0)

        tv_payload = client.get(
            "/api/scm/v1/observabilidad/avance-of-tv",
            headers={"X-Actor-Id": str(seeded["full"].id)},
        ).get_json()
        tv_order = next(item for item in tv_payload["items"] if item["of"] == "OF-OBS-001")
        tv = next(item for item in tv_order["salidas"] if item["corrida_id"] == str(corrida.id))
        assert tv["avance"] == pytest.approx(11.9)
        assert tv["finalizados_kg"] == pytest.approx(11.9)
        assert tv["peso_fisico_kg"] == pytest.approx(11.9)
        assert tv["avance"] != pytest.approx(12.0)

        # A canonical close event remains sufficient without the optional
        # direct-close marker, including after the manga enters terminal
        # receiving state.  There is no reopening in this branch.
        weighing.estado = "REABIERTO"
        weighed_manga.estado = "PENDIENTE_RECEPCION_ALMACEN"
        db.session.add(ScmEvento(
            aggregate_type="MANGA", aggregate_id=str(weighed_manga.id),
            tipo="KG_MANGA_CLOSED_FROM_LAST_CONTROL", actor_id=seeded["full"].id,
            actor_snapshot={"codigo": "OBS-FULL"}, motivo="cierre de control sin marcador",
            before_json={}, after_json={"control_fuente": {"id": str(control.public_id)}},
            occurred_at=datetime(2026, 8, 10, 7, 0, tzinfo=timezone.utc),
        ))
        db.session.commit()
        no_marker_rows = production_reports_service._load_rows(
            db.session, production_reports_service._filters({}, require_dates=False),
            order_states={"EN_EJECUCION"}, exclude_annulled_runs=True,
        )
        no_marker_manga = next(
            row["mangas"][weighed_manga.id]
            for row in no_marker_rows
            if row["corrida"].id == corrida.id
        )
        assert no_marker_manga._report_final_kg == Decimal("11.000")
        assert no_marker_manga._report_open_kg is None


def test_progress_http_serializes_physical_separately_for_real_invalid_ledger(app, client, scm_config):
    from test_scm_production_observability import _manga, _seed_observability_graph, _weighing
    from app.models.molde import Pieza
    from app.models.producto import PiezaColor
    from app.models.registro import RegistroDiarioProduccion
    from app.models.scm_articulos import ScmArticulo
    from app.services.scm_article_service import _ensure_piece_article

    with app.app_context():
        app.config["SECRET_KEY"] = "standard-physical-invalid-ledger-secret"
        seeded = _seed_observability_graph()
        order = ScmOrdenOperacion.query.filter_by(codigo="OF-OBS-001").one()
        fabrication = ScmOrdenFabricacion(orden_operacion_id=order.id)
        db.session.add(fabrication)
        db.session.flush()
        corrida = ScmCorridaFabricacion(
            orden_fabricacion_id=fabrication.orden_operacion_id, codigo="C-STD-PHYSICAL", secuencia=13,
            objetivo_neto_kg=Decimal("10"), estado="EN_EJECUCION",
        )
        piece = Pieza(codigo="PZ-STD-PHYSICAL", nombre="Pieza física", peso_nominal_gr=100)
        db.session.add_all([corrida, piece])
        db.session.flush()
        variant = PiezaColor(
            sku=f"PC-STD-PHYSICAL-{uuid4().hex[:10]}".upper(),
            pieza_id=piece.id, piezas=piece.nombre,
        )
        db.session.add(variant)
        db.session.flush()
        _ensure_piece_article(db.session, variant)
        article = ScmArticulo.query.filter_by(codigo=variant.sku).one()
        output = ScmOrdenOperacionSalida(
            orden_operacion_id=order.id, corrida_fabricacion_id=corrida.id,
            articulo_scm_id=article.id, cantidad_objetivo=100,
            kg_estandar_objetivo=10, excedente_objetivo=0,
        )
        db.session.add(output)
        db.session.flush()
        lot = ScmLoteArticulo(
            codigo="LOTE-STD-PHYSICAL", articulo_id=article.id,
            clase="SALIDA_ORDEN_OPERACION", orden_operacion_salida_id=output.id,
            cantidad_acreditada=1,
        )
        db.session.add(lot)
        new_ot = RegistroDiarioProduccion(
            codigo_ot="OT-STD-PHYSICAL", codigo_ot_sintetico=False,
            estado="EN_EJECUCION", tipo_ot="FABRICACION", fecha=date(2026, 8, 10),
            turno="DIA", maquina_id=1, responsable_id=seeded["full"].id,
            maquinista_previsto_id=seeded["full"].id, created_by_id=seeded["full"].id,
            orden_operacion_id=order.id,
        )
        db.session.add(new_ot)
        db.session.flush()
        work = ScmTrabajoOt(
            orden_trabajo_id=new_ot.id, codigo="TC-STD-PHYSICAL", secuencia=1,
            estado="EN_EJECUCION", orden_operacion_id=order.id,
            cantidad_objetivo_un=100, cantidad_confirmada_un=0,
            created_by_id=seeded["full"].id,
        )
        db.session.add(work)
        db.session.flush()
        db.session.add(ScmTrabajoColor(
            trabajo_ot_id=work.id, corrida_fabricacion_id=corrida.id,
            color_nombre_snapshot="AZUL", molde_codigo_snapshot="ML-STD",
        ))
        weighed_manga = _manga(
            ot=new_ot, work=work, sequence=1, state="PESADA",
            worker=seeded["full"], confirmed=95,
        )
        db.session.add(weighed_manga)
        db.session.flush()
        weighed_manga.lote_articulo = lot
        _weighing(
            manga=weighed_manga, worker=seeded["full"], net="11.500",
            standard="11.500", weighed_at=datetime(2026, 8, 10, 6, 0, tzinfo=timezone.utc),
        )
        event_time = datetime(2026, 8, 10, 6, 0, tzinfo=timezone.utc)
        assignment = ScmAsignacionPersonalTrabajoOt(
            trabajo_ot_id=work.id, trabajador_id=seeded["full"].id,
            estado="CERRADA", finalizada_at=event_time,
            asignada_por_id=seeded["full"].id, finalizada_por_id=seeded["full"].id,
        )
        db.session.add(assignment)
        db.session.flush()
        db.session.add(ScmTramoMangaTrabajo(
            manga_id=weighed_manga.id, trabajo_ot_id=work.id,
            asignacion_personal_trabajo_id=assignment.id, secuencia=1,
            estado="CERRADO", cantidad_inicio_un=0, cantidad_fin_un=95,
            cantidad_atribuida_un=95, cantidad_inicio_kg=0,
            cantidad_fin_kg=Decimal("10.000"), cantidad_atribuida_kg=Decimal("10.000"),
            calidad_evidencia_kg="PENDIENTE", cerrada_at=event_time,
            created_by_id=seeded["full"].id,
        ))
        db.session.commit()

        standard = client.get(
            "/api/scm/v1/observabilidad/avance-of",
            headers={"X-Actor-Id": str(seeded["full"].id)},
        ).get_json()
        item = next(item for item in standard["items"] if item["corrida_id"] == str(corrida.id))
        assert item["kg_finalizados_efectivos"] is None
        assert item["kg_medidos_en_abiertas"] is None
        assert item["kg_medidos_efectivos"] is None
        assert item["avance_medido_kg"] is None
        assert item["subtotal_conocido_kg"] is None
        assert item["peso_fisico_kg"] == pytest.approx(11.5)
        assert item["motivos_atribucion"] == ["ATRIBUCION_INCOMPLETA"]
        assert "ATRIBUCION_INCOMPLETA" in item["coverage"]["motivos"]

        tv = client.get(
            "/api/scm/v1/observabilidad/avance-of-tv",
            headers={"X-Actor-Id": str(seeded["full"].id)},
        ).get_json()
        tv_item = next(item for item in tv["items"] if item["corrida_id"] == str(corrida.id))
        assert tv_item["salidas"][0]["avance"] is None
        assert tv_item["salidas"][0]["peso_fisico_kg"] == pytest.approx(11.5)
        assert tv_item["salidas"][0]["motivos_atribucion"] == ["ATRIBUCION_INCOMPLETA"]


def test_tv_progress_http_consolidates_real_multi_ot_contexts(app, client, scm_config):
    """Endpoint/DB regression: one corrida with two OTs yields one card."""
    from test_scm_production_observability import _seed_observability_graph
    from app.models.scm_articulos import ScmArticulo
    from app.models.molde import Pieza
    from app.models.producto import PiezaColor
    from app.models.registro import RegistroDiarioProduccion
    from uuid import uuid4

    with app.app_context():
        app.config["SECRET_KEY"] = "tv-02-server-secret-for-tests"
        seeded = _seed_observability_graph()
        order = ScmOrdenOperacion.query.filter_by(codigo="OF-OBS-001").one()
        fabrication = ScmOrdenFabricacion(orden_operacion_id=order.id)
        db.session.add(fabrication)
        db.session.flush()
        corrida = ScmCorridaFabricacion(
            orden_fabricacion_id=order.id, codigo="C-OF-TV-MULTI-OT", secuencia=12,
            objetivo_neto_kg=Decimal("10"), estado="EN_EJECUCION",
        )
        piece = Pieza(codigo="PZ-TV-MULTI", nombre="Pieza multi OT", peso_nominal_gr=100)
        db.session.add_all([corrida, piece])
        db.session.flush()
        variant = PiezaColor(sku=f"PC-TV-MULTI-{uuid4().hex[:10]}".upper(), pieza_id=piece.id, piezas=piece.nombre)
        db.session.add(variant)
        db.session.flush()
        from app.services.scm_article_service import _ensure_piece_article
        _ensure_piece_article(db.session, variant)
        article = ScmArticulo.query.filter_by(codigo=variant.sku).one()
        output = ScmOrdenOperacionSalida(
            orden_operacion_id=order.id, corrida_fabricacion_id=corrida.id,
            articulo_scm_id=article.id, cantidad_objetivo=100,
            kg_estandar_objetivo=10, excedente_objetivo=0,
        )
        db.session.add(output)
        db.session.flush()
        lot = ScmLoteArticulo(
            codigo="LOTE-TV-MULTI", articulo_id=article.id,
            clase="SALIDA_ORDEN_OPERACION", orden_operacion_salida_id=output.id,
            cantidad_acreditada=1,
        )
        db.session.add(lot)
        red = ScmTrabajoOt.query.filter_by(codigo="TC-OBS-ROJO").one()
        blue = ScmTrabajoOt.query.filter_by(codigo="TC-OBS-AZUL").one()
        for work in (red, blue):
            color_work = ScmTrabajoColor.query.filter_by(trabajo_ot_id=work.id).one()
            color_work.corrida_fabricacion_id = corrida.id
            color_work.molde_codigo_snapshot = "ML-OBS"
            color_work.color_id_snapshot = 1
        second_ot = RegistroDiarioProduccion(
            codigo_ot="OT-OBS-FAB-2", codigo_ot_sintetico=False,
            estado="EN_EJECUCION", tipo_ot="FABRICACION", fecha=date(2026, 8, 10),
            turno="DIA", maquina_id=1, responsable_id=seeded["full"].id,
            maquinista_previsto_id=seeded["full"].id, created_by_id=seeded["full"].id,
        )
        db.session.add(second_ot)
        db.session.flush()
        second_work = ScmTrabajoOt(
            orden_trabajo_id=second_ot.id, codigo="TC-OBS-VERDE",
            secuencia=1, estado="EN_EJECUCION", orden_operacion_id=order.id,
            cantidad_objetivo_un=100, cantidad_confirmada_un=0,
            created_by_id=seeded["full"].id,
        )
        db.session.add(second_work)
        db.session.flush()
        db.session.add(ScmTrabajoColor(
            trabajo_ot_id=second_work.id, corrida_fabricacion_id=corrida.id,
            color_nombre_snapshot="AZUL", molde_codigo_snapshot="ML-OBS",
            color_id_snapshot=1,
        ))
        weighed_manga = ScmManga.query.filter_by(codigo="M-OT-OBS-FAB-02").one()
        weighed_manga.lote_articulo = lot
        db.session.commit()

        response = client.get(
            "/api/scm/v1/observabilidad/avance-of-tv",
            headers={"X-Actor-Id": str(seeded["full"].id)},
        )

        assert response.status_code == 200
        payload = response.get_json()
        cards = [item for item in payload["items"] if item["corrida_id"] == str(corrida.id)]
        assert len(cards) == 1
        assert cards[0]["group_id"].startswith("tvgrp-") or cards[0]["group_id"] == f"{order.id}|{corrida.id}|ML-OBS|1"
        assert len(cards[0]["ots"]) == 2
        assert len(cards[0]["salidas"]) == 1


def test_tv_loader_preserves_shared_un_segments_for_distinct_work_owners(app, client, scm_config):
    from test_scm_production_observability import _seed_observability_graph
    from app.models.scm_articulos import ScmArticulo
    from app.models.scm_ot import ScmAsignacionPersonalTrabajoOt, ScmTramoMangaTrabajo

    with app.app_context():
        app.config["SECRET_KEY"] = "tv-02-server-secret-for-tests"
        seeded = _seed_observability_graph()
        order = ScmOrdenOperacion.query.filter_by(codigo="OF-OBS-001").one()
        db.session.add(ScmOrdenFabricacion(orden_operacion_id=order.id))
        db.session.flush()
        run_a = ScmCorridaFabricacion(
            orden_fabricacion_id=order.id, codigo="C-TV-UN-A", secuencia=21, estado="EN_EJECUCION"
        )
        run_b = ScmCorridaFabricacion(
            orden_fabricacion_id=order.id, codigo="C-TV-UN-B", secuencia=22, estado="EN_EJECUCION"
        )
        article = ScmArticulo(codigo="WIP-TV-UN", nombre="Salida histórica UN", clase="SUBENSAMBLE_WIP", unidad_inventario="KG")
        db.session.add_all([run_a, run_b, article])
        db.session.flush()
        output_a = ScmOrdenOperacionSalida(
            orden_operacion_id=order.id, corrida_fabricacion_id=run_a.id, articulo_scm_id=article.id,
            cantidad_objetivo=4, kg_estandar_objetivo=None, excedente_objetivo=0,
        )
        output_b = ScmOrdenOperacionSalida(
            orden_operacion_id=order.id, corrida_fabricacion_id=run_b.id, articulo_scm_id=article.id,
            cantidad_objetivo=5, kg_estandar_objetivo=None, excedente_objetivo=0,
        )
        db.session.add_all([output_a, output_b])
        db.session.flush()
        lot = ScmLoteArticulo(
            codigo="LOTE-TV-UN", articulo_id=article.id, clase="SALIDA_ORDEN_OPERACION",
            orden_operacion_salida_id=output_a.id, cantidad_acreditada=1,
        )
        db.session.add(lot)
        red = ScmTrabajoOt.query.filter_by(codigo="TC-OBS-ROJO").one()
        blue = ScmTrabajoOt.query.filter_by(codigo="TC-OBS-AZUL").one()
        ScmTrabajoColor.query.filter_by(trabajo_ot_id=red.id).one().corrida_fabricacion_id = run_a.id
        ScmTrabajoColor.query.filter_by(trabajo_ot_id=blue.id).one().corrida_fabricacion_id = run_b.id
        manga = ScmManga.query.filter_by(codigo="M-OT-OBS-FAB-02").one()
        manga.lote_articulo = lot
        weighing = ScmPesajeManga.query.filter_by(manga_id=manga.id).one()
        from app.models.scm_ot import ScmCorreccionPesajeManga
        correction = ScmCorreccionPesajeManga.query.filter_by(pesaje_id=weighing.id).one()
        correction.result_projection_json = {**correction.result_projection_json, "cantidad_confirmada": "9.000", "peso_fisico_neto_kg": "1.800"}
        assignments = []
        for work in (red, blue):
            assignment = ScmAsignacionPersonalTrabajoOt(
                trabajo_ot_id=work.id, trabajador_id=seeded["full"].id, estado="CERRADA",
                finalizada_at=weighing.pesada_at, asignada_por_id=seeded["full"].id,
                finalizada_por_id=seeded["full"].id,
            )
            assignments.append(assignment)
        db.session.add_all(assignments)
        db.session.flush()
        first_segment = ScmTramoMangaTrabajo(
            manga_id=manga.id, trabajo_ot_id=red.id, asignacion_personal_trabajo_id=assignments[0].id,
            secuencia=1, estado="CERRADO", cantidad_inicio_un=0, cantidad_fin_un=4,
            cantidad_atribuida_un=4, created_by_id=seeded["full"].id,
        )
        second_segment = ScmTramoMangaTrabajo(
            manga_id=manga.id, trabajo_ot_id=blue.id, asignacion_personal_trabajo_id=assignments[1].id,
            secuencia=2, estado="CERRADO", cantidad_inicio_un=4, cantidad_fin_un=9,
            cantidad_atribuida_un=5, created_by_id=seeded["full"].id,
        )
        db.session.add_all([first_segment, second_segment])
        db.session.flush()
        target_plan = ScmAsignacionPlanMangaOt(
            plan_linea_id=manga.plan_linea_id, ot_id=red.orden_trabajo_id, trabajo_ot_id=blue.id,
            cantidad_asignada_un=100, mangas_asignadas=1,
            asignada_por_id=seeded["full"].id,
        )
        db.session.add(target_plan)
        db.session.flush()
        db.session.add(ScmCorreccionAsignacionManga(
            manga_id=manga.id, origen_ot_id=red.orden_trabajo_id, origen_trabajo_ot_id=red.id,
            destino_ot_id=blue.orden_trabajo_id, destino_trabajo_ot_id=blue.id,
            origen_asignacion_id=assignments[0].id, destino_asignacion_id=assignments[1].id,
            destino_asignacion_plan_id=target_plan.id, tramo_objetivo_id=first_segment.id,
            manga_version_antes=manga.version, manga_version_despues=manga.version + 1,
            motivo="Corrección de continuidad hacia trabajo destino", evidencia_json={"source": "test"},
            actor_id=seeded["full"].id, operation_id=uuid4(),
        ))
        db.session.commit()

        # Exercise the actual SQLAlchemy loader: the legacy KG-only segment
        # collection must stay empty while the TV ledger keeps both UN spans.
        loaded = production_reports_service._load_rows(
            db.session, production_reports_service._filters({}, require_dates=False),
            order_states={"EN_EJECUCION"}, exclude_annulled_runs=True,
        )
        rows = {row["corrida"].codigo: row for row in loaded}
        manga_a = rows["C-TV-UN-A"]["mangas"][manga.id]
        manga_b = rows["C-TV-UN-B"]["mangas"][manga.id]
        assert manga_a._report_segments == manga_b._report_segments == []
        assert len(manga_a._report_all_segments) == len(manga_b._report_all_segments) == 2
        assert production_reports_service._progress_tv_actual(rows["C-TV-UN-A"], article.id, "UN") == (None, False)
        assert production_reports_service._progress_tv_actual(rows["C-TV-UN-B"], article.id, "UN") == (Decimal("9"), False)
        assert production_reports_service._progress_tv_actual(rows["C-TV-UN-A"], article.id, "KG") == (None, False)
        assert production_reports_service._progress_tv_actual(rows["C-TV-UN-B"], article.id, "KG") == (None, True)
        origin_kg = production_reports_service._progress_tv_breakdown(rows["C-TV-UN-A"], article.id, "KG")
        destination_kg = production_reports_service._progress_tv_breakdown(rows["C-TV-UN-B"], article.id, "KG")
        assert origin_kg["mangas_total"] == 0
        assert destination_kg["mangas_total"] == 1

        standard_payload = client.get(
            "/api/scm/v1/observabilidad/avance-of",
            headers={"X-Actor-Id": str(seeded["full"].id)},
        ).get_json()
        standard_rows = {
            item["corrida_id"]: item for item in standard_payload["items"]
            if item["corrida_id"] in {str(run_a.id), str(run_b.id)}
        }
        assert standard_rows[str(run_a.id)]["kg_medidos_efectivos"] is None
        assert standard_rows[str(run_b.id)]["kg_medidos_efectivos"] == pytest.approx(8.0)

        tv_payload = client.get(
            "/api/scm/v1/observabilidad/avance-of-tv",
            headers={"X-Actor-Id": str(seeded["full"].id)},
        ).get_json()
        tv_outputs = {
            output["corrida_id"]: output
            for item in tv_payload["items"]
            for output in item["salidas"]
            if output["corrida_id"] in {str(run_a.id), str(run_b.id)}
        }
        assert tv_outputs[str(run_a.id)]["avance"] is None
        assert tv_outputs[str(run_b.id)]["avance"] is None
        assert tv_outputs[str(run_a.id)]["mangas_total"] == 0
        assert tv_outputs[str(run_b.id)]["mangas_total"] == 1


def test_progress_http_hides_weights_without_manga_visibility(app, client, scm_config):
    from test_scm_production_observability import _seed_observability_graph

    with app.app_context():
        seeded = _seed_observability_graph()
        order = ScmOrdenOperacion.query.filter_by(codigo="OF-OBS-001").one()
        fabricacion = ScmOrdenFabricacion(orden_operacion_id=order.id)
        db.session.add(fabricacion)
        db.session.flush()
        corrida = ScmCorridaFabricacion(
            orden_fabricacion_id=fabricacion.orden_operacion_id,
            codigo="C-OBS-RESTRINGIDA",
            secuencia=3,
            objetivo_neto_kg=Decimal("10"),
            estado="EN_EJECUCION",
        )
        db.session.add(corrida)
        db.session.commit()

        response = client.get(
            "/api/scm/v1/observabilidad/avance-of",
            headers={"X-Actor-Id": str(seeded["base"].id)},
        )

        assert response.status_code == 200
        payload = response.get_json()
        item = next(item for item in payload["items"] if item["corrida_id"] == str(corrida.id))
        assert payload["visibilidad"] == {
            "pesaje": False,
            "restriccion": "MANGA_PESAJE_VER requerido para ver pesos",
        }
        assert item["kg_finalizados_efectivos"] is None
        assert item["kg_medidos_en_abiertas"] is None
        assert item["kg_medidos_efectivos"] is None
        assert item["porcentaje"] is None
        assert item["coverage"]["estado"] == "INCOMPLETA"
        assert item["coverage"]["motivos"] == ["MANGA_PESAJE_VER_REQUERIDO"]


def test_progress_http_reports_complete_known_weights(app, client, scm_config):
    from test_scm_production_observability import _seed_observability_graph

    with app.app_context():
        seeded = _seed_observability_graph()
        order = ScmOrdenOperacion.query.filter_by(codigo="OF-OBS-001").one()
        fabricacion = ScmOrdenFabricacion(orden_operacion_id=order.id)
        db.session.add(fabricacion)
        db.session.flush()
        corrida = ScmCorridaFabricacion(
            orden_fabricacion_id=fabricacion.orden_operacion_id,
            codigo="C-OBS-COMPLETA",
            secuencia=3,
            objetivo_neto_kg=Decimal("30"),
            estado="EN_EJECUCION",
        )
        db.session.add(corrida)
        db.session.flush()
        blue = ScmTrabajoOt.query.filter_by(codigo="TC-OBS-AZUL").one()
        color_work = ScmTrabajoColor.query.filter_by(trabajo_ot_id=blue.id).one()
        color_work.corrida_fabricacion_id = corrida.id
        pending = ScmManga.query.filter_by(codigo="M-OT-OBS-FAB-01").one()
        from test_scm_production_observability import NOW, _weighing

        _weighing(
            manga=pending,
            worker=Trabajador.query.filter_by(codigo="TRB-01").one(),
            net="0.500",
            standard="0.500",
            weighed_at=NOW,
        )
        annulled = next(
            manga for manga in blue.mangas if manga.codigo.endswith("-04")
        )
        annulled.estado = "RECIBIDA"
        pesaje = ScmPesajeManga.query.filter_by(manga_id=annulled.id).one()
        db.session.delete(ScmAnulacionPesajeManga.query.filter_by(pesaje_id=pesaje.id).one())
        db.session.commit()

        response = client.get(
            "/api/scm/v1/observabilidad/avance-of",
            headers={"X-Actor-Id": str(seeded["full"].id)},
        )

        assert response.status_code == 200
        item = next(item for item in response.get_json()["items"] if item["corrida_id"] == str(corrida.id))
        # The OT also contains a red work; its manga must not be copied into
        # the blue objective merely because both works share the OT.
        assert item["mangas"] == {"total": 3, "conocidas": 3}
        assert item["kg_medidos_efectivos"] == 29.5
        assert item["coverage"]["estado"] == "COMPLETA"
        assert item["porcentaje"] == pytest.approx(98.3333333333)
